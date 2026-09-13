"""Portable Excel export from validated solver arrays; no external runtime."""
from copy import copy
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


def write_template(root, question, times, fields):
    root = Path(root)
    template = root / 'data' / '附件' / '附件3' / f'result{question}.xlsx'
    wb = load_workbook(template)
    if len(wb.worksheets) != len(fields):
        raise ValueError('Template sheet count does not match the output fields')
    times = np.asarray(times)
    if times.ndim != 1 or not np.isfinite(times).all() or not (np.diff(times) > 0).all():
        raise ValueError('Invalid output time sequence')
    for ws, values in zip(wb.worksheets, fields):
        values = np.asarray(values)
        expected = 22 if question == 4 else 21
        if values.shape != (len(times), expected):
            raise ValueError('Output shape does not match the requested radial grid')
        if np.isinf(values).any() or (question != 4 and np.isnan(values).any()):
            raise ValueError('Unexpected nonfinite output')
        header = ws.cell(1, 1).value
        surface = ws.cell(1, ws.max_column).value
        number_style = copy(ws.cell(2, 2)._style)
        header_style = copy(ws.cell(1, 2)._style)
        ws.delete_rows(1, ws.max_row)
        ws.append([header] + [j / 10 for j in range(21)] + ([surface] if question == 4 else []))
        for cell in ws[1]:
            cell._style = copy(header_style)
        for t, row in zip(times, values):
            ws.append([float(t)] + [None if np.isnan(v) else float(v) for v in row])
        for row in ws.iter_rows(min_row=2, min_col=2):
            for cell in row:
                cell._style = copy(number_style)
                cell.number_format = '0.0000'
        ws.freeze_panes = 'B2'
    wb.properties.creator = 'Team'
    wb.properties.lastModifiedBy = 'Team'
    dest = root / 'results' / f'result{question}.xlsx'
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    wb.close()
    print(f'Excel saved: results/result{question}.xlsx', flush=True)


def export_saved(root, question):
    root = Path(root)
    validation = json.loads((root / 'results' / f'q{question}_validation.json').read_text(encoding='utf-8'))
    if question == 4:
        if validation.get('pass') is not True or not validation.get('checks') or not all(v is True for v in validation['checks'].values()):
            raise RuntimeError('Q4 numerical validation did not pass')
    else:
        from delivery import require_delivery
        require_delivery(validation, root / 'results' / f'q{question}_export_validation.json')
        if question == 3:
            from q3_closeout import require_q3_delivery
            require_q3_delivery(validation, root / 'results' / 'q3_export_validation.json')
    with np.load(root / 'results' / f'q{question}_solution.npz') as z:
        if question in (1, 2):
            write_template(root, question, z['times'], [z['T_out'], z['C_out']])
        elif question == 3:
            write_template(root, question, z['excel_times'], [z['C_out']])
        else:
            write_template(root, question, z['times_s'][1:], [z['C_physical'][1:]])
