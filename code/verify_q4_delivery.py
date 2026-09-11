"""Read-only exhaustive audit of the exported Q4 workbook."""
import hashlib
import json
from pathlib import Path
import numpy as np
from openpyxl import load_workbook
from problem4 import OUT
from utils import RESULTS_DIR

payload=json.loads((OUT/'excel_payload.json').read_text(encoding='utf-8'))
path=RESULTS_DIR/'result4.xlsx'
wb=load_workbook(path,read_only=True,data_only=True)
assert wb.sheetnames==['Sheet1'],wb.sheetnames
ws=wb.active
rows=list(ws.iter_rows())
assert len(rows)==len(payload['rows'])+1
actual_columns=max(map(len,rows))
assert actual_columns==len(payload['headers'])
assert all(len(row)==actual_columns for row in rows)
assert [c.value for c in rows[0]]==payload['headers']
max_error=0.;blanks=0;numeric=0
for row,expected in zip(rows[1:],payload['rows']):
    for j,(cell,value) in enumerate(zip(row,expected)):
        if value is None:
            assert cell.value is None,(cell.coordinate,cell.value)
            blanks+=1
        else:
            assert isinstance(cell.value,(int,float)),(cell.coordinate,cell.value)
            max_error=max(max_error,abs(cell.value-value))
            numeric+=1
            if j>0:
                assert cell.number_format=='0.0000',(cell.coordinate,cell.number_format)
assert max_error<1e-10,max_error
wb.close()
source=np.load(RESULTS_DIR/'q4_solution.npz')
times=source['times_s'][1:]
assert np.array_equal(times[:-1],np.arange(60.,times[-1],60.))
assert source['C_physical'].shape[0]==len(times)+1
n=len(source['xi'])
imax=np.argmax(source['Y'][:,n:],axis=1)
center_gap=float(np.max(np.max(source['Y'][:,n:],axis=1)-source['Y'][:,n]))
assert center_gap<1e-10,center_gap
audit={'pass':True,'sheet':'Sheet1','rows_including_header':len(rows),'columns':actual_columns,
       'numeric_cells':numeric,'domain_blank_cells':blanks,'max_export_abs_error':max_error,
       'initial_state_only_in_source':True,'final_time_s':float(times[-1]),
       'max_position_is_center_all_saved_times':bool(np.all(imax==0)),
       'max_center_gap_all_saved_times':center_gap,
       'event_max_position_xi':float(source['xi'][np.argmax(source['critical_state'][n:])]),
       'report_max_position_xi':float(source['xi'][np.argmax(source['report_state'][n:])]),
       'sha256':hashlib.sha256(path.read_bytes()).hexdigest()}
(RESULTS_DIR/'q4_delivery_validation.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(audit,ensure_ascii=False))
