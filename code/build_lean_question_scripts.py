"""Write standalone answer-only scripts without executing or verifying them."""
import ast
from pathlib import Path
import zipfile

ROOT = Path(__file__).resolve().parents[1]
DEST = ROOT / 'paper' / 'guosai2026' / '支撑材料_按问复现'


def extract(filename, name):
    source = (ROOT / 'code' / filename).read_text(encoding='utf-8-sig')
    node = next(n for n in ast.parse(source).body if isinstance(n, ast.FunctionDef) and n.name == name)
    return ast.get_source_segment(source, node) + '\n\n'


COMMON = '''from pathlib import Path
import csv
import json
import numpy as np
from openpyxl import load_workbook
from scipy.integrate import solve_ivp
from scipy.interpolate import PchipInterpolator
from scipy.sparse import diags, lil_matrix

ROOT = Path(__file__).resolve().parent
ATTACH = ROOT / "data" / "附件"
OUT = ROOT / "results"
R, H, HM, T0, C0 = 0.02, 25.0, 8e-7, 28.0, 2.55


def read_table(name):
    wb = load_workbook(ATTACH / name, read_only=True, data_only=True)
    rows = np.asarray(list(wb.active.values)[1:], dtype=float)
    wb.close()
    return rows


def environment(end):
    rows = read_table("附件1.xlsx")
    rows = rows[rows[:, 0] <= end]
    return (PchipInterpolator(rows[:, 0], rows[:, 1], extrapolate=False),
            PchipInterpolator(rows[:, 0], rows[:, 2], extrapolate=False))


def save_excel(question, times, fields):
    wb = load_workbook(ATTACH / "附件3" / f"result{question}.xlsx")
    for ws, field in zip(wb.worksheets, fields):
        header, surface = ws.cell(1, 1).value, ws.cell(1, ws.max_column).value
        ws.delete_rows(1, ws.max_row)
        ws.append([header] + [j / 10 for j in range(21)] + ([surface] if question == 4 else []))
        for t, row in zip(times, field):
            ws.append([float(t)] + [None if np.isnan(v) else float(v) for v in row])
        for row in ws.iter_rows(min_row=2, min_col=2):
            for cell in row:
                cell.number_format = "0.0000"
        ws.freeze_panes = "B2"
    OUT.mkdir(exist_ok=True)
    wb.properties.creator = wb.properties.lastModifiedBy = "Team"
    wb.save(OUT / f"result{question}.xlsx")
    wb.close()
    print(f"已生成 results/result{question}.xlsx", flush=True)


def save_csv(name, times, values, headers):
    with (OUT / name).open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.writer(stream)
        writer.writerow(["时间"] + headers)
        for t, row in zip(times, values):
            writer.writerow([f"{t:.4f}"] + ["" if np.isnan(v) else f"{v:.4f}" for v in row])


def integrate(rhs, interval, initial, **options):
    solution = solve_ivp(rhs, interval, initial, method="BDF", **options)
    if not solution.success:
        raise RuntimeError(solution.message)
    return solution


def geometry(nodes):
    faces = (nodes[:-1] + nodes[1:]) / 2
    return faces, np.diff(nodes), np.diff(np.r_[nodes[0], faces, nodes[-1]]**2) / 2


'''

COUPLED = '''def coupled_sparsity(n):
    matrix = lil_matrix((2*n, 2*n))
    for i in range(n):
        for j in range(max(0, i-1), min(n, i+2)):
            matrix[i, j] = matrix[i, n+j] = matrix[n+i, j] = matrix[n+i, n+j] = 1
    return matrix.tocsc()


'''

PROPS3 = '''def properties(C, T):
    rho = 650.0 + 128.0 * C
    cp = 1450.0 + 2736.0 * C / (1.0 + C)
    k = 0.21 + 0.38 * C / (1.0 + C)
    D = np.zeros_like(C)
    good = C > 0
    D[good] = 2.4e-3 * np.exp(-0.45 / C[good]) * np.exp(-3850.0 / (T[good] + 273.15))
    return rho * cp, k, D


def coupled_rhs(nodes):
    n = len(nodes)
    faces, dr, weights = geometry(nodes)
    def rhs(y, Ta, Ca):
        T, C = y[:n], y[n:]
        B, k, D = properties(C, T)
        FT = -faces * harmonic_mean(k[:-1], k[1:]) * np.diff(T) / dr
        FC = -faces * harmonic_mean(D[:-1], D[1:]) * np.diff(C) / dr
        td, cd = np.empty(n), np.empty(n)
        td[0], cd[0] = -FT[0] / (B[0]*weights[0]), -FC[0] / weights[0]
        td[1:-1] = (FT[:-1]-FT[1:]) / (B[1:-1]*weights[1:-1])
        cd[1:-1] = (FC[:-1]-FC[1:]) / weights[1:-1]
        td[-1] = (FT[-1]-R*H*(T[-1]-Ta)) / (B[-1]*weights[-1])
        cd[-1] = (FC[-1]-R*HM*(C[-1]-Ca)) / weights[-1]
        return np.r_[td, cd]
    return rhs


'''

Q1 = '''def main():
    Ta, Ca = environment(1800.0)
    r = graded_radial_nodes(8)
    n = len(r)
    faces, dr, weights = geometry(r)
    coefficient = 0.36 * (faces / dr)
    denominator = 820.0 * 2600.0 * weights
    diagonal = -np.r_[coefficient[0], coefficient[:-1]+coefficient[1:], coefficient[-1]+R*H] / denominator
    matrix = diags([coefficient/denominator[1:], diagonal, coefficient/denominator[:-1]], [-1, 0, 1], format="csc")
    dense_matrix = matrix.toarray()
    def heat(t, T):
        value = dense_matrix @ T
        value[-1] += R*H / denominator[-1] * float(Ta(t))
        return value
    def moisture(t, C):
        D = np.zeros_like(C)
        good = C > 0
        D[good] = 7e-9 * np.exp(-0.89/C[good])
        flux = -faces * harmonic_mean(D[:-1], D[1:]) * np.diff(C) / dr
        value = np.empty(n)
        value[0] = -flux[0] / weights[0]
        value[1:-1] = (flux[:-1]-flux[1:]) / weights[1:-1]
        value[-1] = (flux[-1]-R*HM*(C[-1]-float(Ca(t)))) / weights[-1]
        return value
    times = np.arange(1.0, 1801.0)
    common = dict(rtol=1e-10, max_step=0.5, t_eval=times)
    print("问题一：计算温度与含水率...", flush=True)
    temperature = integrate(heat, (0, 1800), np.full(n, T0), atol=1e-10, jac=matrix, **common)
    sparsity = diags([np.ones(n-1), np.ones(n), np.ones(n-1)], [-1, 0, 1], format="csc")
    water = integrate(moisture, (0, 1800), np.full(n, C0), atol=1e-12, jac_sparsity=sparsity, **common)
    ids = [int(np.argmin(abs(r-j*.001))) for j in range(21)]
    T, C = temperature.y[ids].T, water.y[ids].T
    save_excel(1, times, [T, C])
    tt = np.array([100, 300, 600, 900, 1200, 1500, 1800])
    for label, values in (("temperature", T), ("moisture", C)):
        save_csv(f"q1_table_{label}.csv", tt, values[tt-1, ::5], ["0 cm", "0.5 cm", "1 cm", "1.5 cm", "2 cm"])
'''

Q2 = '''def main():
    Ta, Ca = environment(10800.0)
    r = graded_radial_nodes(4)
    n = len(r)
    operator = coupled_rhs(r)
    rhs = lambda t, y: operator(y, float(Ta(t)), float(Ca(t)))
    times = np.arange(1.0, 10801.0)
    print("问题二：计算变物性热湿耦合...", flush=True)
    solution = integrate(rhs, (0, 10800), np.r_[np.full(n, T0), np.full(n, C0)],
                         rtol=1e-9, atol=np.r_[np.full(n, 1e-8), np.full(n, 1e-10)],
                         max_step=1.0, t_eval=times, jac_sparsity=coupled_sparsity(n))
    ids = np.array([int(np.argmin(abs(r-j*.001))) for j in range(21)])
    T, C = solution.y[ids].T, solution.y[n+ids].T
    save_excel(2, times, [T, C])
    tt = np.arange(1800, 10801, 1800)
    for label, values in (("temperature", T), ("moisture", C)):
        save_csv(f"q2_table_{label}.csv", tt/3600, values[tt-1, ::5], ["0 cm", "0.5 cm", "1 cm", "1.5 cm", "2 cm"])
'''

Q3 = '''def main():
    from scipy.optimize import brentq
    Ta, Ca = environment(14400.0)
    r = graded_radial_nodes(8)
    n = len(r)
    operator = coupled_rhs(r)
    rhs = lambda t, y: operator(y, float(Ta(t)) if t <= 14400 else 50.0,
                               float(Ca(t)) if t <= 14400 else 0.05)
    options = dict(rtol=1e-9, atol=np.r_[np.full(n, 1e-8), np.full(n, 1e-10)], jac_sparsity=coupled_sparsity(n))
    event = lambda t, y: float(np.max(y[n:]) - 0.15)
    event.direction, event.terminal = -1.0, True
    print("问题三：从初态独立计算全域达标时间...", flush=True)
    early = integrate(rhs, (0, 14400), np.r_[np.full(n, T0), np.full(n, C0)],
                      t_eval=np.arange(60, 14401, 60), first_step=0.05, max_step=2.0, **options)
    late = integrate(rhs, (14400, 259200), early.y[:, -1], t_eval=np.arange(14460, 259201, 60),
                     first_step=0.2, max_step=60.0, events=event, **options)
    if not len(late.t_events[0]):
        raise RuntimeError("No drying event within 72 hours")
    t_left, y_left = float(late.t[-1]), late.y[:, -1]
    t_raw = float(late.t_events[0][0])
    zoom = integrate(rhs, (t_left, t_raw+max(30, 2*(t_raw-t_left))), y_left,
                     first_step=0.05, max_step=0.5, dense_output=True, events=event, **options)
    if not len(zoom.t_events[0]):
        raise RuntimeError("Local drying-event solve failed")
    g = lambda t: float(np.max(zoom.sol(t)[n:]) - 0.15)
    t_star = float(zoom.t_events[0][0])
    if g(zoom.t[-1]) <= 0:
        t_star = brentq(g, t_left, zoom.t[-1], xtol=1e-4, rtol=1e-12)
    after = integrate(rhs, (t_star, t_star+3600), zoom.sol(t_star),
                      first_step=0.05, max_step=5.0, dense_output=True, **options)
    # Empirical margin from the original full G8 verification; not re-estimated here.
    margin = 3.508819596595336e-7
    target = lambda t: float(np.max(after.sol(t)[n:]) - (0.15-margin))
    crossing = brentq(target, t_star, t_star+3600, xtol=1e-4, rtol=1e-12)
    report = np.ceil((np.floor(crossing)+1)/3600*10000)/10000*3600
    final = after.sol(report)
    if np.max(final[n:]) >= 0.15:
        raise RuntimeError("Final state is not strictly below the drying threshold")
    times = np.r_[early.t, late.t, report]
    states = np.vstack([early.y.T, late.y.T, final])
    ids = np.array([int(np.argmin(abs(r-j*.001))) for j in range(21)])
    C = states[:, n+ids]
    save_excel(3, times, [C])
    tt = np.r_[np.arange(21600, report, 21600), report]
    rows = [int(np.argmin(abs(times-t))) for t in tt]
    save_csv("q3_table_moisture.csv", tt/3600, C[rows, ::5], ["0 cm", "0.5 cm", "1 cm", "1.5 cm", "2 cm"])
    answer = dict(critical_h=t_star/3600, report_h=report/3600, max_moisture=float(np.max(final[n:])), retained_margin=margin)
    (OUT / "q3_answer.json").write_text(json.dumps(answer, indent=2), encoding="utf-8")
    print(f"烘干报告时间：{report/3600:.4f} h", flush=True)
'''

Q4 = '''def main():
    Ta, Ca = environment(14400.0)
    rows = read_table("附件2.xlsx")
    radius = PchipInterpolator(rows[:, 0], rows[:, 1]/100, extrapolate=False)
    xi = graded_radial_nodes(8)/R
    n = len(xi)
    faces, dx, weights = geometry(xi)
    def rhs(t, y, late):
        rr = float(radius(t))
        T, C = y[:n], y[n:]
        B, k, D = properties(C, T)
        ambient_T, ambient_C = (50.0, 0.05) if late else (float(Ta(t)), float(Ca(t)))
        FT = np.r_[0., -faces*harmonic_mean(k[:-1], k[1:])*np.diff(T)/dx, rr*H*(T[-1]-ambient_T)]
        FC = np.r_[0., -faces*harmonic_mean(D[:-1], D[1:])*np.diff(C)/dx, rr*HM*(C[-1]-ambient_C)]
        return np.r_[-np.diff(FT)/(rr**2*weights*B), -np.diff(FC)/(rr**2*weights)]
    event = lambda t, y: float(np.max(y[n:])-0.15)
    event.direction, event.terminal = -1.0, False
    options = dict(rtol=1e-8, atol=np.r_[np.full(n, 1e-8), np.full(n, 1e-10)], jac_sparsity=coupled_sparsity(n))
    initial = np.r_[np.full(n, T0), np.full(n, C0)]
    times, states = [], []
    breaks = np.unique(np.r_[0., 14400., np.arange(21600., min(259200., rows[-1, 0])+1, 21600.)])
    print("问题四：独立计算径向收缩模型...", flush=True)
    t_star = None
    for start, end in zip(breaks[:-1], breaks[1:]):
        late = start >= 14400
        solution = integrate(lambda t, y: rhs(t, y, late), (start, end), initial,
                             max_step=60.0 if late else 10.0, dense_output=True, events=event, **options)
        ts = np.arange(start+60, end+0.1, 60)
        times.extend(ts)
        states.extend(solution.sol(ts).T)
        initial = solution.y[:, -1]
        print(f"已计算至 {end/3600:.1f} h", flush=True)
        if len(solution.t_events[0]):
            t_star = float(solution.t_events[0][0])
            break
    if t_star is None:
        raise RuntimeError("No drying event within the measured radius interval")
    # Keep the two-second offset; this lean version does not compare other grids.
    report = np.ceil((t_star+2.0)/0.36)*0.36
    if report > rows[-1, 0]:
        raise RuntimeError("Reporting time exceeds measured radius coverage")
    times, states = np.asarray(times), np.asarray(states)
    keep = times < report
    times, states = times[keep], states[keep]
    final_solution = integrate(lambda t, y: rhs(t, y, True), (times[-1], report), states[-1],
                               rtol=1e-10, atol=np.r_[np.full(n, 1e-10), np.full(n, 1e-12)],
                               max_step=1.0, jac_sparsity=coupled_sparsity(n))
    final = final_solution.y[:, -1]
    if np.max(final[n:]) >= 0.15:
        raise RuntimeError("Final state is not strictly below the drying threshold")
    times = np.r_[times, report]
    states = np.vstack([states, final])
    fixed = np.arange(21)*.001
    radii = radius(times)
    C = np.full((len(times), 22), np.nan)
    for i, (state, rr) in enumerate(zip(states, radii)):
        good = fixed <= rr+1e-14
        C[i, :21][good] = PchipInterpolator(xi, state[n:], extrapolate=False)(np.minimum(fixed[good]/rr, 1.0))
        C[i, -1] = state[-1]
    save_excel(4, times, [C])
    tt = np.r_[np.arange(21600, report, 21600), report]
    ids = [int(np.argmin(abs(times-t))) for t in tt]
    table = np.column_stack([C[ids][:, [0, 5, 10, 15, 20, 21]], radii[ids]*100])
    save_csv("q4_table6.csv", tt/3600, table, ["0 cm", "0.5 cm", "1 cm", "1.5 cm", "2 cm", "药材表面", "半径/cm"])
    answer = dict(critical_h=t_star/3600, report_h=report/3600, max_moisture=float(np.max(final[n:])), radius_cm=float(radius(report)*100))
    (OUT / "q4_answer.json").write_text(json.dumps(answer, indent=2), encoding="utf-8")
    print(f"烘干报告时间：{report/3600:.4f} h", flush=True)
'''

README = '''# 四问独立求解代码（精简版）

Q1.py至Q4.py分别从原始附件计算对应问题，输出results/result1.xlsx至result4.xlsx及题定表CSV。第三、四问另输出包含临界时间和报告时间的答案JSON。

每问都是普通、可直接阅读的独立Python文件，不需要原项目代码，不调用其他问。只保留本问必要的输入、方程、基准网格求解和数据导出；已删除模块内嵌机制、绘图、多网格复算、守恒审计、解析对照及额外实验。

使用Python 3.12或更高版本，保留data附件结构。安装依赖后选择需要的问题运行，无须按顺序运行：

```text
python -m pip install -r requirements.txt
python Q1.py
python Q2.py
python Q3.py
python Q4.py
```

输入和输出均相对于脚本所在位置，不含个人绝对路径。reference_results中的Excel仅供人工参考，不参与求解。

第一问使用G8网格，第二问使用G4网格，第三、四问使用G8网格；物性、初边值和基准积分容差沿用原模型。Excel保留未舍入数据，以四位小数显示，第四问域外位置留空。

第三问沿用原完整复核确定的含水率裕量3.508819596595336e-7，本脚本不重新估计裕量。第四问按基准事件时间延后2秒并向上取四位小时，不再重新计算其他网格的事件时间。这些报告规则针对随附题目数据，修改参数或附件后不能沿用旧版验证结论。

按用户要求，精简后未运行或核验，不承接此前完整版的逐格一致性结论。运行中仅保留求解失败和未达阈值等必要错误处理。原正式四个Excel未修改。
'''


def main():
    mesh = extract('radial_fvm.py', 'graded_radial_nodes')
    harmonic = extract('radial_fvm.py', 'harmonic_mean')
    for q, body in ((1, Q1), (2, Q2), (3, Q3), (4, Q4)):
        header = f'# -*- coding: utf-8 -*-\n"""问题{q}：独立求解并导出答案。运行 python Q{q}.py。\n仅依赖numpy、scipy、openpyxl和同目录data附件。精简版未重新运行核验。\n"""\n'
        props = PROPS3 if q in (2, 3) else extract('problem4.py', 'properties') if q == 4 else ''
        text = header + COMMON + mesh + harmonic + (COUPLED if q != 1 else '') + props + body
        text += '\n\nif __name__ == "__main__":\n    main()\n'
        (DEST / f'Q{q}.py').write_text(text, encoding='utf-8')
    (DEST / 'README.md').write_text(README, encoding='utf-8')
    (DEST / 'requirements.txt').write_text('numpy==2.5.1\nscipy==1.18.0\nopenpyxl==3.1.5\n', encoding='utf-8')
    included = [DEST / f'Q{q}.py' for q in range(1, 5)] + [DEST / 'README.md', DEST / 'requirements.txt']
    for directory in ('data', 'reference_results'):
        included += sorted(p for p in (DEST / directory).rglob('*') if p.is_file())
    with zipfile.ZipFile(DEST.with_suffix('.zip'), 'w', zipfile.ZIP_DEFLATED) as z:
        for path in included:
            z.write(path, DEST.name + '/' + path.relative_to(DEST).as_posix())
    print('Q1-Q4 and ZIP updated. No solver or verification was executed.')


if __name__ == '__main__':
    main()
