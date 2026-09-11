"""Read-only audit of the supplied A-problem data; no PDE solution."""
import json
import math
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
data = ROOT / 'data' / '附件'
report = {}
for name in ['附件1.xlsx', '附件2.xlsx']:
    wb = load_workbook(data / name, read_only=True, data_only=True)
    rows = list(wb.active.values)
    a = np.asarray(rows[1:], dtype=float)
    info = dict(headers=rows[0], rows=len(a), columns=a.shape[1],
                missing_or_nonfinite=int((~np.isfinite(a)).sum()),
                duplicate_times=int(len(a)-len(np.unique(a[:, 0]))),
                time_strictly_increasing=bool(np.all(np.diff(a[:, 0]) > 0)),
                time_step_values=np.unique(np.diff(a[:, 0])).tolist(),
                minimum=np.min(a, axis=0).tolist(), maximum=np.max(a, axis=0).tolist(),
                first=a[0].tolist(), last=a[-1].tolist())
    if a.shape[1] == 3:
        late = a[a[:, 0] >= 10800]
        info['last_hour_mean_T_C'] = late[:, 1:].mean(axis=0).tolist()
        info['last_hour_std_T_C'] = late[:, 1:].std(axis=0, ddof=1).tolist()
        info['at_1800s'] = a[a[:, 0] == 1800][0].tolist()
    else:
        info['radius_increases'] = int((np.diff(a[:, 1]) > 0).sum())
        info['nonpositive_radii'] = int((a[:, 1] <= 0).sum())
    report[name] = info
    wb.close()
for path in sorted((data / '附件3').glob('*.xlsx')):
    wb = load_workbook(path, read_only=True, data_only=True)
    report[path.name] = {s.title: list(s.values) for s in wb.worksheets}
    wb.close()
R, L, h, k, rho, cp, hm, C = .02, .25, 25, .36, 820, 2600, 8e-7, 2.55
D = 7e-9 * math.exp(-.89 / C)
report['scale_checks_not_solution'] = dict(
    length_over_diameter=L/(2*R), end_area_over_lateral_area=R/L,
    alpha=k/(rho*cp), Bi_radius=h*R/k,
    Bi_volume_over_area=h*(R*L/(2*(L+R)))/k,
    initial_D1=D, initial_mass_Bi=hm*R/D,
    thermal_radial_scale_seconds=R*R/(k/(rho*cp)),
    moisture_radial_scale_hours=R*R/D/3600,
    D23_at_C_015_T50=2.4e-3*math.exp(-.45/.15)*math.exp(-3850/323.15),
    D23_at_C_255_T50=2.4e-3*math.exp(-.45/2.55)*math.exp(-3850/323.15))
out = ROOT/'reports'/'A_INPUT_AUDIT.json'
out.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(report, ensure_ascii=False, indent=2))
