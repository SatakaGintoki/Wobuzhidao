"""Read-only Q2 source audit and coefficient/geometry scale checks; no PDE solve."""
from pathlib import Path
import hashlib
import json
import math
import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


def audit():
    paths = [ROOT / 'data/A.pdf', ROOT / 'data/附件/附件1.xlsx',
             ROOT / 'data/附件/附件3/result2.xlsx']
    source = {str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest() for p in paths}
    wb = load_workbook(paths[1], read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.values)
    x = np.array(rows[1:], dtype=float)
    selected = x[x[:, 0] <= 10800]
    oven = dict(sheet=ws.title, headers=list(rows[0]), count=len(x),
                time_range_s=x[[0, -1], 0].tolist(), intervals_s=np.unique(np.diff(x[:, 0])).tolist(),
                finite=bool(np.isfinite(x).all()), unique_times=bool(len(np.unique(x[:, 0])) == len(x)),
                q2_count=len(selected), q2_first=selected[0].tolist(), q2_last=selected[-1].tolist(),
                q2_min=selected.min(axis=0).tolist(), q2_max=selected.max(axis=0).tolist())
    wb.close()
    wb = load_workbook(paths[2], read_only=False, data_only=False)
    template = {ws.title: dict(rows=ws.max_row, cols=ws.max_column,
                 values=[list(row) for row in ws.values], merges=[str(r) for r in ws.merged_cells.ranges])
                for ws in wb}
    wb.close()
    def props(c, t):
        rho = 650 + 128*c
        cp = 1450 + 2736*c/(1+c)
        k = 0.21 + 0.38*c/(1+c)
        d = 2.4e-3*math.exp(-0.45/c)*math.exp(-3850/(t+273.15))
        return dict(C=c, T_C=t, T_K=t+273.15, rho=rho, cp=cp, k=k,
                    capacity=rho*cp, alpha=k/(rho*cp), D=d,
                    dlogD_dT=3850/(t+273.15)**2, dlogD_dC=0.45/c**2)
    initial = props(2.55, 28.)
    cmin = float(selected[:, 2].min())
    tmax = float(selected[:, 1].max())
    # alpha=(0.21+0.59C)/[(650+128C)(1450+4186C)], decreasing for C>=0.
    amax = props(cmin, 28.)['alpha']
    dmax = props(2.55, tmax)['D']
    scales = dict(alpha_max_on_model_bounds=amax, D_max_on_model_bounds=dmax,
                  heat_length_3h_m=math.sqrt(amax*10800), moisture_length_3h_m=math.sqrt(dmax*10800),
                  end_distance_m=0.125,
                  one_end_step_erfc_scale=math.erfc(0.125/(2*math.sqrt(amax*10800))),
                  one_end_moisture_erfc_scale=math.erfc(0.125/(2*math.sqrt(dmax*10800))),
                  thermal_Bi_R_initial=25*0.02/initial['k'],
                  moisture_Bi_R_initial=8e-7*0.02/initial['D'],
                  initial_dry_density_if_wet_density=initial['rho']/3.55,
                  dry_density_inferred_at_C1=(650+128)/2,
                  note='Scales only; erfc comparison is not a 2D error bound or a Q2 solution.')
    q1_paths = ['code/problem1.py', 'code/radial_fvm.py', 'code/utils.py', 'code/q1_analytic.py',
                'results/q1_solution.npz', 'results/result1.xlsx', 'paper/drafts/problem1_model_consolidated.md']
    keep = {p: hashlib.sha256((ROOT/p).read_bytes()).hexdigest() for p in q1_paths}
    result = dict(status='INPUT_AND_SCALE_AUDIT_ONLY_NO_PDE', sources_sha256=source,
                  oven=oven, template=template, initial_properties=initial,
                  coefficient_samples=[props(2.55, 50.), props(1., 50.), props(0.15, 50.)],
                  scales=scales, q1_preservation_sha256=keep)
    path = ROOT / 'reports/Q2_INPUT_AUDIT.json'
    path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    audit()
