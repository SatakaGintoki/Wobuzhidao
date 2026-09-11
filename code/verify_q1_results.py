"""Read-only acceptance checks. Writes audit JSON/previews, never result1."""
from pathlib import Path
import hashlib
import json
import re
import sys
import numpy as np

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'results/verification/q1'
OUT.mkdir(parents=True, exist_ok=True)


def dump(name, obj):
    (OUT/name).write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(obj, ensure_ascii=True, indent=2), flush=True)


def files():
    from openpyxl import load_workbook
    import pypdfium2 as pdfium
    from pypdf import PdfReader
    from PIL import Image, ImageOps, ImageDraw
    paths = [ROOT/'data/A.pdf', ROOT/'data/附件/附件1.xlsx',
             ROOT/'data/附件/附件3/result1.xlsx']
    paths += list((ROOT/'results').glob('q1*'))+[ROOT/'results/result1.xlsx']
    paths += list((ROOT/'figures').glob('q1*.pdf'))
    paths += [ROOT/'code'/n for n in ('problem1.py','radial_fvm.py','utils.py','q1_analytic.py')]
    before = {str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in paths if p.is_file()}
    base = np.load(ROOT/'results/q1_solution.npz')
    times, radii = base['times'], base['output_radii_m']
    report = {'scope':'q1 result acceptance, read only', 'source_sha256':before,
              'npz_N':int(base['N']), 'npz_keys':list(base.files), 'workbook':{}, 'csv':{}, 'figures':[]}
    report['axes'] = {'times_1_to_1800':bool(np.array_equal(times,np.arange(1,1801))),
                      'radii_0_to_2cm':bool(np.allclose(radii,np.arange(21)*.001,rtol=0,atol=1e-15))}
    template = load_workbook(ROOT/'data/附件/附件3/result1.xlsx',read_only=True,data_only=False)
    report['template']={s.title:{'dimensions':[s.max_row,s.max_column],
                               'first_rows':list(s.values)[:4]} for s in template}
    template.close()
    wb = load_workbook(ROOT/'results/result1.xlsx',data_only=False)
    report['sheet_order_ok'] = wb.sheetnames == ['温度','水分浓度']
    selected_times=np.array([100,300,600,900,1200,1500,1800])
    validation=json.loads((ROOT/'results/q1_validation.json').read_text(encoding='utf-8'))
    source_report=(ROOT/'reports/RESULTS_REPORT.md').read_text(encoding='utf-8')
    table_lines=[line for line in source_report.splitlines() if re.match(r'^\| (100|300|600|900|1200|1500|1800) \|',line)]
    for index,(sheet,key,paperkey) in enumerate([('温度','T','paper_table_T'),('水分浓度','C','paper_table_C')]):
        ws=wb[sheet]
        arr=np.array([[c.value for c in row] for row in ws.iter_rows(min_row=2,max_row=1801,min_col=2,max_col=22)],dtype=float)
        formats=[c.number_format for row in ws.iter_rows(min_row=2,max_row=1801,min_col=2,max_col=22) for c in row]
        celltypes=[c.data_type for row in ws.iter_rows(min_row=2,max_row=1801,min_col=2,max_col=22) for c in row]
        diff=np.abs(arr-base[key+'_out'])
        step=int(base['N'])//20
        report['workbook'][sheet]={
            'dimensions':[ws.max_row,ws.max_column],
            'times_ok':bool(np.array_equal([ws.cell(i,1).value for i in range(2,1802)],times)),
            'radii_ok':bool(np.allclose([ws.cell(1,j).value for j in range(2,23)],radii*100,atol=1e-14)),
            'all_finite':bool(np.isfinite(arr).all()),'numeric_cells':celltypes.count('n'),
            'four_decimal_formatted_cells':formats.count('0.0000'),
            'max_difference_vs_npz':float(diff.max()),
            'npz_output_vs_full':float(np.max(np.abs(base[key+'_out']-base[key+'_full'][:,::step]))),
            'minimum':float(arr.min()),'maximum':float(arr.max()),
            'final_center':float(arr[-1,0]),'final_surface':float(arr[-1,-1]),
            'hidden_rows':sum(bool(x.hidden) for x in ws.row_dimensions.values()),
            'hidden_cols':sum(bool(x.hidden) for x in ws.column_dimensions.values()),
            'a1':ws['A1'].value, 'freeze_panes':ws.freeze_panes}
        csvpath=ROOT/'results'/('q1_table_temperature.csv' if key=='T' else 'q1_table_moisture.csv')
        csv=np.loadtxt(csvpath,delimiter=',',skiprows=1)
        expected=base[key+'_out'][selected_times-1][:,::5]
        printed=np.array([[float(x.strip()) for x in line.split('|')[2:-1]] for line in table_lines[index*7:(index+1)*7]])
        jsonvalues=[validation[paperkey][f't{t}_r{r:g}'] for t in selected_times for r in np.arange(5)*.5]
        report['csv'][key]={'shape':list(csv.shape),
            'times_match':bool(np.array_equal(csv[:,0],selected_times)),
            'max_difference':float(np.max(np.abs(csv[:,1:]-expected))),
            'paper_report_four_decimal_matches':bool(np.array_equal(printed,np.round(expected,4))),
            'validation_json_table_max_difference':float(np.max(np.abs(np.array(jsonvalues).reshape(7,5)-expected)))}
    wb.close()
    # Recompute the existing global balance using the saved one-second series.
    env=load_workbook(ROOT/'data/附件/附件1.xlsx',read_only=True,data_only=True)
    e=np.array(list(env.active.values)[1:],dtype=float);env.close()
    r=base['r'];dr=.02/int(base['N']);a=np.maximum(0,r-dr/2);b=np.minimum(.02,r+dr/2);w=(b*b-a*a)/2
    t=np.r_[0,times];Ta=np.interp(t,e[:,0],e[:,1]);Ca=np.interp(t,e[:,0],e[:,2])
    Ts=np.r_[28,base['T_full'][:,-1]];Cs=np.r_[2.55,base['C_full'][:,-1]]
    E=2*np.pi*.25*820*2600*np.dot(w,base['T_full'][-1]-28)
    Q=2*np.pi*.02*.25*25*np.trapezoid(Ta-Ts,t)
    loss=np.dot(w,2.55-base['C_full'][-1]);flux=.02*8e-7*np.trapezoid(Cs-Ca,t)
    report['conservation_recomputed']={'heat_E_J':float(E),'heat_Q_J':float(Q),
        'heat_relative_residual':float(abs(E-Q)/max(abs(E),abs(Q))),
        'moisture_relative_residual':float(abs(loss-flux)/max(abs(loss),abs(flux)))}
    previews=[]
    for p in sorted((ROOT/'figures').glob('q1*.pdf')):
        reader=PdfReader(p)
        texts=[pg.extract_text() or '' for pg in reader.pages]
        doc=pdfium.PdfDocument(str(p))
        for j in range(len(doc)):
            page=doc[j];bitmap=page.render(scale=1.7);im=bitmap.to_pil().convert('RGB')
            dest=OUT/f'{p.stem}_page{j+1}.png';im.save(dest);previews.append((p.stem,im.copy()))
            bitmap.close();page.close()
        doc.close()
        report['figures'].append({'name':p.name,'pages':len(reader.pages),
                                   'text':'\n'.join(texts),'nonempty':all(texts)})
    contact=Image.new('RGB',(1320,3*460),'white');draw=ImageDraw.Draw(contact)
    for i,(name,im) in enumerate(previews):
        im.thumbnail((650,420));x=(i%2)*660;y=(i//2)*460
        contact.paste(im,(x,y+25));draw.text((x+10,y+5),name,fill='black')
    contact.save(OUT/'figure_contact.png')
    report['source_files_unchanged']=all(hashlib.sha256((ROOT/p).read_bytes()).hexdigest()==h for p,h in before.items())
    dump('file_audit.json',report)


def numerical():
    from problem1 import load_q1_inputs, solve_moisture
    from radial_fvm import RadialFVM
    from q1_analytic import BesselDuhamel
    from utils import ALPHA, H, R, K, T0
    base=np.load(ROOT/'results/q1_solution.npz')
    Ta,Ca,breaks=load_q1_inputs()
    ref=BesselDuhamel(H*R/K,ALPHA,R,n_terms=120)
    tr=ref.evaluate_many(base['output_radii_m'],base['times'],Ta,T0)
    report={'temperature_max_difference_to_120_mode_reference':float(np.max(np.abs(tr-base['T_out']))),'moisture_short_runs':[]}
    solved={}
    for N in (640,1280,2560):
        fvm=RadialFVM(N)
        t,y,_=solve_moisture(fvm,Ca,np.array([0.,60.,100.]),1e-10,1e-12,.1)
        vals=y[t>0][:,::N//20]
        solved[N]=vals
        entry={'N':N,'C_surface_t1':float(vals[0,-1]),'C_surface_t100':float(vals[-1,-1]),
               'max_difference_vs_saved_N640_first100s':float(np.max(np.abs(vals-base['C_out'][:100]))),
               'four_decimal_changed_cells_vs_saved':int(np.sum(np.round(vals,4)!=np.round(base['C_out'][:100],4)))}
        report['moisture_short_runs'].append(entry)
        print(json.dumps(entry),flush=True)
    for lo,hi in ((640,1280),(1280,2560)):
        diff=np.abs(solved[lo]-solved[hi]);i,j=np.unravel_index(np.argmax(diff),diff.shape)
        report[f'C_{lo}_vs_{hi}_first100s']={'max_abs':float(diff.max()),'time_s':int(i+1),'radius_cm':float(j*.1)}
    np.savez_compressed(OUT/'short_moisture_checks.npz',times=np.arange(1,101),**{f'C_N{n}':y for n,y in solved.items()})
    dump('numerical_audit.json',report)


if __name__=='__main__':
    numerical() if '--numerical' in sys.argv else files()
