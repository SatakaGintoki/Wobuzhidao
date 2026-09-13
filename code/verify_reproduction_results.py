"""Compare every answer cell with the supplied original result workbooks."""
import argparse
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


def compare(root, questions):
    records = []
    for q in questions:
        actual = load_workbook(root / 'results' / f'result{q}.xlsx', read_only=True, data_only=True)
        expected = load_workbook(root / 'reference_results' / f'result{q}.xlsx', read_only=True, data_only=True)
        if actual.sheetnames != expected.sheetnames:
            raise AssertionError(f'Q{q}: sheet names differ')
        for a, b in zip(actual, expected):
            if (a.max_row, a.max_column) != (b.max_row, b.max_column):
                raise AssertionError(f'Q{q}: sheet dimensions differ')
            maximum = 0.0
            changed = count = 0
            for row_id, (ra, rb) in enumerate(zip(a.values, b.values), 1):
                for col, (va, vb) in enumerate(zip(ra, rb), 1):
                    if isinstance(va, (float, int)) and isinstance(vb, (float, int)):
                        if not np.isfinite([va, vb]).all():
                            raise AssertionError(f'Q{q}: nonfinite cell')
                        error = abs(va - vb)
                        if row_id == 1 or col == 1:
                            if error > 1e-8:
                                raise AssertionError(f'Q{q}: time or radius axis differs at {row_id},{col}')
                        else:
                            count += 1
                            maximum = max(maximum, error)
                            changed += f'{va:.4f}' != f'{vb:.4f}'
                    elif va != vb:
                        raise AssertionError(f'Q{q}: text/domain blank differs at {row_id},{col}')
            records.append({'question': q, 'sheet': a.title, 'shape': [a.max_row, a.max_column],
                            'numeric_answer_cells': count, 'maximum_absolute_difference': maximum,
                            'changed_at_four_decimals': changed, 'pass': changed == 0 and maximum < 2e-5})
        actual.close()
        expected.close()
    summary = {'scope': 'All time/radius axes, all answer cells, sheet names and outside-domain blanks',
               'checks': records, 'pass': all(r['pass'] for r in records)}
    (root / 'reproduction_check.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    if not summary['pass']:
        raise SystemExit('Recomputed values differ from the reference; see reproduction_check.json')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--questions', nargs='+', type=int, choices=range(1, 5), default=[1, 2, 3, 4])
    args = parser.parse_args()
    compare(Path(__file__).resolve().parent, args.questions)
