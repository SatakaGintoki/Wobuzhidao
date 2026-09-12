"""Problem 1: radial FVM + BDF for preheating temperature and moisture."""
from __future__ import annotations

import json
import sys
import time
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from scipy.integrate import solve_ivp, quad

CODE_DIR = Path(__file__).resolve().parent
if str(CODE_DIR) not in sys.path:
    sys.path.insert(0, str(CODE_DIR))

from q1_analytic import BesselDuhamel, converged_reference
from delivery import make_gate, require_delivery
from radial_fvm import RadialFVM
from radial_fvm import graded_radial_nodes
from utils import (
    ALPHA,
    C0,
    CP,
    FIGURES_DIR,
    H,
    HM,
    K,
    L,
    R,
    RESULTS_DIR,
    RHO,
    T0,
    T_END,
    PchipOven,
    load_oven_table,
    moisture_diffusivity,
)

OUTPUT_TIMES = np.arange(1.0, T_END + 0.5, 1.0)
TABLE_TIMES = np.array([100, 300, 600, 900, 1200, 1500, 1800], dtype=float)
TABLE_RADII_CM = np.array([0.0, 0.5, 1.0, 1.5, 2.0])
OUTPUT_RADII_M = 0.001 * np.arange(21)  # 0, 0.1, ..., 2.0 cm
PAPER_TIMES = TABLE_TIMES.copy()


def load_q1_inputs() -> tuple[PchipOven, PchipOven, np.ndarray]:
    t, Ta, Ca = load_oven_table()
    mask = t <= T_END + 1e-12
    t, Ta, Ca = t[mask], Ta[mask], Ca[mask]
    if t.size != 31 or abs(t[-1] - T_END) > 1e-12:
        raise ValueError(f"expected 31 oven records on [0, 1800], got {t.size}, t_end={t[-1]}")
    if abs(Ta[0] - T0) > 1e-12:
        raise ValueError(f"Ta(0)={Ta[0]} != T0={T0}")
    Ta_fun = PchipOven(t, Ta)
    Ca_fun = PchipOven(t, Ca)
    for tj, yj in zip(t, Ta):
        if abs(float(Ta_fun(tj)) - float(yj)) > 1e-12:
            raise RuntimeError(f"PCHIP misses Ta knot t={tj}")
    for tj, yj in zip(t, Ca):
        if abs(float(Ca_fun(tj)) - float(yj)) > 1e-12:
            raise RuntimeError(f"PCHIP misses Ca knot t={tj}")
    return Ta_fun, Ca_fun, t


def integrate_mol(
    rhs,
    y0: np.ndarray,
    t_breaks: np.ndarray,
    *,
    rtol: float,
    atol: float,
    max_step: float,
    jac=None,
    jac_sparsity=None,
    field_name: str = "field",
):
    y = np.asarray(y0, dtype=float).copy()
    times = [0.0]
    states = [y.copy()]
    pieces: list[tuple[float, float, object]] = []
    t_end = float(t_breaks[-1])
    for i in range(len(t_breaks) - 1):
        t0 = float(t_breaks[i])
        t1 = float(min(t_breaks[i + 1], t_end))
        if t1 <= t0:
            continue
        t_eval = np.arange(np.floor(t0) + 1.0, np.floor(t1) + 1.0 + 1e-12, 1.0)
        t_eval = t_eval[(t_eval > t0 + 1e-14) & (t_eval <= t1 + 1e-14)]
        kwargs = dict(
            fun=rhs,
            t_span=(t0, t1),
            y0=y,
            method="BDF",
            rtol=rtol,
            atol=atol,
            max_step=max_step,
            dense_output=True,
            vectorized=False,
        )
        if t_eval.size:
            kwargs["t_eval"] = t_eval
        if jac is not None:
            kwargs["jac"] = jac
        if jac_sparsity is not None:
            kwargs["jac_sparsity"] = jac_sparsity
        sol = solve_ivp(**kwargs)
        if not sol.success:
            raise RuntimeError(f"{field_name} BDF failed on [{t0}, {t1}]: {sol.message}")
        pieces.append((t0, t1, sol))
        if sol.t.size:
            for k, tk in enumerate(sol.t):
                if times and abs(tk - times[-1]) < 1e-12:
                    states[-1] = sol.y[:, k].copy()
                    continue
                times.append(float(tk))
                states.append(sol.y[:, k].copy())
        y = sol.y[:, -1].copy() if sol.y.size else y
    t_arr = np.asarray(times, dtype=float)
    Y = np.stack(states, axis=0)
    return t_arr, Y, pieces


def dense_state(pieces, t: float) -> np.ndarray:
    t = float(t)
    for t0, t1, sol in pieces:
        if t0 - 1e-12 <= t <= t1 + 1e-12:
            return np.asarray(sol.sol(np.clip(t, t0, t1)), dtype=float)
    raise ValueError(f"time {t} not covered by dense output")


def extract_output_grid(fvm: RadialFVM, Y: np.ndarray) -> np.ndarray:
    idx = fvm.output_indices()
    return Y[:, idx]


def conservation_heat(fvm: RadialFVM, pieces, Ta_fun: PchipOven, T_init: np.ndarray, t: float) -> dict:
    T_t = dense_state(pieces, t)
    E = 2.0 * np.pi * L * RHO * CP * np.dot(fvm.W, T_t - T_init)
    flux = lambda x: float(Ta_fun(x)) - float(dense_state(pieces, x)[-1])
    Qraw, qerr = quad(flux, 0, t, points=Ta_fun.t[(Ta_fun.t > 0) & (Ta_fun.t < t)], epsabs=1e-9, epsrel=1e-10, limit=500)
    Q = 2.0 * np.pi * R * L * H * Qraw
    denom = max(abs(Q), abs(E), 1e-12)
    return {
        "E_relative_J": E,
        "Q_surface_J": Q,
        "relative_residual": abs(E - Q) / denom,
        "quadrature_abs_error": qerr, "quadrature": "adaptive QUADPACK on PCHIP intervals",
    }


def conservation_moisture(fvm: RadialFVM, pieces, Ca_fun: PchipOven, C_init: np.ndarray, t: float) -> dict:
    C_t = dense_state(pieces, t)
    MC = float(np.dot(fvm.W, C_t))
    MC0 = float(np.dot(fvm.W, C_init))
    flux = lambda x: HM * (float(dense_state(pieces, x)[-1]) - float(Ca_fun(x)))
    I, qerr = quad(flux, 0, t, points=Ca_fun.t[(Ca_fun.t > 0) & (Ca_fun.t < t)], epsabs=1e-14, epsrel=1e-10, limit=500)
    rhs = MC0 - R * I
    denom = max(abs(MC0 - MC), abs(R * I), 1e-18)
    return {
        "MC": MC,
        "MC0": MC0,
        "minus_R_int_hm": rhs,
        "relative_residual": abs(MC - rhs) / denom,
        "quadrature_abs_error": qerr, "quadrature": "adaptive QUADPACK on PCHIP intervals",
    }


def compare_on_output(A: np.ndarray, B: np.ndarray, times: np.ndarray) -> dict:
    diff = np.abs(A - B)
    k = int(np.argmax(diff))
    n_t, n_r = diff.shape
    i, j = divmod(k, n_r)
    return {
        "max_abs": float(diff.max()),
        "max_time_s": float(times[i]),
        "max_radius_cm": 0.1 * j,
        "mean_abs": float(diff.mean()),
        "paper_table_max_abs": float(np.abs(A[np.isin(times, TABLE_TIMES)][:, ::5] - B[np.isin(times, TABLE_TIMES)][:, ::5]).max())
        if A.shape[1] == 21
        else None,
    }


def four_decimal_stats(A: np.ndarray, B: np.ndarray, times: np.ndarray) -> dict:
    Ar = np.round(A, 4)
    Br = np.round(B, 4)
    changed = Ar != Br
    paper_t = np.isin(times, TABLE_TIMES)
    paper_r = slice(None, None, 5) if A.shape[1] == 21 else slice(None)
    paper_changed = Ar[paper_t][:, paper_r] != Br[paper_t][:, paper_r]
    early = times <= 100.0
    return {
        "changed_cells": int(np.sum(changed)),
        "total_cells": int(A.size),
        "paper_changed_cells": int(np.sum(paper_changed)),
        "paper_total_cells": int(paper_changed.size),
        "early100_changed_cells": int(np.sum(changed[early])),
        "max_abs": float(np.max(np.abs(A - B))),
        "paper_table_max_abs": float(np.max(np.abs(A[paper_t][:, paper_r] - B[paper_t][:, paper_r]))),
    }


def solve_temperature(fvm: RadialFVM, Ta_fun: PchipOven, t_breaks, rtol, atol, max_step):
    T_init = np.full(fvm.n_nodes, T0)
    def rhs(t, T):
        return fvm.heat_rhs(t, T, float(Ta_fun(t)))
    def jac(t, T):
        return fvm.heat_jac(t, T)
    return integrate_mol(
        rhs, T_init, t_breaks,
        rtol=rtol, atol=atol, max_step=max_step,
        jac=jac, jac_sparsity=fvm.heat_jac_sparsity(),
        field_name="temperature",
    )


def solve_moisture(fvm: RadialFVM, Ca_fun: PchipOven, t_breaks, rtol, atol, max_step, D_const=None):
    C_init = np.full(fvm.n_nodes, C0)
    def rhs(t, C):
        return fvm.moisture_rhs(t, C, float(Ca_fun(t)), D_const=D_const)
    return integrate_mol(
        rhs, C_init, t_breaks,
        rtol=rtol, atol=atol, max_step=max_step,
        jac_sparsity=fvm.moisture_jac_sparsity(),
        field_name="moisture",
    )


def assert_physical(T_out, C_out, Ta_fun, Ca_fun) -> dict:
    Ta_max = float(np.max(Ta_fun.y))
    Ta_min = float(np.min(Ta_fun.y))
    Ca_min = float(np.min(Ca_fun.y))
    info = {
        "T_min": float(T_out.min()),
        "T_max": float(T_out.max()),
        "C_min": float(C_out.min()),
        "C_max": float(C_out.max()),
        "T_below_28": float(max(0.0, 28.0 - T_out.min())),
        "T_above_Ta_max": float(max(0.0, T_out.max() - Ta_max)),
        "C_above_C0": float(max(0.0, C_out.max() - C0)),
        "C_below_Ca_min": float(max(0.0, Ca_min - C_out.min())),
        "Ta_range": [Ta_min, Ta_max],
        "Ca_range": [Ca_min, float(np.max(Ca_fun.y))],
    }
    info["T_bounds_ok"] = info["T_below_28"] <= 1e-6 and info["T_above_Ta_max"] <= 1e-6
    info["C_bounds_ok"] = info["C_above_C0"] <= 1e-8 and info["C_below_Ca_min"] <= 1e-8
    info["C_nonnegative"] = bool(C_out.min() >= -1e-8)
    return info


def write_result_workbook(path: Path, T_out: np.ndarray, C_out: np.ndarray, times: np.ndarray):
    radii_cm = 0.1 * np.arange(21)
    wb = Workbook()
    for title, arr in (("温度", T_out), ("水分浓度", C_out)):
        ws = wb.create_sheet(title)
        ws["A1"] = "时间\\到药材中心的距离"
        for j, rc in enumerate(radii_cm, start=2):
            ws.cell(1, j, float(rc))
        for i, t in enumerate(times):
            ws.cell(i + 2, 1, int(round(float(t))))
            for j in range(21):
                cell = ws.cell(i + 2, j + 2, float(arr[i, j]))
                cell.number_format = "0.0000"
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_paper_csv(path: Path, arr_out: np.ndarray, times: np.ndarray, radii_cm=TABLE_RADII_CM):
    t_idx = {float(t): i for i, t in enumerate(times)}
    r_idx = {round(0.1 * j, 10): j for j in range(21)}
    lines = ["t_s," + ",".join(f"r_{rc:g}cm" for rc in radii_cm)]
    for t in TABLE_TIMES:
        row = [f"{t:.0f}"]
        i = t_idx[float(t)]
        for rc in radii_cm:
            j = r_idx[float(rc)]
            row.append(f"{arr_out[i, j]:.8f}")
        lines.append(",".join(row))
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def setup_mpl():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({
        "font.sans-serif": ["Microsoft YaHei", "SimHei", "Noto Sans CJK SC", "DejaVu Sans"],
        "axes.unicode_minus": False,
        "font.size": 9,
        "axes.labelsize": 10,
        "legend.fontsize": 8,
        "xtick.labelsize": 8,
        "ytick.labelsize": 8,
        "axes.linewidth": 0.8,
        "pdf.fonttype": 42,
        "ps.fonttype": 42,
    })
    return plt


def plot_figures(plt, fvm, times, T_full, C_full, T_out, C_out, Ta_fun, Ca_fun):
    FIGURES_DIR.mkdir(parents=True, exist_ok=True)
    r_cm = fvm.r * 100.0
    r_out_cm = OUTPUT_RADII_M * 100.0
    t_marks = PAPER_TIMES
    idx = {float(t): i for i, t in enumerate(times)}

    fig, ax = plt.subplots(figsize=(6.4, 3.8))
    for t in t_marks:
        ax.plot(r_cm, T_full[idx[float(t)]], lw=1.2, label=fr"$t={int(t)}\,\mathrm{{s}}$")
    ax.set_xlabel("半径 $r$ (cm)")
    ax.set_ylabel("温度 $T$ (°C)")
    ax.set_xlim(0, 2)
    ax.legend(frameon=False, ncol=2)
    fig.tight_layout()
    fig.savefig(FIGURES_DIR / "q1_T_profiles.pdf")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(6.4, 3.8))
    for t in t_marks:
        ax.plot(r_cm, C_full[idx[float(t)]], lw=1.2, label=fr"$t={int(t)}\,\mathrm{{s}}$")
    ax.set_xlabel("半径 $r$ (cm)")
    ax.set_ylabel(r"干基含水率 $C$ (kg/kg)")
    ax.set_xlim(0, 2)
    ax.legend(frameon=False, ncol=2)
    fig.tight_layout()
    fig.savefig(FIGURES_DIR / "q1_C_profiles.pdf")
    plt.close(fig)

    t_plot = np.concatenate([[0.0], times])
    T_c = np.concatenate([[T0], T_out[:, 0]])
    T_s = np.concatenate([[T0], T_out[:, -1]])
    Ta = np.array([Ta_fun(t) for t in t_plot])
    fig, ax = plt.subplots(figsize=(6.4, 3.8))
    ax.plot(t_plot, Ta, lw=1.0, ls="--", label="烘房 $T_a$")
    ax.plot(t_plot, T_s, lw=1.2, label="表面")
    ax.plot(t_plot, T_c, lw=1.2, label="中心")
    ax.set_xlabel("时间 $t$ (s)")
    ax.set_ylabel("温度 (°C)")
    ax.set_xlim(0, 1800)
    ax.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(FIGURES_DIR / "q1_T_center_surface.pdf")
    plt.close(fig)

    C_c = np.concatenate([[C0], C_out[:, 0]])
    C_s = np.concatenate([[C0], C_out[:, -1]])
    fig, ax = plt.subplots(figsize=(6.4, 3.8))
    ax.plot(t_plot, C_s, lw=1.2, label="表面")
    ax.plot(t_plot, C_c, lw=1.2, label="中心")
    ax.set_xlabel("时间 $t$ (s)")
    ax.set_ylabel(r"干基含水率 $C$ (kg/kg)")
    ax.set_xlim(0, 1800)
    ax.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(FIGURES_DIR / "q1_C_center_surface.pdf")
    plt.close(fig)

    fig, ax1 = plt.subplots(figsize=(6.4, 3.6))
    ax2 = ax1.twinx()
    t_dense = np.linspace(0.0, T_END, 1801)
    ax1.plot(t_dense, np.array([Ta_fun(t) for t in t_dense]), color="C0", lw=1.2, label="烘房温度")
    ax2.plot(t_dense, np.array([Ca_fun(t) for t in t_dense]), color="C1", lw=1.2, ls="--", label="烘房水分")
    ax1.scatter(Ta_fun.t, Ta_fun.y, s=14, color="C0", zorder=3)
    ax2.scatter(Ca_fun.t, Ca_fun.y, s=14, color="C1", zorder=3)
    ax1.set_xlabel("时间 $t$ (s)")
    ax1.set_ylabel("烘房温度 (°C)", color="C0")
    ax2.set_ylabel("烘房水分 (kg/kg)", color="C1")
    ax1.set_xlim(0, 1800)
    h1, l1 = ax1.get_legend_handles_labels()
    h2, l2 = ax2.get_legend_handles_labels()
    ax1.legend(h1 + h2, l1 + l2, frameon=False)
    fig.tight_layout()
    fig.savefig(FIGURES_DIR / "q1_oven_input.pdf")
    plt.close(fig)
    _ = r_out_cm


def main():
    t_wall0 = time.perf_counter()
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    Ta_fun, Ca_fun, knots = load_q1_inputs()
    # PCHIP is C1; integrate as one span rather than 60 s linear kinks.
    t_breaks = np.array([0.0, T_END], dtype=float)
    print(
        "oven points", knots.size,
        "kind", Ta_fun.kind,
        "Ta(1800)=", Ta_fun(T_END),
        "Ca(1800)=", Ca_fun(T_END),
    )
    print("alpha", ALPHA, "D(C0)", moisture_diffusivity(C0), "Bi_T", H * R / K)

    mesh_levels = [2, 4, 8]
    rtol, atol_T, atol_C, max_step = 1e-10, 1e-10, 1e-12, 0.5
    solutions = {}
    for level in mesh_levels:
        r = graded_radial_nodes(level)
        fvm = RadialFVM(r=r)
        tag = f"G{level}"
        print(
            f"\n=== {tag} nodes={fvm.n_nodes} surface_dr={fvm.dr_face[-1]:.4e} m ==="
        )
        print(f"=== {tag} temperature ===")
        tT, YT, piecesT = solve_temperature(fvm, Ta_fun, t_breaks, rtol, atol_T, max_step)
        print(f"=== {tag} moisture ===")
        tC, YC, piecesC = solve_moisture(fvm, Ca_fun, t_breaks, rtol, atol_C, max_step)
        if not np.allclose(tT, tC):
            raise RuntimeError("temperature and moisture output times differ")
        times = tT[tT > 0]
        YT_pos = YT[tT > 0]
        YC_pos = YC[tC > 0]
        T_out = extract_output_grid(fvm, YT_pos)
        C_out = extract_output_grid(fvm, YC_pos)
        solutions[level] = dict(
            tag=tag,
            level=level,
            fvm=fvm,
            times=times,
            YT=YT_pos,
            YC=YC_pos,
            T_out=T_out,
            C_out=C_out,
            piecesT=piecesT,
            piecesC=piecesC,
            T_init=np.full(fvm.n_nodes, T0),
            C_init=np.full(fvm.n_nodes, C0),
            n_nodes=fvm.n_nodes,
            surface_dr=float(fvm.dr_face[-1]),
        )
        print(
            "done", tag,
            "T range", T_out.min(), T_out.max(),
            "C range", C_out.min(), C_out.max(),
            "C(t=1,s)", C_out[0, -1],
            "C(t=100,s)", C_out[99, -1],
        )

    conv = {}
    fourdec = {}
    for coarse, fine in ((2, 4), (4, 8)):
        t = solutions[fine]["times"]
        conv[f"T_G{coarse}_vs_G{fine}"] = compare_on_output(
            solutions[coarse]["T_out"], solutions[fine]["T_out"], t
        )
        conv[f"C_G{coarse}_vs_G{fine}"] = compare_on_output(
            solutions[coarse]["C_out"], solutions[fine]["C_out"], t
        )
        fourdec[f"T_G{coarse}_vs_G{fine}"] = four_decimal_stats(
            solutions[coarse]["T_out"], solutions[fine]["T_out"], t
        )
        fourdec[f"C_G{coarse}_vs_G{fine}"] = four_decimal_stats(
            solutions[coarse]["C_out"], solutions[fine]["C_out"], t
        )
        print("conv T", coarse, fine, conv[f"T_G{coarse}_vs_G{fine}"])
        print("conv C", coarse, fine, conv[f"C_G{coarse}_vs_G{fine}"])
        print("4dp T", coarse, fine, fourdec[f"T_G{coarse}_vs_G{fine}"])
        print("4dp C", coarse, fine, fourdec[f"C_G{coarse}_vs_G{fine}"])

    fd_T = fourdec["T_G4_vs_G8"]
    fd_C = fourdec["C_G4_vs_G8"]
    paper_4dp_stable = (
        fd_T["paper_changed_cells"] == 0 and fd_C["paper_changed_cells"] == 0
    )
    full_4dp_stable = fd_T["changed_cells"] == 0 and fd_C["changed_cells"] == 0
    abs_2e5 = (
        conv["T_G4_vs_G8"]["max_abs"] <= 2e-5
        and conv["C_G4_vs_G8"]["max_abs"] <= 2e-5
    )
    pub_level = 4 if paper_4dp_stable and full_4dp_stable else 8
    pub = solutions[pub_level]
    fvm = pub["fvm"]
    times = pub["times"]
    print(
        "publishing", pub["tag"],
        "n_nodes", pub["n_nodes"],
        "paper_4dp_stable", paper_4dp_stable,
        "full_4dp_stable", full_4dp_stable,
        "abs_2e5", abs_2e5,
    )

    print("time refinement on published mesh")
    tT2, YT2, _ = solve_temperature(fvm, Ta_fun, t_breaks, rtol / 5, atol_T / 5, 0.25)
    tC2, YC2, _ = solve_moisture(fvm, Ca_fun, t_breaks, rtol / 5, atol_C / 5, 0.25)
    T2 = extract_output_grid(fvm, YT2[tT2 > 0])
    C2 = extract_output_grid(fvm, YC2[tC2 > 0])
    time_sens = {
        "T": compare_on_output(pub["T_out"], T2, times),
        "C": compare_on_output(pub["C_out"], C2, times),
    }
    time_4dp = {
        "T": four_decimal_stats(pub["T_out"], T2, times),
        "C": four_decimal_stats(pub["C_out"], C2, times),
    }
    print("time sens", time_sens)
    print("time 4dp", time_4dp)

    print("Bessel temperature reference")
    Bi = H * R / K
    T_ref, heat_truncation = converged_reference(Bi, ALPHA, R, OUTPUT_RADII_M, times, Ta_fun, T0, start_terms=120)
    bessel_T = compare_on_output(pub["T_out"], T_ref, times)
    print("Bessel T", bessel_T)

    print("frozen-D moisture probe")
    D0 = float(moisture_diffusivity(C0))
    Bi_m = HM * R / D0
    C_ref, moisture_truncation = converged_reference(Bi_m, D0, R, OUTPUT_RADII_M, times, Ca_fun, C0)
    tCf, YCf, _ = solve_moisture(fvm, Ca_fun, t_breaks, rtol, atol_C, max_step, D_const=D0)
    C_frozen = extract_output_grid(fvm, YCf[tCf > 0])
    frozen = compare_on_output(C_frozen, C_ref, times)
    print("frozen D vs analytic", frozen)
    nl_vs_frozen = compare_on_output(pub["C_out"], C_frozen, times)
    print("nonlinear vs frozen D", nl_vs_frozen)
    D_field = moisture_diffusivity(pub["YC"])
    D_stats = {
        "D0": D0,
        "D_min": float(np.min(D_field)),
        "D_max": float(np.max(D_field)),
        "D_min_over_D0": float(np.min(D_field) / D0),
        "D_max_over_D0": float(np.max(D_field) / D0),
        "C_min_field": float(np.min(pub["YC"])),
        "C_max_field": float(np.max(pub["YC"])),
        "surface_C_min": float(np.min(pub["C_out"][:, -1])),
        "center_C_end": float(pub["C_out"][-1, 0]),
        "surface_C_end": float(pub["C_out"][-1, -1]),
        "surface_C_t1": float(pub["C_out"][0, -1]),
        "surface_C_t100": float(pub["C_out"][99, -1]),
    }
    print("D stats", D_stats)

    bounds = assert_physical(pub["YT"], pub["YC"], Ta_fun, Ca_fun)
    print("bounds", bounds)

    cons_T = conservation_heat(fvm, pub["piecesT"], Ta_fun, pub["T_init"], T_END)
    cons_C = conservation_moisture(fvm, pub["piecesC"], Ca_fun, pub["C_init"], T_END)
    print("conservation T", cons_T)
    print("conservation C", cons_C)

    delivery = make_gate({
        "space_abs": abs_2e5,
        "space_paper": paper_4dp_stable,
        "all_finite": np.isfinite(pub["YT"]).all() and np.isfinite(pub["YC"]).all(),
        "heat_reference_truncation": heat_truncation["passed"],
        "heat_reference_error": bessel_T["max_abs"] <= 2e-5,
        "time_abs": time_sens["T"]["max_abs"] <= 2e-5 and time_sens["C"]["max_abs"] <= 2e-5,
        "time_paper": time_4dp["T"]["paper_changed_cells"] == 0 and time_4dp["C"]["paper_changed_cells"] == 0,
        "temperature_bounds": bounds["T_bounds_ok"],
        "moisture_bounds": bounds["C_bounds_ok"],
        "moisture_reference_truncation": moisture_truncation["passed"],
        "frozen_reference_error": frozen["max_abs"] <= 2e-5,
        "heat_balance": cons_T["relative_residual"] <= 1e-6,
        "moisture_balance": cons_C["relative_residual"] <= 1e-6,
    })
    print("delivery", delivery)

    paper_T = {
        f"t{int(t)}_r{rc:g}": float(pub["T_out"][int(t) - 1, int(round(rc / 0.1))])
        for t in TABLE_TIMES for rc in TABLE_RADII_CM
    }
    paper_C = {
        f"t{int(t)}_r{rc:g}": float(pub["C_out"][int(t) - 1, int(round(rc / 0.1))])
        for t in TABLE_TIMES for rc in TABLE_RADII_CM
    }

    validation = {
        "result_version": "q1-closeout-v1",
        "interpolation": Ta_fun.kind,
        "published_mesh": pub["tag"],
        "published_n_nodes": pub["n_nodes"],
        "published_surface_dr_m": pub["surface_dr"],
        "mesh_levels": mesh_levels,
        "mesh_n_nodes": {f"G{lv}": solutions[lv]["n_nodes"] for lv in mesh_levels},
        "solver": "FVM+BDF",
        "rtol": rtol,
        "atol_T": atol_T,
        "atol_C": atol_C,
        "max_step": max_step,
        "t_breaks": t_breaks.tolist(),
        "grid_convergence": conv,
        "four_decimal_grid": fourdec,
        "time_sensitivity": time_sens,
        "four_decimal_time": time_4dp,
        "delivery_gate": delivery,
        "reference_truncation_T": heat_truncation,
        "bessel_temperature": bessel_T,
        "reference_truncation_C": moisture_truncation,
        "frozen_D_moisture_vs_analytic": frozen,
        "nonlinear_C_vs_frozen_D": nl_vs_frozen,
        "D_variation": D_stats,
        "conservation_T_1800s": cons_T,
        "conservation_C_1800s": cons_C,
        "physical_bounds": bounds,
        "paper_table_T": paper_T,
        "paper_table_C": paper_C,
        "elapsed_s": time.perf_counter() - t_wall0,
        "python": sys.version,
    }
    require_delivery(validation, RESULTS_DIR / "diagnostics/q1_delivery/latest.json")
    np.savez(
        RESULTS_DIR / "q1_solution.npz",
        N=pub["n_nodes"] - 1,
        mesh_level=pub_level,
        n_nodes=pub["n_nodes"],
        times=times,
        r=fvm.r,
        checkpoint_times=np.r_[0., TABLE_TIMES],
        T_checkpoints=np.vstack([pub["T_init"], pub["YT"][TABLE_TIMES.astype(int)-1]]),
        C_checkpoints=np.vstack([pub["C_init"], pub["YC"][TABLE_TIMES.astype(int)-1]]),
        T_full=pub["YT"],
        C_full=pub["YC"],
        T_out=pub["T_out"],
        C_out=pub["C_out"],
        output_radii_m=OUTPUT_RADII_M,
        interpolation=np.array(Ta_fun.kind),
        rtol=rtol,
        atol_T=atol_T,
        atol_C=atol_C,
        max_step=max_step,
        surface_dr=pub["surface_dr"],
    )
    write_result_workbook(RESULTS_DIR / "result1.xlsx", pub["T_out"], pub["C_out"], times)
    write_paper_csv(RESULTS_DIR / "q1_table_temperature.csv", pub["T_out"], times)
    write_paper_csv(RESULTS_DIR / "q1_table_moisture.csv", pub["C_out"], times)

    plt = setup_mpl()
    plot_figures(plt, fvm, times, pub["YT"], pub["YC"], pub["T_out"], pub["C_out"], Ta_fun, Ca_fun)

    (RESULTS_DIR / "q1_validation.json").write_text(
        json.dumps(validation, ensure_ascii=False, indent=2, default=float),
        encoding="utf-8",
    )
    print("elapsed_s", validation["elapsed_s"])
    print("wrote results to", RESULTS_DIR)



if __name__ == "__main__":
    main()
