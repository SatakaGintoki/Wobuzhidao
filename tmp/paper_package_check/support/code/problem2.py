"""Problem 2: coupled variable-property radial drying for 0–3 h."""
from __future__ import annotations

import json
import sys
import time
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from scipy.integrate import solve_ivp

CODE_DIR = Path(__file__).resolve().parent
if str(CODE_DIR) not in sys.path:
    sys.path.insert(0, str(CODE_DIR))

from appendix3 import (
    conductivity,
    conductivity_deriv,
    diffusivity_partials,
    heat_capacity,
    heat_capacity_deriv,
    moisture_diffusivity_q2,
)
from q1_analytic import BesselDuhamel, converged_reference
from delivery import make_gate, require_delivery
from radial_coupled import CoupledRadialFVM
from radial_fvm import graded_radial_nodes
from utils import (
    C0,
    FIGURES_DIR,
    H,
    HM,
    L,
    R,
    RESULTS_DIR,
    T0,
    PchipOven,
    load_oven_table,
)

T_END = 10800.0
OUTPUT_RADII_M = 0.001 * np.arange(21)
TABLE_HOURS = np.array([0.5, 1.0, 1.5, 2.0, 2.5, 3.0])
TABLE_TIMES = 3600.0 * TABLE_HOURS
TABLE_RADII_CM = np.array([0.0, 0.5, 1.0, 1.5, 2.0])
TREF = 28.0


def load_q2_inputs() -> tuple[PchipOven, PchipOven, np.ndarray]:
    t, Ta, Ca = load_oven_table()
    mask = t <= T_END + 1e-12
    t, Ta, Ca = t[mask], Ta[mask], Ca[mask]
    if t.size != 181 or abs(t[-1] - T_END) > 1e-12:
        raise ValueError(f"expected 181 oven records on [0, 10800], got {t.size}, t_end={t[-1]}")
    if abs(Ta[0] - T0) > 1e-12:
        raise ValueError(f"Ta(0)={Ta[0]} != T0={T0}")
    Ta_fun = PchipOven(t, Ta)
    Ca_fun = PchipOven(t, Ca)
    for tj, yj in zip(t, Ta):
        if abs(float(Ta_fun(tj)) - float(yj)) > 1e-12:
            raise RuntimeError(f"PCHIP misses Ta knot t={tj}")
    for tj, yj in zip(t, Ca):
        if abs(float(Ca_fun(tj)) - float(yj)) > 1e-12:
            raise RuntimeError(f"PCHIP misses Ca knot t={tj}")
    return Ta_fun, Ca_fun, t


def integrate_coupled(fvm: CoupledRadialFVM, y0, t_end, Ta_fun, Ca_fun, rtol, atol, max_step,
                      D_const=None, source=None, t_eval=None):
    y0 = np.asarray(y0, dtype=float)
    n = fvm.n_nodes
    atol_arr = np.asarray(atol, dtype=float)
    if atol_arr.ndim == 0:
        atol_arr = np.concatenate([np.full(n, float(atol_arr)), np.full(n, float(atol_arr) * 0.01)])

    def rhs(t, y):
        ST = SC = None
        if source is not None:
            ST, SC = source(t, fvm.r)
        return fvm.rhs(t, y, float(Ta_fun(t)), float(Ca_fun(t)), D_const=D_const, ST=ST, SC=SC)

    if t_eval is None:
        t_eval = np.arange(1.0, t_end + 0.5, 1.0)
    sol = solve_ivp(
        rhs,
        (0.0, float(t_end)),
        y0,
        method="BDF",
        rtol=rtol,
        atol=atol_arr,
        max_step=max_step,
        t_eval=t_eval,
        dense_output=False,
        vectorized=False,
        jac_sparsity=fvm.jac_sparsity(),
    )
    if not sol.success:
        raise RuntimeError(f"coupled BDF failed: {sol.message}")
    times = np.concatenate([[0.0], sol.t])
    Y = np.vstack([y0, sol.y.T])
    return times, Y


def extract_output(fvm: CoupledRadialFVM, Y: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    n = fvm.n_nodes
    idx = fvm.output_indices()
    T = Y[:, :n][:, idx]
    C = Y[:, n:][:, idx]
    return T, C


def compare_on_output(A, B, times) -> dict:
    diff = np.abs(A - B)
    k = int(np.argmax(diff))
    n_t, n_r = diff.shape
    i, j = divmod(k, n_r)
    paper_t = np.isin(times, TABLE_TIMES)
    paper_r = slice(None, None, 5) if n_r == 21 else slice(None)
    return {
        "max_abs": float(diff.max()),
        "max_time_s": float(times[i]),
        "max_radius_cm": 0.1 * j,
        "mean_abs": float(diff.mean()),
        "paper_table_max_abs": float(np.max(np.abs(A[paper_t][:, paper_r] - B[paper_t][:, paper_r]))),
    }


def four_decimal_stats(A, B, times) -> dict:
    Ar, Br = np.round(A, 4), np.round(B, 4)
    changed = Ar != Br
    paper_t = np.isin(times, TABLE_TIMES)
    paper_r = slice(None, None, 5) if A.shape[1] == 21 else slice(None)
    paper_changed = Ar[paper_t][:, paper_r] != Br[paper_t][:, paper_r]
    return {
        "changed_cells": int(np.sum(changed)),
        "total_cells": int(A.size),
        "paper_changed_cells": int(np.sum(paper_changed)),
        "paper_total_cells": int(paper_changed.size),
        "max_abs": float(np.max(np.abs(A - B))),
        "paper_table_max_abs": float(np.max(np.abs(A[paper_t][:, paper_r] - B[paper_t][:, paper_r]))),
    }


def manufactured_fields(r: np.ndarray, t: float) -> tuple[np.ndarray, np.ndarray]:
    x2 = (r / R) ** 2
    T = 28.0 + 2.0 * np.exp(-t / 600.0) * (1.0 + x2)
    C = 1.5 + 0.2 * np.exp(-t / 500.0) * (1.0 + x2)
    return T, C


def manufactured_env(r_surface: float, t: float) -> tuple[float, float]:
    T, C = manufactured_fields(np.array([r_surface]), t)
    Ts, Cs = float(T[0]), float(C[0])
    u = np.exp(-t / 600.0)
    v = np.exp(-t / 500.0)
    Tr = 4.0 / R * u
    Cr = 4.0 / R * v * 0.1  # C amplitude 0.2 vs T amplitude 2: Cr = 0.2*v*2r/R^2 at r=R -> 0.4 v / R
    Cr = 0.4 / R * v
    kn = float(conductivity(Cs))
    Dn = float(moisture_diffusivity_q2(Cs, Ts))
    Ta = Ts + kn * Tr / H
    Ca = Cs + Dn * Cr / HM
    return Ta, Ca


class ManufacturedOven:
    def __init__(self, kind: str):
        self.kind = kind

    def __call__(self, t):
        Ta, Ca = manufactured_env(R, float(t))
        return Ta if self.kind == "T" else Ca


def manufactured_sources(t: float, r: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    T, C = manufactured_fields(r, t)
    x2 = (r / R) ** 2
    u = np.exp(-t / 600.0)
    v = np.exp(-t / 500.0)
    Tt = 2.0 * (-1.0 / 600.0) * u * (1.0 + x2)
    Ct = 0.2 * (-1.0 / 500.0) * v * (1.0 + x2)
    Tr = 4.0 * r / (R ** 2) * u
    Trr = np.full_like(r, 4.0 / (R ** 2) * u)
    Cr = 0.4 * r / (R ** 2) * v
    Crr = np.full_like(r, 0.4 / (R ** 2) * v)
    lapT = np.empty_like(r)
    lapC = np.empty_like(r)
    center = r < 1e-16
    lapT[center] = 2.0 * Trr[center]
    lapC[center] = 2.0 * Crr[center]
    away = ~center
    lapT[away] = Trr[away] + Tr[away] / r[away]
    lapC[away] = Crr[away] + Cr[away] / r[away]
    kn = np.asarray(conductivity(C), dtype=float)
    kp = np.asarray(conductivity_deriv(C), dtype=float)
    Bn = np.asarray(heat_capacity(C), dtype=float)
    D, dD_dC, dD_dT = diffusivity_partials(C, T)
    ST = Bn * Tt - (kn * lapT + kp * Cr * Tr)
    SC = Ct - (D * lapC + dD_dC * Cr ** 2 + dD_dT * Tr * Cr)
    return ST, SC


def conservation_moisture(fvm, times, Y, Ca_fun) -> dict:
    n = fvm.n_nodes
    C = Y[:, n:]
    M = C @ fvm.W
    M0 = float(M[0])
    CN = C[:, -1]
    Ca = np.array([float(Ca_fun(t)) for t in times])
    I = float(np.trapezoid(HM * (CN - Ca), times))
    rhs = M0 - R * I
    Mt = float(M[-1])
    denom = max(abs(M0 - Mt), abs(R * I), 1e-18)
    return {"M0": M0, "M_end": Mt, "rhs": rhs, "relative_residual": abs(Mt - rhs) / denom}


def conservation_heat(fvm, times, Y, Ta_fun) -> dict:
    n = fvm.n_nodes
    T = Y[:, :n]
    C = Y[:, n:]
    Bn = np.asarray(heat_capacity(C), dtype=float)
    Bp = np.asarray(heat_capacity_deriv(C), dtype=float)
    E = 2.0 * np.pi * L * np.sum(fvm.W[None, :] * Bn * (T - TREF), axis=1)
    TN = T[:, -1]
    Ta = np.array([float(Ta_fun(t)) for t in times])
    Q = 2.0 * np.pi * R * L * H * float(np.trapezoid(Ta - TN, times))
    dC = np.gradient(C, times, axis=0)
    corr = 2.0 * np.pi * L * np.sum(fvm.W[None, :] * Bp * (T - TREF) * dC, axis=1)
    Icorr = float(np.trapezoid(corr, times))
    lhs = float(E[-1] - E[0] - Icorr)
    denom = max(abs(Q), abs(lhs), 1e-12)
    return {
        "E_end_minus_E0": float(E[-1] - E[0]),
        "composition_correction": Icorr,
        "Q_surface": Q,
        "relative_residual": abs(lhs - Q) / denom,
    }


def assert_physical(T_out, C_out, Ta_fun, Ca_fun) -> dict:
    Ta_max = float(np.max(Ta_fun.y))
    Ca_min = float(np.min(Ca_fun.y))
    info = {
        "T_min": float(T_out.min()),
        "T_max": float(T_out.max()),
        "C_min": float(C_out.min()),
        "C_max": float(C_out.max()),
        "T_below_28": float(max(0.0, 28.0 - T_out.min())),
        "T_above_Ta_max": float(max(0.0, T_out.max() - Ta_max)),
        "C_above_C0": float(max(0.0, C_out.max() - C0)),
        "C_below_Ca_min": float(max(0.0, Ca_min - C_out.min())),
        "Ta_range": [float(np.min(Ta_fun.y)), Ta_max],
        "Ca_range": [Ca_min, float(np.max(Ca_fun.y))],
    }
    info["T_bounds_ok"] = info["T_below_28"] <= 1e-6 and info["T_above_Ta_max"] <= 1e-4
    info["C_bounds_ok"] = info["C_above_C0"] <= 1e-8 and info["C_below_Ca_min"] <= 1e-8
    return info


def write_result_workbook(path: Path, T_out, C_out, times):
    radii_cm = 0.1 * np.arange(21)
    wb = Workbook()
    for title, arr in (("温度", T_out), ("水分浓度", C_out)):
        ws = wb.create_sheet(title)
        ws["A1"] = "时间\\到药材中心的距离"
        for j, rc in enumerate(radii_cm, start=2):
            ws.cell(1, j, float(rc))
        for i, t in enumerate(times):
            ws.cell(i + 2, 1, int(round(float(t))))
            for j in range(21):
                cell = ws.cell(i + 2, j + 2, float(arr[i, j]))
                cell.number_format = "0.0000"
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_paper_csv(path: Path, arr_out, times):
    t_idx = {float(t): i for i, t in enumerate(times)}
    r_idx = {round(0.1 * j, 10): j for j in range(21)}
    lines = ["t_h," + ",".join(f"r_{rc:g}cm" for rc in TABLE_RADII_CM)]
    for th, t in zip(TABLE_HOURS, TABLE_TIMES):
        row = [f"{th:g}"]
        i = t_idx[float(t)]
        for rc in TABLE_RADII_CM:
            j = r_idx[float(rc)]
            row.append(f"{arr_out[i, j]:.8f}")
        lines.append(",".join(row))
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def setup_mpl():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({
        "font.sans-serif": ["Microsoft YaHei", "SimHei", "Noto Sans CJK SC", "DejaVu Sans"],
        "axes.unicode_minus": False,
        "font.size": 9,
        "axes.labelsize": 10,
        "legend.fontsize": 8,
        "xtick.labelsize": 8,
        "ytick.labelsize": 8,
        "axes.linewidth": 0.8,
        "pdf.fonttype": 42,
        "ps.fonttype": 42,
    })
    return plt


def plot_figures(plt, fvm, times, Y, T_out, C_out, Ta_fun, Ca_fun):
    FIGURES_DIR.mkdir(parents=True, exist_ok=True)
    n = fvm.n_nodes
    r_cm = fvm.r * 100.0
    idx = {float(t): i for i, t in enumerate(times)}
    marks = TABLE_TIMES

    fig, ax = plt.subplots(figsize=(6.4, 3.8))
    for t in marks:
        ax.plot(r_cm, Y[idx[float(t)], :n], lw=1.2, label=fr"$t={t/3600:g}\,\mathrm{{h}}$")
    ax.set_xlabel("半径 $r$ (cm)")
    ax.set_ylabel("温度 $T$ (°C)")
    ax.set_xlim(0, 2)
    ax.legend(frameon=False, ncol=3, loc="lower center", bbox_to_anchor=(0.5, 1.02))
    fig.tight_layout()
    fig.savefig(FIGURES_DIR / "q2_T_profiles.pdf")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(6.4, 3.8))
    for t in marks:
        ax.plot(r_cm, Y[idx[float(t)], n:], lw=1.2, label=fr"$t={t/3600:g}\,\mathrm{{h}}$")
    ax.set_xlabel("半径 $r$ (cm)")
    ax.set_ylabel(r"干基含水率 $C$ (kg/kg)")
    ax.set_xlim(0, 2)
    ax.legend(frameon=False, ncol=2)
    fig.tight_layout()
    fig.savefig(FIGURES_DIR / "q2_C_profiles.pdf")
    plt.close(fig)

    t_plot = times
    fig, ax = plt.subplots(figsize=(6.4, 3.8))
    ax.plot(t_plot / 3600.0, [Ta_fun(t) for t in t_plot], lw=1.0, ls="--", label="烘房 $T_a$")
    ax.plot(t_plot / 3600.0, T_out[:, -1], lw=1.2, label="表面")
    ax.plot(t_plot / 3600.0, T_out[:, 0], lw=1.2, label="中心")
    ax.set_xlabel("时间 $t$ (h)")
    ax.set_ylabel("温度 (°C)")
    ax.set_xlim(0, 3)
    ax.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(FIGURES_DIR / "q2_T_center_surface.pdf")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(6.4, 3.8))
    ax.plot(t_plot / 3600.0, C_out[:, -1], lw=1.2, label="表面")
    ax.plot(t_plot / 3600.0, C_out[:, 0], lw=1.2, label="中心")
    ax.set_xlabel("时间 $t$ (h)")
    ax.set_ylabel(r"干基含水率 $C$ (kg/kg)")
    ax.set_xlim(0, 3)
    ax.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(FIGURES_DIR / "q2_C_center_surface.pdf")
    plt.close(fig)

    D_c = moisture_diffusivity_q2(C_out[:, 0], T_out[:, 0])
    D_s = moisture_diffusivity_q2(C_out[:, -1], T_out[:, -1])
    fig, ax = plt.subplots(figsize=(6.4, 3.6))
    ax.plot(t_plot / 3600.0, D_s, lw=1.2, label="表面 $D$")
    ax.plot(t_plot / 3600.0, D_c, lw=1.2, label="中心 $D$")
    ax.set_xlabel("时间 $t$ (h)")
    ax.set_ylabel(r"$D$ (m$^2$/s)")
    ax.set_xlim(0, 3)
    ax.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(FIGURES_DIR / "q2_D_center_surface.pdf")
    plt.close(fig)


def run_mms(level: int = 4) -> dict:
    r = graded_radial_nodes(level)
    fvm = CoupledRadialFVM(r)
    T0f, C0f = manufactured_fields(fvm.r, 0.0)
    y0 = fvm.pack(T0f, C0f)
    Ta_fun = ManufacturedOven("T")
    Ca_fun = ManufacturedOven("C")
    times, Y = integrate_coupled(
        fvm, y0, 100.0, Ta_fun, Ca_fun,
        rtol=1e-8, atol=1e-9, max_step=0.5,
        source=manufactured_sources,
        t_eval=np.arange(1.0, 100.5, 1.0),
    )
    n = fvm.n_nodes
    T_num, C_num = Y[:, :n], Y[:, n:]
    T_ex = np.stack([manufactured_fields(fvm.r, t)[0] for t in times])
    C_ex = np.stack([manufactured_fields(fvm.r, t)[1] for t in times])
    return {
        "level": level,
        "n_nodes": n,
        "T_max_abs": float(np.max(np.abs(T_num - T_ex))),
        "C_max_abs": float(np.max(np.abs(C_num - C_ex))),
        "T_end_max_abs": float(np.max(np.abs(T_num[-1] - T_ex[-1]))),
        "C_end_max_abs": float(np.max(np.abs(C_num[-1] - C_ex[-1]))),
    }


def main():
    t_wall0 = time.perf_counter()
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    Ta_fun, Ca_fun, knots = load_q2_inputs()
    print("oven points", knots.size, "kind", Ta_fun.kind, "Ta(10800)=", Ta_fun(T_END), "Ca(10800)=", Ca_fun(T_END))
    D0 = float(moisture_diffusivity_q2(C0, T0))
    print("initial D", D0, "k", float(conductivity(C0)), "B", float(heat_capacity(C0)))

    print("manufactured-solution probe")
    mms = run_mms(4)
    print("MMS", mms)

    mesh_levels = [2, 4]
    rtol, atol_T, atol_C, max_step = 1e-9, 1e-8, 1e-10, 1.0
    solutions = {}
    for level in mesh_levels:
        r = graded_radial_nodes(level)
        fvm = CoupledRadialFVM(r)
        tag = f"G{level}"
        print(f"\n=== {tag} nodes={fvm.n_nodes} ===")
        y0 = fvm.pack(np.full(fvm.n_nodes, T0), np.full(fvm.n_nodes, C0))
        atol = np.concatenate([np.full(fvm.n_nodes, atol_T), np.full(fvm.n_nodes, atol_C)])
        times, Y = integrate_coupled(fvm, y0, T_END, Ta_fun, Ca_fun, rtol, atol, max_step)
        times_pos = times[times > 0]
        Y_pos = Y[times > 0]
        T_out, C_out = extract_output(fvm, Y_pos)
        solutions[level] = dict(
            tag=tag, fvm=fvm, times=times_pos, Y=Y_pos, Y_full=Y, times_full=times,
            T_out=T_out, C_out=C_out, n_nodes=fvm.n_nodes,
        )
        print("done", tag, "T", T_out.min(), T_out.max(), "C", C_out.min(), C_out.max(),
              "C(t=1,s)", C_out[0, -1], "C(3h,s)", C_out[-1, -1], "C(3h,c)", C_out[-1, 0])

    t = solutions[4]["times"]
    conv = {
        "T_G2_vs_G4": compare_on_output(solutions[2]["T_out"], solutions[4]["T_out"], t),
        "C_G2_vs_G4": compare_on_output(solutions[2]["C_out"], solutions[4]["C_out"], t),
    }
    fourdec = {
        "T_G2_vs_G4": four_decimal_stats(solutions[2]["T_out"], solutions[4]["T_out"], t),
        "C_G2_vs_G4": four_decimal_stats(solutions[2]["C_out"], solutions[4]["C_out"], t),
    }
    print("conv", conv)
    print("4dp", fourdec)
    paper_ok = fourdec["T_G2_vs_G4"]["paper_changed_cells"] == 0 and fourdec["C_G2_vs_G4"]["paper_changed_cells"] == 0
    abs_ok = conv["T_G2_vs_G4"]["max_abs"] <= 2e-5 and conv["C_G2_vs_G4"]["max_abs"] <= 2e-5
    need_g8 = not (paper_ok and abs_ok)
    if need_g8:
        print("G2 vs G4 not enough; running G8")
        level = 8
        r = graded_radial_nodes(level)
        fvm = CoupledRadialFVM(r)
        y0 = fvm.pack(np.full(fvm.n_nodes, T0), np.full(fvm.n_nodes, C0))
        atol = np.concatenate([np.full(fvm.n_nodes, atol_T), np.full(fvm.n_nodes, atol_C)])
        times, Y = integrate_coupled(fvm, y0, T_END, Ta_fun, Ca_fun, rtol, atol, max_step)
        times_pos = times[times > 0]
        Y_pos = Y[times > 0]
        T_out, C_out = extract_output(fvm, Y_pos)
        solutions[8] = dict(
            tag="G8", fvm=fvm, times=times_pos, Y=Y_pos, Y_full=Y, times_full=times,
            T_out=T_out, C_out=C_out, n_nodes=fvm.n_nodes,
        )
        t = solutions[8]["times"]
        conv["T_G4_vs_G8"] = compare_on_output(solutions[4]["T_out"], solutions[8]["T_out"], t)
        conv["C_G4_vs_G8"] = compare_on_output(solutions[4]["C_out"], solutions[8]["C_out"], t)
        fourdec["T_G4_vs_G8"] = four_decimal_stats(solutions[4]["T_out"], solutions[8]["T_out"], t)
        fourdec["C_G4_vs_G8"] = four_decimal_stats(solutions[4]["C_out"], solutions[8]["C_out"], t)
        print("conv G4 G8", conv["T_G4_vs_G8"], conv["C_G4_vs_G8"])
        pub_level = 8
        paper_ok = fourdec["T_G4_vs_G8"]["paper_changed_cells"] == 0 and fourdec["C_G4_vs_G8"]["paper_changed_cells"] == 0
        abs_ok = conv["T_G4_vs_G8"]["max_abs"] <= 2e-5 and conv["C_G4_vs_G8"]["max_abs"] <= 2e-5
    else:
        pub_level = 4 if fourdec["C_G2_vs_G4"]["changed_cells"] or fourdec["T_G2_vs_G4"]["changed_cells"] else 2
        if pub_level == 2:
            pub_level = 4
    pub = solutions[pub_level]
    fvm = pub["fvm"]
    times = pub["times"]
    print("publishing", pub["tag"], "paper_ok", paper_ok, "abs_ok", abs_ok)

    print("time refinement")
    y0 = fvm.pack(np.full(fvm.n_nodes, T0), np.full(fvm.n_nodes, C0))
    atol = np.concatenate([np.full(fvm.n_nodes, atol_T / 5), np.full(fvm.n_nodes, atol_C / 5)])
    t2, Y2 = integrate_coupled(fvm, y0, T_END, Ta_fun, Ca_fun, rtol / 5, atol, 0.5)
    T2, C2 = extract_output(fvm, Y2[t2 > 0])
    time_sens = {"T": compare_on_output(pub["T_out"], T2, times), "C": compare_on_output(pub["C_out"], C2, times)}
    time_4dp = {"T": four_decimal_stats(pub["T_out"], T2, times), "C": four_decimal_stats(pub["C_out"], C2, times)}
    print("time sens", time_sens)

    print("frozen-D moisture probe")
    D_freeze = D0
    Bi_m = HM * R / D_freeze
    C_ref, moisture_truncation = converged_reference(Bi_m, D_freeze, R, OUTPUT_RADII_M, times, Ca_fun, C0)
    y0f = fvm.pack(np.full(fvm.n_nodes, T0), np.full(fvm.n_nodes, C0))
    atol = np.concatenate([np.full(fvm.n_nodes, atol_T), np.full(fvm.n_nodes, atol_C)])
    tf, Yf = integrate_coupled(fvm, y0f, T_END, Ta_fun, Ca_fun, rtol, atol, max_step, D_const=D_freeze)
    _, C_frozen = extract_output(fvm, Yf[tf > 0])
    frozen = compare_on_output(C_frozen, C_ref, times)
    nl_vs_frozen = compare_on_output(pub["C_out"], C_frozen, times)
    print("frozen vs analytic", frozen)
    print("nonlinear vs frozen", nl_vs_frozen)

    n = fvm.n_nodes
    D_field = moisture_diffusivity_q2(pub["Y"][:, n:], pub["Y"][:, :n])
    D_stats = {
        "D0": D0,
        "D_min": float(np.min(D_field)),
        "D_max": float(np.max(D_field)),
        "D_min_over_D0": float(np.min(D_field) / D0),
        "D_max_over_D0": float(np.max(D_field) / D0),
        "center_C_end": float(pub["C_out"][-1, 0]),
        "surface_C_end": float(pub["C_out"][-1, -1]),
        "center_T_end": float(pub["T_out"][-1, 0]),
        "surface_T_end": float(pub["T_out"][-1, -1]),
    }
    print("D stats", D_stats)

    bounds = assert_physical(pub["Y_full"][:, :n], pub["Y_full"][:, n:], Ta_fun, Ca_fun)
    cons_C = conservation_moisture(fvm, pub["times_full"], pub["Y_full"], Ca_fun)
    cons_T = conservation_heat(fvm, pub["times_full"], pub["Y_full"], Ta_fun)
    print("bounds", bounds)
    print("conservation C", cons_C)
    print("conservation T", cons_T)

    delivery = make_gate({
        "space_abs": abs_ok,
        "space_paper": paper_ok,
        "all_finite": np.isfinite(pub["Y_full"]).all(),
        "manufactured_solution": mms["T_max_abs"] <= 2e-5 and mms["C_max_abs"] <= 2e-5,
        "time_abs": time_sens["T"]["max_abs"] <= 2e-5 and time_sens["C"]["max_abs"] <= 2e-5,
        "time_paper": time_4dp["T"]["paper_changed_cells"] == 0 and time_4dp["C"]["paper_changed_cells"] == 0,
        "temperature_bounds": bounds["T_bounds_ok"],
        "moisture_bounds": bounds["C_bounds_ok"],
        "moisture_reference_truncation": moisture_truncation["passed"],
        "frozen_reference_error": frozen["max_abs"] <= 2e-5,
        "heat_balance": cons_T["relative_residual"] <= 1e-6,
        "moisture_balance": cons_C["relative_residual"] <= 1e-6,
    })
    print("delivery", delivery)

    paper_T = {f"h{th:g}_r{rc:g}": float(pub["T_out"][int(t) - 1, int(round(rc / 0.1))])
               for th, t in zip(TABLE_HOURS, TABLE_TIMES) for rc in TABLE_RADII_CM}
    paper_C = {f"h{th:g}_r{rc:g}": float(pub["C_out"][int(t) - 1, int(round(rc / 0.1))])
               for th, t in zip(TABLE_HOURS, TABLE_TIMES) for rc in TABLE_RADII_CM}

    validation = {
        "result_version": "q2-closeout-v1",
        "interpolation": Ta_fun.kind,
        "published_mesh": pub["tag"],
        "published_n_nodes": pub["n_nodes"],
        "mesh_levels": list(solutions.keys()),
        "solver": "coupled-FVM+BDF",
        "rtol": rtol, "atol_T": atol_T, "atol_C": atol_C, "max_step": max_step,
        "mms": mms,
        "grid_convergence": conv,
        "four_decimal_grid": fourdec,
        "time_sensitivity": time_sens,
        "four_decimal_time": time_4dp,
        "delivery_gate": delivery,
        "reference_truncation_C": moisture_truncation,
        "frozen_D_moisture_vs_analytic": frozen,
        "nonlinear_C_vs_frozen_D": nl_vs_frozen,
        "D_variation": D_stats,
        "conservation_T_10800s": cons_T,
        "conservation_C_10800s": cons_C,
        "physical_bounds": bounds,
        "paper_table_T": paper_T,
        "paper_table_C": paper_C,
        "elapsed_s": time.perf_counter() - t_wall0,
        "python": sys.version,
        "end_face_2d": "skipped_by_plan",
    }
    require_delivery(validation, RESULTS_DIR / "diagnostics/q2_delivery/latest.json")
    np.savez(
        RESULTS_DIR / "q2_solution.npz",
        mesh_level=pub_level,
        n_nodes=pub["n_nodes"],
        times=times,
        r=fvm.r,
        output_radii_m=OUTPUT_RADII_M,
        checkpoint_times=np.r_[0., 1., 60., 100., TABLE_TIMES],
        T_checkpoints=pub["Y_full"][np.r_[0., 1., 60., 100., TABLE_TIMES].astype(int), :n],
        C_checkpoints=pub["Y_full"][np.r_[0., 1., 60., 100., TABLE_TIMES].astype(int), n:],
        T_out=pub["T_out"],
        C_out=pub["C_out"],
        interpolation=np.array(Ta_fun.kind),
        rtol=rtol, atol_T=atol_T, atol_C=atol_C, max_step=max_step,
    )
    write_result_workbook(RESULTS_DIR / "result2.xlsx", pub["T_out"], pub["C_out"], times)
    write_paper_csv(RESULTS_DIR / "q2_table_temperature.csv", pub["T_out"], times)
    write_paper_csv(RESULTS_DIR / "q2_table_moisture.csv", pub["C_out"], times)
    plt = setup_mpl()
    plot_figures(plt, fvm, times, pub["Y"], pub["T_out"], pub["C_out"], Ta_fun, Ca_fun)

    (RESULTS_DIR / "q2_validation.json").write_text(
        json.dumps(validation, ensure_ascii=False, indent=2, default=float),
        encoding="utf-8",
    )
    print("elapsed_s", validation["elapsed_s"])
    print("wrote results to", RESULTS_DIR)



if __name__ == "__main__":
    main()
