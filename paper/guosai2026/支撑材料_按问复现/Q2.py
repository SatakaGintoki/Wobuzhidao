# -*- coding: utf-8 -*-
"""问题2：独立求解并导出答案。运行 python Q2.py。
仅依赖numpy、scipy、openpyxl和同目录data附件。精简版未重新运行核验。
"""
from pathlib import Path
import csv
import json
import numpy as np
from openpyxl import load_workbook
from scipy.integrate import solve_ivp
from scipy.interpolate import PchipInterpolator
from scipy.sparse import diags, lil_matrix

ROOT = Path(__file__).resolve().parent
ATTACH = ROOT / "data" / "附件"
OUT = ROOT / "results"
R, H, HM, T0, C0 = 0.02, 25.0, 8e-7, 28.0, 2.55


def read_table(name):
    wb = load_workbook(ATTACH / name, read_only=True, data_only=True)
    rows = np.asarray(list(wb.active.values)[1:], dtype=float)
    wb.close()
    return rows


def environment(end):
    rows = read_table("附件1.xlsx")
    rows = rows[rows[:, 0] <= end]
    return (PchipInterpolator(rows[:, 0], rows[:, 1], extrapolate=False),
            PchipInterpolator(rows[:, 0], rows[:, 2], extrapolate=False))


def save_excel(question, times, fields):
    wb = load_workbook(ATTACH / "附件3" / f"result{question}.xlsx")
    for ws, field in zip(wb.worksheets, fields):
        header, surface = ws.cell(1, 1).value, ws.cell(1, ws.max_column).value
        ws.delete_rows(1, ws.max_row)
        ws.append([header] + [j / 10 for j in range(21)] + ([surface] if question == 4 else []))
        for t, row in zip(times, field):
            ws.append([float(t)] + [None if np.isnan(v) else float(v) for v in row])
        for row in ws.iter_rows(min_row=2, min_col=2):
            for cell in row:
                cell.number_format = "0.0000"
        ws.freeze_panes = "B2"
    OUT.mkdir(exist_ok=True)
    wb.properties.creator = wb.properties.lastModifiedBy = "Team"
    wb.save(OUT / f"result{question}.xlsx")
    wb.close()
    print(f"已生成 results/result{question}.xlsx", flush=True)


def save_csv(name, times, values, headers):
    with (OUT / name).open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.writer(stream)
        writer.writerow(["时间"] + headers)
        for t, row in zip(times, values):
            writer.writerow([f"{t:.4f}"] + ["" if np.isnan(v) else f"{v:.4f}" for v in row])


def integrate(rhs, interval, initial, **options):
    solution = solve_ivp(rhs, interval, initial, method="BDF", **options)
    if not solution.success:
        raise RuntimeError(solution.message)
    return solution


def geometry(nodes):
    faces = (nodes[:-1] + nodes[1:]) / 2
    return faces, np.diff(nodes), np.diff(np.r_[nodes[0], faces, nodes[-1]]**2) / 2


def graded_radial_nodes(level: int) -> np.ndarray:
    """Nodes denser near the surface; every 0.1 cm output radius remains a node.

    Interval subdivisions (before multiplying by ``level``):
    0–1.5 cm: 8; 1.5–1.8 cm: 16; 1.8–1.9 cm: 32; 1.9–2.0 cm: 64.
    Surface spacing at level L is 1e-3 / (64 L) metres.
    """
    level = int(level)
    if level < 1:
        raise ValueError("graded mesh level must be a positive integer")
    n_sub = np.full(20, 8 * level, dtype=int)
    n_sub[15:] = 16 * level
    n_sub[18] = 32 * level
    n_sub[19] = 64 * level
    pts = [0.0]
    for j in range(20):
        r0 = 0.001 * j
        r1 = 0.001 * (j + 1)
        m = int(n_sub[j])
        for k in range(1, m + 1):
            pts.append(r0 + (r1 - r0) * k / m)
    r = np.asarray(pts, dtype=float)
    r[0] = 0.0
    r[-1] = R
    return r

def harmonic_mean(a: np.ndarray, b: np.ndarray) -> np.ndarray:
    out = np.zeros_like(a, dtype=float)
    both = (a > 0.0) & (b > 0.0)
    out[both] = 2.0 * a[both] * b[both] / (a[both] + b[both])
    return out

def coupled_sparsity(n):
    matrix = lil_matrix((2*n, 2*n))
    for i in range(n):
        for j in range(max(0, i-1), min(n, i+2)):
            matrix[i, j] = matrix[i, n+j] = matrix[n+i, j] = matrix[n+i, n+j] = 1
    return matrix.tocsc()


def properties(C, T):
    rho = 650.0 + 128.0 * C
    cp = 1450.0 + 2736.0 * C / (1.0 + C)
    k = 0.21 + 0.38 * C / (1.0 + C)
    D = np.zeros_like(C)
    good = C > 0
    D[good] = 2.4e-3 * np.exp(-0.45 / C[good]) * np.exp(-3850.0 / (T[good] + 273.15))
    return rho * cp, k, D


def coupled_rhs(nodes):
    n = len(nodes)
    faces, dr, weights = geometry(nodes)
    def rhs(y, Ta, Ca):
        T, C = y[:n], y[n:]
        B, k, D = properties(C, T)
        FT = -faces * harmonic_mean(k[:-1], k[1:]) * np.diff(T) / dr
        FC = -faces * harmonic_mean(D[:-1], D[1:]) * np.diff(C) / dr
        td, cd = np.empty(n), np.empty(n)
        td[0], cd[0] = -FT[0] / (B[0]*weights[0]), -FC[0] / weights[0]
        td[1:-1] = (FT[:-1]-FT[1:]) / (B[1:-1]*weights[1:-1])
        cd[1:-1] = (FC[:-1]-FC[1:]) / weights[1:-1]
        td[-1] = (FT[-1]-R*H*(T[-1]-Ta)) / (B[-1]*weights[-1])
        cd[-1] = (FC[-1]-R*HM*(C[-1]-Ca)) / weights[-1]
        return np.r_[td, cd]
    return rhs


def main():
    Ta, Ca = environment(10800.0)
    r = graded_radial_nodes(4)
    n = len(r)
    operator = coupled_rhs(r)
    rhs = lambda t, y: operator(y, float(Ta(t)), float(Ca(t)))
    times = np.arange(1.0, 10801.0)
    print("问题二：计算变物性热湿耦合...", flush=True)
    solution = integrate(rhs, (0, 10800), np.r_[np.full(n, T0), np.full(n, C0)],
                         rtol=1e-9, atol=np.r_[np.full(n, 1e-8), np.full(n, 1e-10)],
                         max_step=1.0, t_eval=times, jac_sparsity=coupled_sparsity(n))
    ids = np.array([int(np.argmin(abs(r-j*.001))) for j in range(21)])
    T, C = solution.y[ids].T, solution.y[n+ids].T
    save_excel(2, times, [T, C])
    tt = np.arange(1800, 10801, 1800)
    for label, values in (("temperature", T), ("moisture", C)):
        save_csv(f"q2_table_{label}.csv", tt/3600, values[tt-1, ::5], ["0 cm", "0.5 cm", "1 cm", "1.5 cm", "2 cm"])


if __name__ == "__main__":
    main()
