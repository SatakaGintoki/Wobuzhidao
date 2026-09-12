"""Regression and exhaustive delivery checks for Q3 closeout."""
from __future__ import annotations
import copy
import hashlib
import json
import tempfile
from pathlib import Path
from unittest.mock import patch
import numpy as np
from openpyxl import load_workbook
import problem3
import q3_closeout
from utils import ROOT,RESULTS_DIR
from q3_balance import combine


def main():
    v=json.loads((RESULTS_DIR/'q3_validation.json').read_text(encoding='utf-8'))
    assert v['result_version']=='q3-closeout-v1'
    assert v['delivery_gate']['meets_delivery_gate'] and all(v['delivery_gate']['checks'].values())
    checked=[]
    with tempfile.TemporaryDirectory(prefix='q3_gate_',dir=ROOT/'tmp') as td:
        root=Path(td);out=root/'result3.xlsx';out.write_bytes(b'ORIGINAL')
        for name in v['delivery_gate']['checks']:
            bad=copy.deepcopy(v);bad['delivery_gate']['checks'][name]=False
            # Deliberately leave the top-level flag true to test recomputation.
            with patch.object(problem3,'RESULTS_DIR',root),patch.object(q3_closeout,'export_workbook') as writer:
                try:
                    problem3.write_result3_xlsx(out,np.zeros((1,21)),np.array([60.]),validation=bad)
                    raise AssertionError(f'failed gate allowed writing: {name}')
                except RuntimeError:
                    pass
                assert not writer.called and out.read_bytes()==b'ORIGINAL'
            checked.append(name)
            with patch.object(q3_closeout,'DIAG',root/'diagnostics'):
                try:
                    q3_closeout.publish_results(None, None, bad, None)
                    raise AssertionError('failed gate reached production publisher')
                except RuntimeError:
                    pass
            missing=copy.deepcopy(v);del missing['delivery_gate']['checks'][name]
            with patch.object(q3_closeout,'DIAG',root/'diagnostics'):
                try:
                    q3_closeout.publish_results(None, None, missing, None)
                    raise AssertionError('missing check reached production publisher')
                except RuntimeError:
                    pass
        empty=copy.deepcopy(v);empty['delivery_gate']['checks']={}
        with patch.object(problem3,'RESULTS_DIR',root),patch.object(q3_closeout,'export_workbook') as writer:
            try:
                problem3.write_result3_xlsx(out,np.zeros((1,21)),np.array([60.]),validation=empty)
                raise AssertionError('empty gate passed')
            except RuntimeError:
                pass
            assert not writer.called
        with patch.object(problem3,'RESULTS_DIR',root),patch.object(q3_closeout,'export_workbook') as writer:
            problem3.write_result3_xlsx(out,np.zeros((1,21)),np.array([60.]),validation=v)
            assert writer.call_count==1
    strict=problem3.find_strict_report(lambda t:.15-1e-6*(t-100.2),100.2,300.,margin=1e-5)
    assert strict['t_rep_s']==int(strict['t_rep_s'])
    assert strict['M_hour4']<.15-1e-5
    assert abs(strict['t_hour4_h']*10000-round(strict['t_hour4_h']*10000))<1e-8
    for evaluator, hi in [(lambda t:float('nan'),300.), (lambda t:.15-1e-6*(t-100.2),100.21)]:
        try:
            problem3.find_strict_report(evaluator,100.2,hi)
            raise AssertionError('invalid/out-of-range strict report passed')
        except RuntimeError:
            pass
    parts=v['continuous_balances']['parts']
    combine(parts)
    for field,delta in [('start_s',.001),('M_start',1e-5),('E_start',1.)]:
        bad=copy.deepcopy(parts);bad[1][field]+=delta
        try:
            combine(bad)
            raise AssertionError('discontinuous balance segments passed')
        except ValueError:
            pass
    # Independent weighted RHS identity checks both signs and the B-prime term.
    from appendix3 import heat_capacity,heat_capacity_deriv
    from radial_coupled import CoupledRadialFVM
    from utils import R,L,H,HM
    fvm=CoupledRadialFVM(np.linspace(0,R,41));n=fvm.n_nodes
    T=35+8*(fvm.r/R)**2;C=1.2-.5*(fvm.r/R)**2
    y=np.r_[T,C];dy=fvm.rhs(0,y,50.,.05);factor=2*np.pi*L
    storage_rate=factor*np.dot(fvm.W,heat_capacity(C)*dy[:n]+heat_capacity_deriv(C)*(T-28)*dy[n:])
    correction=factor*np.dot(fvm.W,heat_capacity_deriv(C)*(T-28)*dy[n:])
    def energy(z):
        return factor*np.dot(fvm.W,heat_capacity(z[n:])*(z[:n]-28))
    finite_difference=(energy(y+1e-3*dy)-energy(y-1e-3*dy))/(2e-3)
    assert np.isclose(finite_difference,storage_rate,rtol=1e-7)
    assert np.isclose(storage_rate-correction,factor*R*H*(50-T[-1]),rtol=1e-12)
    assert np.isclose(fvm.W@dy[n:],-R*HM*(C[-1]-.05),rtol=1e-12)
    source=np.load(RESULTS_DIR/'q3_solution.npz')
    wb=load_workbook(RESULTS_DIR/'result3.xlsx',read_only=True,data_only=True)
    assert wb.sheetnames==['水分浓度']
    rows=list(wb.active.iter_rows())
    assert len(rows)==len(source['excel_times'])+1
    assert all(len(row)==22 for row in rows)
    assert np.allclose([c.value for c in rows[0][1:]],np.arange(21)*.1,atol=1e-14,rtol=0)
    arr=np.array([[c.value for c in row] for row in rows[1:]],dtype=float)
    delta=float(np.max(abs(arr[:,1:]-source['C_out'])))
    assert delta<1e-14
    assert np.allclose(arr[:,0],source['excel_times'],atol=1e-8,rtol=0)
    assert all(c.number_format=='0.0000' for row in rows[1:] for c in row[1:])
    wb.close()
    assert np.array_equal(source['excel_times'][:-1],np.arange(60.,source['t_rep'],60.))
    assert float(source['t_rep'])==v['t_rep_s']
    n=int(source['n_nodes'])
    assert source['Y'].shape==(len(source['times']),2*n)
    assert np.max(source['y_end'][n:])+v['moisture_margin']<.15
    indices=np.array([np.argmin(abs(source['r']-r)) for r in problem3.OUTPUT_RADII_M])
    times,iy,ix=np.intersect1d(source['times'],source['excel_times'],return_indices=True)
    assert len(times)==len(source['excel_times'])-1
    assert np.array_equal(source['Y'][iy,n:][:,indices],source['C_out'][ix])
    assert np.array_equal(source['Y'][iy,:n][:,indices],source['T_out'][ix])
    assert np.array_equal(source['y_end'][n:][indices],source['C_out'][-1])
    assert np.array_equal(source['Y'][0,:n],np.full(n,28.))
    assert np.array_equal(source['Y'][0,n:],np.full(n,2.55))
    assert abs(source['table_hours'][-1]-v['t_rep_h'])<1e-12
    table=np.loadtxt(RESULTS_DIR/'q3_table_moisture.csv',delimiter=',',skiprows=1)
    assert np.max(abs(table[:,1:]-source['C_table']))<6e-11
    assert np.max(abs(table[:,0]-source['table_hours']))<6e-9
    arc=RESULTS_DIR/'archive/q3-pre-closeout-20260911'
    manifest=json.loads((arc/'manifest.json').read_text(encoding='utf-8-sig'))
    archive_checks={p['path']:hashlib.sha256((arc/p['path']).read_bytes()).hexdigest().lower()==p['sha256'].lower() for p in manifest}
    assert all(archive_checks.values()),archive_checks
    old=np.load(arc/'results/q3_solution.npz')
    # Compare common regular 6-hour rows; the two report times differ deliberately.
    table_changed=int(np.count_nonzero(np.round(old['C_table'][:-1],4)!=np.round(source['C_table'][:-1],4)))
    matched,io,inew=np.intersect1d(old['excel_times'],source['excel_times'],return_indices=True)
    common_delta=float(np.max(abs(old['C_out'][io]-source['C_out'][inew])))
    protected=json.loads((arc/'protected.json').read_text(encoding='utf-8-sig'))
    protected_checks={p['path']:hashlib.sha256((ROOT/p['path']).read_bytes()).hexdigest().lower()==p['sha256'].lower() for p in protected}
    assert all(protected_checks.values()),protected_checks
    source_checks={p:hashlib.sha256((ROOT/p).read_bytes()).hexdigest()==sha for p,sha in v['source_sha256'].items()}
    assert all(source_checks.values()),source_checks
    js=json.loads((q3_closeout.DIAG/'js_gate_test.json').read_text(encoding='utf-8'))
    assert js['pass'] and set(js['required_checks'])==q3_closeout.REQUIRED_CHECKS
    baseline={}
    for q in (1,2,3,4):
        data=json.loads((RESULTS_DIR/f'q{q}_validation.json').read_text(encoding='utf-8'))
        gate=data.get('delivery_gate',{})
        baseline[f'q{q}']=bool(gate.get('meets_delivery_gate',data.get('pass',False)))
    assert all(baseline.values()),baseline
    report_rows=[]
    for line in (ROOT/'reports/Q3_VERIFY_REPORT.md').read_text(encoding='utf-8').splitlines():
        cells=[c.strip() for c in line.strip('|').split('|')]
        if not line.startswith('|') or len(cells)!=6:
            continue
        try:
            report_rows.append([float(c) for c in cells])
        except ValueError:
            continue
    report_table=np.asarray(report_rows)
    assert report_table.shape==(len(source['table_hours']),6)
    assert np.allclose(report_table[:,0],source['table_hours'],rtol=0,atol=1e-10)
    assert np.array_equal(report_table[:,1:],np.round(source['C_table'],4))
    provenance=json.loads((RESULTS_DIR/'diagnostics/q3_delivery/figure_provenance.json').read_text(encoding='utf-8'))
    assert all(hashlib.sha256((ROOT/p).read_bytes()).hexdigest()==sha for p,sha in provenance['sha256'].items())
    visual=json.loads((RESULTS_DIR/'diagnostics/q3_delivery/visual_review.json').read_text(encoding='utf-8'))
    assert visual['pass']
    figure_hashes={str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest()
                   for p in sorted((ROOT/'figures').glob('q3_*.pdf'))}
    assert len(figure_hashes)==5
    prior=json.loads((RESULTS_DIR/'verification/closeout/files.json').read_text(encoding='utf-8'))
    prior_hashes={path:hashlib.sha256((ROOT/path).read_bytes()).hexdigest()==sha
                  for q in ('q1','q2') for path,sha in prior[q]['source_sha256'].items()}
    assert all(prior_hashes.values()),prior_hashes
    audit={'pass':True,'gate_failures_blocked':checked,'missing_required_checks_blocked':checked,
           'publisher_failure_blocked':True,'js_gate_test':js,'balance_continuity_regression':True,
           'weighted_rhs_balance_identity':True,'source_hashes_match':source_checks,
           'internal_fields_and_initial_state_match':True,'empty_gate_blocked':True,'valid_gate_reaches_writer':True,
           'strict_report_rounding_regression':True,'workbook_rows_including_header':len(rows),'columns':22,
           'export_max_abs_difference':delta,'regular_table_four_decimal_changes_from_old':table_changed,
           'common_excel_max_C_change_from_old':common_delta,'old_critical_h':float(old['t_star'])/3600,
           'new_critical_h':v['t_star_h'],'old_report_h':float(old['t_rep'])/3600,'new_report_h':v['t_rep_h'],
           'protected_files_unchanged':protected_checks,'archive_hashes_match':archive_checks,
           'q12_prior_evidence_hashes_match':prior_hashes,'four_question_numerical_gates':baseline,
           'report_table_matches_source':True,'figure_provenance_matches':True,
           'figure_sha256':figure_hashes,'visual_review':visual,
           'result3_sha256':hashlib.sha256((RESULTS_DIR/'result3.xlsx').read_bytes()).hexdigest()}
    (RESULTS_DIR/'q3_closeout_delivery_validation.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(audit,ensure_ascii=False),flush=True)


if __name__=='__main__':main()
