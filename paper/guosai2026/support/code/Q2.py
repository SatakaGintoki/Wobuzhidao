# -*- coding: utf-8 -*-
"""问题 2：从原始附件复算答案数据。

运行 python Q2.py；检查环境 python Q2.py --check。
本文件内嵌本问所需的全部项目模块，模块源码在下方按文件名分段列出。
运行时自动展开到相对目录 _runtime/Q2/code，以支持源码哈希核验。
输入、输出均相对于本文件的位置，与当前工作目录无关。
方程、物性、网格、容差和验证判据沿用原正式求解程序。
"""
import argparse
import importlib
import os
from pathlib import Path
import sys

PACKAGE_ROOT = Path(__file__).resolve().parent
QUESTION = 2
SOURCES = {}

# ========================================================================
# Module: utils.py
# ========================================================================
SOURCES['utils'] = r'''"""Shared paths and attachment readers for Problem A."""
from __future__ import annotations

from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.interpolate import PchipInterpolator

ROOT = Path(__file__).resolve().parents[3]
DATA_DIR = ROOT / "data"
ATTACH_DIR = DATA_DIR / "附件"
RESULTS_DIR = ROOT / "results"
FIGURES_DIR = ROOT / "figures"
REPORTS_DIR = ROOT / "reports"

# Appendix 2 (Problem 1)
R = 0.02
L = 0.25
RHO = 820.0
CP = 2600.0
K = 0.36
H = 25.0
HM = 8.0e-7
T0 = 28.0
C0 = 2.55
D_PREFACTOR = 7.0e-9
D_EXP_COEF = 0.89
T_END = 1800.0
ALPHA = K / (RHO * CP)


def moisture_diffusivity(C: np.ndarray | float, *, protect: bool = True) -> np.ndarray | float:
    """D(C) = 7e-9 exp(-0.89/C). Nonpositive C uses D=0 only as a solver-domain guard."""
    C_arr = np.asarray(C, dtype=float)
    out = np.zeros_like(C_arr, dtype=float)
    positive = C_arr > 0.0
    if np.any(positive):
        out[positive] = D_PREFACTOR * np.exp(-D_EXP_COEF / C_arr[positive])
    if not protect and np.any(~positive):
        raise ValueError("moisture diffusivity evaluated at nonpositive C without protection")
    if np.isscalar(C):
        return float(out)
    return out


def load_oven_table() -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    path = ATTACH_DIR / "附件1.xlsx"
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = list(wb.active.values)
    wb.close()
    data = np.asarray(rows[1:], dtype=float)
    t, Ta, Ca = data[:, 0], data[:, 1], data[:, 2]
    if t[0] != 0.0:
        raise ValueError(f"oven table must start at t=0, got {t[0]}")
    return t, Ta, Ca


class PiecewiseLinear:
    """Piecewise linear interpolant; kept for comparison, not used in Q1 production."""

    def __init__(self, t: np.ndarray, y: np.ndarray):
        self.t = np.asarray(t, dtype=float)
        self.y = np.asarray(y, dtype=float)
        self.kind = "linear"
        if self.t.ndim != 1 or self.y.shape != self.t.shape:
            raise ValueError("interpolant requires matching 1-D arrays")
        if np.any(np.diff(self.t) <= 0):
            raise ValueError("breakpoints must be strictly increasing")

    def __call__(self, t: float | np.ndarray) -> float | np.ndarray:
        return np.interp(t, self.t, self.y)

    def slope(self, t: float) -> float:
        t = float(t)
        if t <= self.t[0] or t >= self.t[-1]:
            if t >= self.t[-1]:
                dt = self.t[-1] - self.t[-2]
                return (self.y[-1] - self.y[-2]) / dt
            return (self.y[1] - self.y[0]) / (self.t[1] - self.t[0])
        i = int(np.searchsorted(self.t, t, side="right") - 1)
        i = min(i, len(self.t) - 2)
        return (self.y[i + 1] - self.y[i]) / (self.t[i + 1] - self.t[i])


class PchipOven:
    """Monotone piecewise cubic Hermite interpolant (PCHIP) through attachment knots."""

    def __init__(self, t: np.ndarray, y: np.ndarray):
        self.t = np.asarray(t, dtype=float)
        self.y = np.asarray(y, dtype=float)
        self.kind = "pchip"
        if self.t.ndim != 1 or self.y.shape != self.t.shape:
            raise ValueError("interpolant requires matching 1-D arrays")
        if np.any(np.diff(self.t) <= 0):
            raise ValueError("breakpoints must be strictly increasing")
        self._pchip = PchipInterpolator(self.t, self.y, extrapolate=False)
        self._deriv = self._pchip.derivative()

    def __call__(self, t: float | np.ndarray) -> float | np.ndarray:
        tt = np.clip(np.asarray(t, dtype=float), self.t[0], self.t[-1])
        val = self._pchip(tt)
        if np.ndim(t) == 0:
            return float(val)
        return np.asarray(val, dtype=float)

    def slope(self, t: float) -> float:
        tt = float(np.clip(t, self.t[0], self.t[-1]))
        return float(self._deriv(tt))
'''

# ========================================================================
# Module: delivery.py
# ========================================================================
SOURCES['delivery'] = r'''"""Fail-closed numerical delivery checks shared by Q1 and Q2."""
import json
from pathlib import Path


def make_gate(checks):
    checks = {name: bool(value) for name, value in checks.items()}
    return {"checks": checks, "failed_checks": [k for k, v in checks.items() if not v],
            "meets_delivery_gate": bool(checks) and all(checks.values())}


def require_delivery(validation, diagnostic_path):
    """Persist diagnostics first; never reach any production writer on failure."""
    path = Path(diagnostic_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(validation, ensure_ascii=False, indent=2, default=float), encoding="utf-8")
    gate = validation["delivery_gate"]
    checks = gate.get("checks", {})
    if not checks or not all(value is True for value in checks.values()) or gate.get("meets_delivery_gate") is not True:
        raise RuntimeError("Delivery blocked; see " + str(path) + ": " + str(gate.get("failed_checks", [])))
'''

# ========================================================================
# Module: q1_analytic.py
# ========================================================================
SOURCES['q1_analytic'] = r'''"""Independent Bessel–Duhamel reference for the linear radial heat (or frozen-D moisture) problem."""
from __future__ import annotations

from typing import Any

import numpy as np
from scipy.integrate import solve_ivp
from scipy.optimize import brentq
from scipy.special import j0, j1
from scipy.sparse import diags


def robin_eigenvalues(Bi: float, n_roots: int) -> np.ndarray:
    """Positive roots of lambda * J1(lambda) = Bi * J0(lambda)."""
    if Bi <= 0:
        raise ValueError("Biot number must be positive")
    n_roots = int(n_roots)

    def f(lam: float) -> float:
        return lam * j1(lam) - Bi * j0(lam)

    lam_max = (n_roots + 4) * np.pi
    grid = np.linspace(1e-10, lam_max, max(8000, 200 * n_roots))
    vals = grid * j1(grid) - Bi * j0(grid)
    roots: list[float] = []
    for i in range(len(grid) - 1):
        v0, v1 = vals[i], vals[i + 1]
        if not np.isfinite(v0) or not np.isfinite(v1):
            continue
        a, b = float(grid[i]), float(grid[i + 1])
        if v0 == 0.0:
            cand = a
        elif v0 * v1 < 0.0:
            cand = float(brentq(f, a, b, xtol=1e-14, maxiter=200))
        else:
            continue
        if not roots or abs(cand - roots[-1]) > 1e-10:
            roots.append(cand)
        if len(roots) >= n_roots:
            break
    if len(roots) < n_roots:
        raise RuntimeError(f"found only {len(roots)} Robin Bessel roots, need {n_roots}")
    return np.asarray(roots, dtype=float)


def expansion_coefficients(lam: np.ndarray) -> np.ndarray:
    j0l = j0(lam)
    j1l = j1(lam)
    return 2.0 * j1l / (lam * (j0l**2 + j1l**2))


class BesselDuhamel:
    def __init__(self, Bi: float, diffusivity: float, R: float, n_terms: int = 120):
        self.Bi = float(Bi)
        self.alpha = float(diffusivity)
        self.R = float(R)
        self.lam = robin_eigenvalues(self.Bi, n_terms)
        self.A = expansion_coefficients(self.lam)
        self.beta = self.alpha * self.lam**2 / self.R**2

    def radial_modes(self, r: np.ndarray) -> np.ndarray:
        """Return array shape (n_terms, n_r) of A_n J0(lambda_n r/R)."""
        r = np.asarray(r, dtype=float)
        return self.A[:, None] * j0(self.lam[:, None] * (r[None, :] / self.R))

    def evaluate(self, r: np.ndarray, t: float, env: Any, u0: float) -> np.ndarray:
        return self.evaluate_many(r, np.array([float(t)]), env, u0)[0]

    def evaluate_many(self, r: np.ndarray, times: np.ndarray, env: Any, u0: float) -> np.ndarray:
        """T = env(t) - sum A_n J0 I_n, with I' = -beta I + env'(t), I(0)=env(0)-u0.

        Valid for any C1 interpolant that exposes env.y, env.slope, and env(t).
        """
        r = np.asarray(r, dtype=float)
        times = np.asarray(times, dtype=float)
        if times.size == 0:
            return np.empty((0, r.size), dtype=float)
        modes = self.radial_modes(r)
        beta = self.beta
        I0 = np.full(beta.size, float(env.y[0] - u0))
        t_end = float(np.max(times))
        out = np.empty((times.size, r.size), dtype=float)
        zero = times == 0.0
        if np.all(zero):
            out[:] = u0
            return out
        t_eval = times[~zero]
        sol = solve_ivp(
            lambda t, I: -beta * I + env.slope(t),
            (0.0, max(t_end, 1e-30)),
            I0,
            method="BDF",
            jac=diags(-beta, format="csc"),
            rtol=1e-11,
            atol=1e-13,
            t_eval=t_eval,
            max_step=1.0,
            vectorized=False,
        )
        if not sol.success:
            raise RuntimeError(f"Bessel Duhamel ODE failed: {sol.message}")
        Ihist = sol.y.T
        env_t = np.array([float(env(t)) for t in t_eval], dtype=float)
        nonzero_vals = env_t[:, None] - Ihist @ modes
        out[zero] = u0
        out[~zero] = nonzero_vals
        return out


def converged_reference(Bi, diffusivity, R, r, times, env, u0, start_terms=320, tolerance=1e-8):
    """Double the modal truncation and explicitly check every requested point."""
    previous = BesselDuhamel(Bi, diffusivity, R, start_terms).evaluate_many(r, times, env, u0)
    history = []
    for terms in (start_terms * 2, start_terms * 4, start_terms * 8):
        current = BesselDuhamel(Bi, diffusivity, R, terms).evaluate_many(r, times, env, u0)
        error = float(np.max(np.abs(current - previous)))
        history.append({"terms": terms, "previous_terms": terms // 2, "max_abs_difference": error})
        if np.isfinite(error) and error <= tolerance:
            return current, {"passed": True, "tolerance": tolerance, "history": history, "terms": terms}
        previous = current
    return current, {"passed": False, "tolerance": tolerance, "history": history, "terms": terms}
'''

# ========================================================================
# Module: radial_fvm.py
# ========================================================================
SOURCES['radial_fvm'] = r'''"""Node-centered radial finite-volume operators for Problem 1."""
from __future__ import annotations

import numpy as np
from scipy.sparse import diags

from utils import CP, H, HM, K, R, RHO, moisture_diffusivity

OUTPUT_RADII_M = 0.001 * np.arange(21)


def graded_radial_nodes(level: int) -> np.ndarray:
    """Nodes denser near the surface; every 0.1 cm output radius remains a node.

    Interval subdivisions (before multiplying by ``level``):
    0–1.5 cm: 8; 1.5–1.8 cm: 16; 1.8–1.9 cm: 32; 1.9–2.0 cm: 64.
    Surface spacing at level L is 1e-3 / (64 L) metres.
    """
    level = int(level)
    if level < 1:
        raise ValueError("graded mesh level must be a positive integer")
    n_sub = np.full(20, 8 * level, dtype=int)
    n_sub[15:] = 16 * level
    n_sub[18] = 32 * level
    n_sub[19] = 64 * level
    pts = [0.0]
    for j in range(20):
        r0 = 0.001 * j
        r1 = 0.001 * (j + 1)
        m = int(n_sub[j])
        for k in range(1, m + 1):
            pts.append(r0 + (r1 - r0) * k / m)
    r = np.asarray(pts, dtype=float)
    r[0] = 0.0
    r[-1] = R
    return r


class RadialFVM:
    def __init__(self, N: int | None = None, r: np.ndarray | None = None):
        if r is not None:
            self.r = np.asarray(r, dtype=float)
            if self.r.ndim != 1 or self.r.size < 3:
                raise ValueError("r must be a 1-D array with at least 3 nodes")
            if abs(self.r[0]) > 1e-15 or abs(self.r[-1] - R) > 1e-14:
                raise ValueError("radial nodes must start at 0 and end at R")
            if np.any(np.diff(self.r) <= 0.0):
                raise ValueError("radial nodes must be strictly increasing")
        else:
            if N is None or int(N) < 2:
                raise ValueError("N must be at least 2")
            self.r = np.linspace(0.0, R, int(N) + 1)
        self.n_nodes = int(self.r.size)
        self.N = self.n_nodes - 1
        self.dr_face = np.diff(self.r)
        self.r_half = 0.5 * (self.r[:-1] + self.r[1:])
        a = np.empty(self.n_nodes)
        b = np.empty(self.n_nodes)
        a[0] = 0.0
        b[0] = self.r_half[0]
        a[-1] = self.r_half[-1]
        b[-1] = R
        if self.n_nodes > 2:
            a[1:-1] = self.r_half[:-1]
            b[1:-1] = self.r_half[1:]
        self.W = 0.5 * (b**2 - a**2)
        self.dr = float(self.dr_face.min())
        if abs(self.W.sum() - 0.5 * R**2) > 1e-14:
            raise RuntimeError("radial weights do not sum to R^2/2")

        self.M_heat = self._heat_matrix()
        self.M_heat_sparse = diags(
            [np.diag(self.M_heat, -1), np.diag(self.M_heat, 0), np.diag(self.M_heat, 1)],
            offsets=[-1, 0, 1],
            shape=self.M_heat.shape,
            format="csc",
        )
        self._heat_force_coef = (R * H) / (RHO * CP * self.W[-1])
        self._heat_force_index = self.n_nodes - 1

    def output_indices(self) -> np.ndarray:
        """Indices of 0, 0.1, ..., 2.0 cm nodes."""
        idx = np.empty(OUTPUT_RADII_M.size, dtype=int)
        for j, rc in enumerate(OUTPUT_RADII_M):
            k = int(np.argmin(np.abs(self.r - rc)))
            if abs(self.r[k] - rc) > 1e-12:
                raise ValueError(f"grid misses output radius {rc}, nearest {self.r[k]}")
            idx[j] = k
        return idx

    def _interior_face_coeff(self) -> np.ndarray:
        return self.r_half / self.dr_face

    def _heat_matrix(self) -> np.ndarray:
        n = self.n_nodes
        M = np.zeros((n, n))
        face = K * self._interior_face_coeff()
        den = RHO * CP * self.W
        M[0, 0] = -face[0] / den[0]
        M[0, 1] = face[0] / den[0]
        for i in range(1, n - 1):
            M[i, i - 1] = face[i - 1] / den[i]
            M[i, i + 1] = face[i] / den[i]
            M[i, i] = -(face[i - 1] + face[i]) / den[i]
        i = n - 1
        M[i, i - 1] = face[i - 1] / den[i]
        M[i, i] = (-face[i - 1] - R * H) / den[i]
        return M

    def heat_rhs(self, t: float, T: np.ndarray, Ta: float) -> np.ndarray:
        out = self.M_heat @ T
        out[self._heat_force_index] += self._heat_force_coef * Ta
        return out

    def heat_jac(self, t: float, T: np.ndarray):
        return self.M_heat_sparse

    def heat_jac_sparsity(self):
        n = self.n_nodes
        ones_m = np.ones(n - 1, dtype=float)
        ones_d = np.ones(n, dtype=float)
        return diags([ones_m, ones_d, ones_m], offsets=[-1, 0, 1], shape=(n, n), format="csc")

    def moisture_rhs(self, t: float, C: np.ndarray, Ca: float, D_const: float | None = None) -> np.ndarray:
        if D_const is None:
            D_node = moisture_diffusivity(C, protect=True)
        else:
            D_node = np.full_like(C, float(D_const))
        D_face = harmonic_mean(D_node[:-1], D_node[1:])
        F = -self.r_half * D_face * np.diff(C) / self.dr_face
        Cdot = np.empty_like(C)
        Cdot[0] = -F[0] / self.W[0]
        Cdot[1:-1] = (F[:-1] - F[1:]) / self.W[1:-1]
        F_surf = R * HM * (C[-1] - Ca)
        Cdot[-1] = (F[-1] - F_surf) / self.W[-1]
        return Cdot

    def moisture_jac_sparsity(self):
        return self.heat_jac_sparsity()

    def surface_heat_influx_density(self, T_s: float, Ta: float) -> float:
        """Heat into the solid per unit surface area, W/m^2."""
        return H * (Ta - T_s)

    def surface_moisture_outflux(self, C_s: float, Ca: float) -> float:
        """Outward moisture flux in the normalized (no rho_d) units, m/s * kg/kg."""
        return HM * (C_s - Ca)


def harmonic_mean(a: np.ndarray, b: np.ndarray) -> np.ndarray:
    out = np.zeros_like(a, dtype=float)
    both = (a > 0.0) & (b > 0.0)
    out[both] = 2.0 * a[both] * b[both] / (a[both] + b[both])
    return out
'''

# ========================================================================
# Module: reproduction_export.py
# ========================================================================
SOURCES['reproduction_export'] = r'''"""Portable Excel export from validated solver arrays; no external runtime."""
from copy import copy
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


def write_template(root, question, times, fields):
    root = Path(root)
    template = root / 'data' / '附件' / '附件3' / f'result{question}.xlsx'
    wb = load_workbook(template)
    if len(wb.worksheets) != len(fields):
        raise ValueError('Template sheet count does not match the output fields')
    times = np.asarray(times)
    if times.ndim != 1 or not np.isfinite(times).all() or not (np.diff(times) > 0).all():
        raise ValueError('Invalid output time sequence')
    for ws, values in zip(wb.worksheets, fields):
        values = np.asarray(values)
        expected = 22 if question == 4 else 21
        if values.shape != (len(times), expected):
            raise ValueError('Output shape does not match the requested radial grid')
        if np.isinf(values).any() or (question != 4 and np.isnan(values).any()):
            raise ValueError('Unexpected nonfinite output')
        header = ws.cell(1, 1).value
        surface = ws.cell(1, ws.max_column).value
        number_style = copy(ws.cell(2, 2)._style)
        header_style = copy(ws.cell(1, 2)._style)
        ws.delete_rows(1, ws.max_row)
        ws.append([header] + [j / 10 for j in range(21)] + ([surface] if question == 4 else []))
        for cell in ws[1]:
            cell._style = copy(header_style)
        for t, row in zip(times, values):
            ws.append([float(t)] + [None if np.isnan(v) else float(v) for v in row])
        for row in ws.iter_rows(min_row=2, min_col=2):
            for cell in row:
                cell._style = copy(number_style)
                cell.number_format = '0.0000'
        ws.freeze_panes = 'B2'
    wb.properties.creator = 'Team'
    wb.properties.lastModifiedBy = 'Team'
    dest = root / 'results' / f'result{question}.xlsx'
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    wb.close()
    print(f'Excel saved: results/result{question}.xlsx', flush=True)


def export_saved(root, question):
    root = Path(root)
    validation = json.loads((root / 'results' / f'q{question}_validation.json').read_text(encoding='utf-8'))
    if question == 4:
        if validation.get('pass') is not True or not validation.get('checks') or not all(v is True for v in validation['checks'].values()):
            raise RuntimeError('Q4 numerical validation did not pass')
    else:
        from delivery import require_delivery
        require_delivery(validation, root / 'results' / f'q{question}_export_validation.json')
        if question == 3:
            from q3_closeout import require_q3_delivery
            require_q3_delivery(validation, root / 'results' / 'q3_export_validation.json')
    with np.load(root / 'results' / f'q{question}_solution.npz') as z:
        if question in (1, 2):
            write_template(root, question, z['times'], [z['T_out'], z['C_out']])
        elif question == 3:
            write_template(root, question, z['excel_times'], [z['C_out']])
        else:
            write_template(root, question, z['times_s'][1:], [z['C_physical'][1:]])
'''

# ========================================================================
# Module: appendix3.py
# ========================================================================
SOURCES['appendix3'] = r'''"""Appendix 3 effective properties for Problems 2 and 3."""
from __future__ import annotations

import numpy as np

RHO_A = 650.0
RHO_B = 128.0
CP_A = 1450.0
CP_B = 2736.0
K_A = 0.21
K_B = 0.38
D_PREFACTOR = 2.4e-3
D_C_COEF = 0.45
D_T_COEF = 3850.0


def rho(C: np.ndarray | float) -> np.ndarray | float:
    return RHO_A + RHO_B * np.asarray(C, dtype=float)


def cp(C: np.ndarray | float) -> np.ndarray | float:
    C_arr = np.asarray(C, dtype=float)
    return CP_A + CP_B * C_arr / (1.0 + C_arr)


def conductivity(C: np.ndarray | float) -> np.ndarray | float:
    C_arr = np.asarray(C, dtype=float)
    return K_A + K_B * C_arr / (1.0 + C_arr)


def heat_capacity(C: np.ndarray | float) -> np.ndarray | float:
    return rho(C) * cp(C)


def conductivity_deriv(C: np.ndarray | float) -> np.ndarray | float:
    C_arr = np.asarray(C, dtype=float)
    return K_B / (1.0 + C_arr) ** 2


def heat_capacity_deriv(C: np.ndarray | float) -> np.ndarray | float:
    C_arr = np.asarray(C, dtype=float)
    return RHO_B * cp(C_arr) + rho(C_arr) * CP_B / (1.0 + C_arr) ** 2


def moisture_diffusivity_q2(
    C: np.ndarray | float,
    T_celsius: np.ndarray | float,
    *,
    protect: bool = True,
) -> np.ndarray | float:
    """D(C,T) with T in Celsius; Kelvin is used only inside the Arrhenius factor."""
    C_arr = np.asarray(C, dtype=float)
    T_arr = np.asarray(T_celsius, dtype=float)
    out = np.zeros(np.broadcast(C_arr, T_arr).shape, dtype=float)
    C_b, T_b = np.broadcast_arrays(C_arr, T_arr)
    positive = C_b > 0.0
    if np.any(positive):
        Tk = T_b[positive] + 273.15
        out[positive] = D_PREFACTOR * np.exp(-D_C_COEF / C_b[positive]) * np.exp(-D_T_COEF / Tk)
    if not protect and np.any(~positive):
        raise ValueError("q2 diffusivity evaluated at nonpositive C without protection")
    if np.ndim(C) == 0 and np.ndim(T_celsius) == 0:
        return float(np.asarray(out))
    return out


def diffusivity_partials(C: np.ndarray, T_celsius: np.ndarray) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    D = np.asarray(moisture_diffusivity_q2(C, T_celsius, protect=True), dtype=float)
    C_arr = np.asarray(C, dtype=float)
    T_arr = np.asarray(T_celsius, dtype=float)
    dD_dC = np.zeros_like(D)
    dD_dT = np.zeros_like(D)
    ok = C_arr > 0.0
    dD_dC[ok] = D[ok] * (D_C_COEF / C_arr[ok] ** 2)
    Tk = T_arr + 273.15
    dD_dT[ok] = D[ok] * (D_T_COEF / Tk[ok] ** 2)
    return D, dD_dC, dD_dT
'''

# ========================================================================
# Module: radial_coupled.py
# ========================================================================
SOURCES['radial_coupled'] = r'''"""Coupled variable-property radial FVM for Problem 2."""
from __future__ import annotations

import numpy as np
from scipy.sparse import lil_matrix

from appendix3 import conductivity, heat_capacity, moisture_diffusivity_q2
from radial_fvm import RadialFVM, harmonic_mean
from utils import H, HM, R


class CoupledRadialFVM:
    def __init__(self, r: np.ndarray):
        geom = RadialFVM(r=r)
        self.r = geom.r
        self.W = geom.W
        self.r_half = geom.r_half
        self.dr_face = geom.dr_face
        self.n_nodes = geom.n_nodes
        self.N = geom.N
        self._geom = geom

    def output_indices(self) -> np.ndarray:
        return self._geom.output_indices()

    def pack(self, T: np.ndarray, C: np.ndarray) -> np.ndarray:
        return np.concatenate([T, C])

    def unpack(self, y: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
        n = self.n_nodes
        return y[:n], y[n:]

    def rhs(
        self,
        t: float,
        y: np.ndarray,
        Ta: float,
        Ca: float,
        *,
        D_const: float | None = None,
        ST: np.ndarray | None = None,
        SC: np.ndarray | None = None,
    ) -> np.ndarray:
        T, C = self.unpack(y)
        kn = np.asarray(conductivity(C), dtype=float)
        Bn = np.asarray(heat_capacity(C), dtype=float)
        if D_const is None:
            Dn = np.asarray(moisture_diffusivity_q2(C, T, protect=True), dtype=float)
        else:
            Dn = np.full_like(C, float(D_const))
        k_face = harmonic_mean(kn[:-1], kn[1:])
        D_face = harmonic_mean(Dn[:-1], Dn[1:])
        FT = -self.r_half * k_face * np.diff(T) / self.dr_face
        FC = -self.r_half * D_face * np.diff(C) / self.dr_face
        Tdot = np.empty_like(T)
        Cdot = np.empty_like(C)
        Tdot[0] = -FT[0] / (Bn[0] * self.W[0])
        Cdot[0] = -FC[0] / self.W[0]
        Tdot[1:-1] = (FT[:-1] - FT[1:]) / (Bn[1:-1] * self.W[1:-1])
        Cdot[1:-1] = (FC[:-1] - FC[1:]) / self.W[1:-1]
        FTsurf = R * H * (T[-1] - Ta)
        FCsurf = R * HM * (C[-1] - Ca)
        Tdot[-1] = (FT[-1] - FTsurf) / (Bn[-1] * self.W[-1])
        Cdot[-1] = (FC[-1] - FCsurf) / self.W[-1]
        if ST is not None:
            Tdot = Tdot + ST / Bn
        if SC is not None:
            Cdot = Cdot + SC
        return self.pack(Tdot, Cdot)

    def jac_sparsity(self):
        n = self.n_nodes
        J = lil_matrix((2 * n, 2 * n), dtype=float)
        for i in range(n):
            js = [i]
            if i > 0:
                js.append(i - 1)
            if i + 1 < n:
                js.append(i + 1)
            for j in js:
                J[i, j] = 1.0
                J[i, n + j] = 1.0
                J[n + i, j] = 1.0
                J[n + i, n + j] = 1.0
        return J.tocsc()
'''

# ========================================================================
# Module: problem2.py
# ========================================================================
SOURCES['problem2'] = r'''"""Problem 2: coupled variable-property radial drying for 0–3 h."""
from __future__ import annotations

import json
import sys
import time
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from scipy.integrate import solve_ivp

CODE_DIR = Path(__file__).resolve().parent
if str(CODE_DIR) not in sys.path:
    sys.path.insert(0, str(CODE_DIR))

from appendix3 import (
    conductivity,
    conductivity_deriv,
    diffusivity_partials,
    heat_capacity,
    heat_capacity_deriv,
    moisture_diffusivity_q2,
)
from q1_analytic import BesselDuhamel, converged_reference
from delivery import make_gate, require_delivery
from radial_coupled import CoupledRadialFVM
from radial_fvm import graded_radial_nodes
from utils import (
    C0,
    FIGURES_DIR,
    H,
    HM,
    L,
    R,
    RESULTS_DIR,
    T0,
    PchipOven,
    load_oven_table,
)

T_END = 10800.0
OUTPUT_RADII_M = 0.001 * np.arange(21)
TABLE_HOURS = np.array([0.5, 1.0, 1.5, 2.0, 2.5, 3.0])
TABLE_TIMES = 3600.0 * TABLE_HOURS
TABLE_RADII_CM = np.array([0.0, 0.5, 1.0, 1.5, 2.0])
TREF = 28.0


def load_q2_inputs() -> tuple[PchipOven, PchipOven, np.ndarray]:
    t, Ta, Ca = load_oven_table()
    mask = t <= T_END + 1e-12
    t, Ta, Ca = t[mask], Ta[mask], Ca[mask]
    if t.size != 181 or abs(t[-1] - T_END) > 1e-12:
        raise ValueError(f"expected 181 oven records on [0, 10800], got {t.size}, t_end={t[-1]}")
    if abs(Ta[0] - T0) > 1e-12:
        raise ValueError(f"Ta(0)={Ta[0]} != T0={T0}")
    Ta_fun = PchipOven(t, Ta)
    Ca_fun = PchipOven(t, Ca)
    for tj, yj in zip(t, Ta):
        if abs(float(Ta_fun(tj)) - float(yj)) > 1e-12:
            raise RuntimeError(f"PCHIP misses Ta knot t={tj}")
    for tj, yj in zip(t, Ca):
        if abs(float(Ca_fun(tj)) - float(yj)) > 1e-12:
            raise RuntimeError(f"PCHIP misses Ca knot t={tj}")
    return Ta_fun, Ca_fun, t


def integrate_coupled(fvm: CoupledRadialFVM, y0, t_end, Ta_fun, Ca_fun, rtol, atol, max_step,
                      D_const=None, source=None, t_eval=None):
    y0 = np.asarray(y0, dtype=float)
    n = fvm.n_nodes
    atol_arr = np.asarray(atol, dtype=float)
    if atol_arr.ndim == 0:
        atol_arr = np.concatenate([np.full(n, float(atol_arr)), np.full(n, float(atol_arr) * 0.01)])

    def rhs(t, y):
        ST = SC = None
        if source is not None:
            ST, SC = source(t, fvm.r)
        return fvm.rhs(t, y, float(Ta_fun(t)), float(Ca_fun(t)), D_const=D_const, ST=ST, SC=SC)

    if t_eval is None:
        t_eval = np.arange(1.0, t_end + 0.5, 1.0)
    sol = solve_ivp(
        rhs,
        (0.0, float(t_end)),
        y0,
        method="BDF",
        rtol=rtol,
        atol=atol_arr,
        max_step=max_step,
        t_eval=t_eval,
        dense_output=False,
        vectorized=False,
        jac_sparsity=fvm.jac_sparsity(),
    )
    if not sol.success:
        raise RuntimeError(f"coupled BDF failed: {sol.message}")
    times = np.concatenate([[0.0], sol.t])
    Y = np.vstack([y0, sol.y.T])
    return times, Y


def extract_output(fvm: CoupledRadialFVM, Y: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    n = fvm.n_nodes
    idx = fvm.output_indices()
    T = Y[:, :n][:, idx]
    C = Y[:, n:][:, idx]
    return T, C


def compare_on_output(A, B, times) -> dict:
    diff = np.abs(A - B)
    k = int(np.argmax(diff))
    n_t, n_r = diff.shape
    i, j = divmod(k, n_r)
    paper_t = np.isin(times, TABLE_TIMES)
    paper_r = slice(None, None, 5) if n_r == 21 else slice(None)
    return {
        "max_abs": float(diff.max()),
        "max_time_s": float(times[i]),
        "max_radius_cm": 0.1 * j,
        "mean_abs": float(diff.mean()),
        "paper_table_max_abs": float(np.max(np.abs(A[paper_t][:, paper_r] - B[paper_t][:, paper_r]))),
    }


def four_decimal_stats(A, B, times) -> dict:
    Ar, Br = np.round(A, 4), np.round(B, 4)
    changed = Ar != Br
    paper_t = np.isin(times, TABLE_TIMES)
    paper_r = slice(None, None, 5) if A.shape[1] == 21 else slice(None)
    paper_changed = Ar[paper_t][:, paper_r] != Br[paper_t][:, paper_r]
    return {
        "changed_cells": int(np.sum(changed)),
        "total_cells": int(A.size),
        "paper_changed_cells": int(np.sum(paper_changed)),
        "paper_total_cells": int(paper_changed.size),
        "max_abs": float(np.max(np.abs(A - B))),
        "paper_table_max_abs": float(np.max(np.abs(A[paper_t][:, paper_r] - B[paper_t][:, paper_r]))),
    }


def manufactured_fields(r: np.ndarray, t: float) -> tuple[np.ndarray, np.ndarray]:
    x2 = (r / R) ** 2
    T = 28.0 + 2.0 * np.exp(-t / 600.0) * (1.0 + x2)
    C = 1.5 + 0.2 * np.exp(-t / 500.0) * (1.0 + x2)
    return T, C


def manufactured_env(r_surface: float, t: float) -> tuple[float, float]:
    T, C = manufactured_fields(np.array([r_surface]), t)
    Ts, Cs = float(T[0]), float(C[0])
    u = np.exp(-t / 600.0)
    v = np.exp(-t / 500.0)
    Tr = 4.0 / R * u
    Cr = 4.0 / R * v * 0.1  # C amplitude 0.2 vs T amplitude 2: Cr = 0.2*v*2r/R^2 at r=R -> 0.4 v / R
    Cr = 0.4 / R * v
    kn = float(conductivity(Cs))
    Dn = float(moisture_diffusivity_q2(Cs, Ts))
    Ta = Ts + kn * Tr / H
    Ca = Cs + Dn * Cr / HM
    return Ta, Ca


class ManufacturedOven:
    def __init__(self, kind: str):
        self.kind = kind

    def __call__(self, t):
        Ta, Ca = manufactured_env(R, float(t))
        return Ta if self.kind == "T" else Ca


def manufactured_sources(t: float, r: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    T, C = manufactured_fields(r, t)
    x2 = (r / R) ** 2
    u = np.exp(-t / 600.0)
    v = np.exp(-t / 500.0)
    Tt = 2.0 * (-1.0 / 600.0) * u * (1.0 + x2)
    Ct = 0.2 * (-1.0 / 500.0) * v * (1.0 + x2)
    Tr = 4.0 * r / (R ** 2) * u
    Trr = np.full_like(r, 4.0 / (R ** 2) * u)
    Cr = 0.4 * r / (R ** 2) * v
    Crr = np.full_like(r, 0.4 / (R ** 2) * v)
    lapT = np.empty_like(r)
    lapC = np.empty_like(r)
    center = r < 1e-16
    lapT[center] = 2.0 * Trr[center]
    lapC[center] = 2.0 * Crr[center]
    away = ~center
    lapT[away] = Trr[away] + Tr[away] / r[away]
    lapC[away] = Crr[away] + Cr[away] / r[away]
    kn = np.asarray(conductivity(C), dtype=float)
    kp = np.asarray(conductivity_deriv(C), dtype=float)
    Bn = np.asarray(heat_capacity(C), dtype=float)
    D, dD_dC, dD_dT = diffusivity_partials(C, T)
    ST = Bn * Tt - (kn * lapT + kp * Cr * Tr)
    SC = Ct - (D * lapC + dD_dC * Cr ** 2 + dD_dT * Tr * Cr)
    return ST, SC


def conservation_moisture(fvm, times, Y, Ca_fun) -> dict:
    n = fvm.n_nodes
    C = Y[:, n:]
    M = C @ fvm.W
    M0 = float(M[0])
    CN = C[:, -1]
    Ca = np.array([float(Ca_fun(t)) for t in times])
    I = float(np.trapezoid(HM * (CN - Ca), times))
    rhs = M0 - R * I
    Mt = float(M[-1])
    denom = max(abs(M0 - Mt), abs(R * I), 1e-18)
    return {"M0": M0, "M_end": Mt, "rhs": rhs, "relative_residual": abs(Mt - rhs) / denom}


def conservation_heat(fvm, times, Y, Ta_fun) -> dict:
    n = fvm.n_nodes
    T = Y[:, :n]
    C = Y[:, n:]
    Bn = np.asarray(heat_capacity(C), dtype=float)
    Bp = np.asarray(heat_capacity_deriv(C), dtype=float)
    E = 2.0 * np.pi * L * np.sum(fvm.W[None, :] * Bn * (T - TREF), axis=1)
    TN = T[:, -1]
    Ta = np.array([float(Ta_fun(t)) for t in times])
    Q = 2.0 * np.pi * R * L * H * float(np.trapezoid(Ta - TN, times))
    dC = np.gradient(C, times, axis=0)
    corr = 2.0 * np.pi * L * np.sum(fvm.W[None, :] * Bp * (T - TREF) * dC, axis=1)
    Icorr = float(np.trapezoid(corr, times))
    lhs = float(E[-1] - E[0] - Icorr)
    denom = max(abs(Q), abs(lhs), 1e-12)
    return {
        "E_end_minus_E0": float(E[-1] - E[0]),
        "composition_correction": Icorr,
        "Q_surface": Q,
        "relative_residual": abs(lhs - Q) / denom,
    }


def assert_physical(T_out, C_out, Ta_fun, Ca_fun) -> dict:
    Ta_max = float(np.max(Ta_fun.y))
    Ca_min = float(np.min(Ca_fun.y))
    info = {
        "T_min": float(T_out.min()),
        "T_max": float(T_out.max()),
        "C_min": float(C_out.min()),
        "C_max": float(C_out.max()),
        "T_below_28": float(max(0.0, 28.0 - T_out.min())),
        "T_above_Ta_max": float(max(0.0, T_out.max() - Ta_max)),
        "C_above_C0": float(max(0.0, C_out.max() - C0)),
        "C_below_Ca_min": float(max(0.0, Ca_min - C_out.min())),
        "Ta_range": [float(np.min(Ta_fun.y)), Ta_max],
        "Ca_range": [Ca_min, float(np.max(Ca_fun.y))],
    }
    info["T_bounds_ok"] = info["T_below_28"] <= 1e-6 and info["T_above_Ta_max"] <= 1e-4
    info["C_bounds_ok"] = info["C_above_C0"] <= 1e-8 and info["C_below_Ca_min"] <= 1e-8
    return info


def write_result_workbook(path: Path, T_out, C_out, times):
    radii_cm = 0.1 * np.arange(21)
    wb = Workbook()
    for title, arr in (("温度", T_out), ("水分浓度", C_out)):
        ws = wb.create_sheet(title)
        ws["A1"] = "时间\\到药材中心的距离"
        for j, rc in enumerate(radii_cm, start=2):
            ws.cell(1, j, float(rc))
        for i, t in enumerate(times):
            ws.cell(i + 2, 1, int(round(float(t))))
            for j in range(21):
                cell = ws.cell(i + 2, j + 2, float(arr[i, j]))
                cell.number_format = "0.0000"
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_paper_csv(path: Path, arr_out, times):
    t_idx = {float(t): i for i, t in enumerate(times)}
    r_idx = {round(0.1 * j, 10): j for j in range(21)}
    lines = ["t_h," + ",".join(f"r_{rc:g}cm" for rc in TABLE_RADII_CM)]
    for th, t in zip(TABLE_HOURS, TABLE_TIMES):
        row = [f"{th:g}"]
        i = t_idx[float(t)]
        for rc in TABLE_RADII_CM:
            j = r_idx[float(rc)]
            row.append(f"{arr_out[i, j]:.8f}")
        lines.append(",".join(row))
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def setup_mpl():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({
        "font.sans-serif": ["Microsoft YaHei", "SimHei", "Noto Sans CJK SC", "DejaVu Sans"],
        "axes.unicode_minus": False,
        "font.size": 9,
        "axes.labelsize": 10,
        "legend.fontsize": 8,
        "xtick.labelsize": 8,
        "ytick.labelsize": 8,
        "axes.linewidth": 0.8,
        "pdf.fonttype": 42,
        "ps.fonttype": 42,
    })
    return plt


def plot_figures(plt, fvm, times, Y, T_out, C_out, Ta_fun, Ca_fun):
    FIGURES_DIR.mkdir(parents=True, exist_ok=True)
    n = fvm.n_nodes
    r_cm = fvm.r * 100.0
    idx = {float(t): i for i, t in enumerate(times)}
    marks = TABLE_TIMES

    fig, ax = plt.subplots(figsize=(6.4, 3.8))
    for t in marks:
        ax.plot(r_cm, Y[idx[float(t)], :n], lw=1.2, label=fr"$t={t/3600:g}\,\mathrm{{h}}$")
    ax.set_xlabel("半径 $r$ (cm)")
    ax.set_ylabel("温度 $T$ (°C)")
    ax.set_xlim(0, 2)
    ax.legend(frameon=False, ncol=3, loc="lower center", bbox_to_anchor=(0.5, 1.02))
    fig.tight_layout()
    fig.savefig(FIGURES_DIR / "q2_T_profiles.pdf")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(6.4, 3.8))
    for t in marks:
        ax.plot(r_cm, Y[idx[float(t)], n:], lw=1.2, label=fr"$t={t/3600:g}\,\mathrm{{h}}$")
    ax.set_xlabel("半径 $r$ (cm)")
    ax.set_ylabel(r"干基含水率 $C$ (kg/kg)")
    ax.set_xlim(0, 2)
    ax.legend(frameon=False, ncol=2)
    fig.tight_layout()
    fig.savefig(FIGURES_DIR / "q2_C_profiles.pdf")
    plt.close(fig)

    t_plot = times
    fig, ax = plt.subplots(figsize=(6.4, 3.8))
    ax.plot(t_plot / 3600.0, [Ta_fun(t) for t in t_plot], lw=1.0, ls="--", label="烘房 $T_a$")
    ax.plot(t_plot / 3600.0, T_out[:, -1], lw=1.2, label="表面")
    ax.plot(t_plot / 3600.0, T_out[:, 0], lw=1.2, label="中心")
    ax.set_xlabel("时间 $t$ (h)")
    ax.set_ylabel("温度 (°C)")
    ax.set_xlim(0, 3)
    ax.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(FIGURES_DIR / "q2_T_center_surface.pdf")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(6.4, 3.8))
    ax.plot(t_plot / 3600.0, C_out[:, -1], lw=1.2, label="表面")
    ax.plot(t_plot / 3600.0, C_out[:, 0], lw=1.2, label="中心")
    ax.set_xlabel("时间 $t$ (h)")
    ax.set_ylabel(r"干基含水率 $C$ (kg/kg)")
    ax.set_xlim(0, 3)
    ax.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(FIGURES_DIR / "q2_C_center_surface.pdf")
    plt.close(fig)

    D_c = moisture_diffusivity_q2(C_out[:, 0], T_out[:, 0])
    D_s = moisture_diffusivity_q2(C_out[:, -1], T_out[:, -1])
    fig, ax = plt.subplots(figsize=(6.4, 3.6))
    ax.plot(t_plot / 3600.0, D_s, lw=1.2, label="表面 $D$")
    ax.plot(t_plot / 3600.0, D_c, lw=1.2, label="中心 $D$")
    ax.set_xlabel("时间 $t$ (h)")
    ax.set_ylabel(r"$D$ (m$^2$/s)")
    ax.set_xlim(0, 3)
    ax.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(FIGURES_DIR / "q2_D_center_surface.pdf")
    plt.close(fig)


def run_mms(level: int = 4) -> dict:
    r = graded_radial_nodes(level)
    fvm = CoupledRadialFVM(r)
    T0f, C0f = manufactured_fields(fvm.r, 0.0)
    y0 = fvm.pack(T0f, C0f)
    Ta_fun = ManufacturedOven("T")
    Ca_fun = ManufacturedOven("C")
    times, Y = integrate_coupled(
        fvm, y0, 100.0, Ta_fun, Ca_fun,
        rtol=1e-8, atol=1e-9, max_step=0.5,
        source=manufactured_sources,
        t_eval=np.arange(1.0, 100.5, 1.0),
    )
    n = fvm.n_nodes
    T_num, C_num = Y[:, :n], Y[:, n:]
    T_ex = np.stack([manufactured_fields(fvm.r, t)[0] for t in times])
    C_ex = np.stack([manufactured_fields(fvm.r, t)[1] for t in times])
    return {
        "level": level,
        "n_nodes": n,
        "T_max_abs": float(np.max(np.abs(T_num - T_ex))),
        "C_max_abs": float(np.max(np.abs(C_num - C_ex))),
        "T_end_max_abs": float(np.max(np.abs(T_num[-1] - T_ex[-1]))),
        "C_end_max_abs": float(np.max(np.abs(C_num[-1] - C_ex[-1]))),
    }


def main():
    t_wall0 = time.perf_counter()
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    Ta_fun, Ca_fun, knots = load_q2_inputs()
    print("oven points", knots.size, "kind", Ta_fun.kind, "Ta(10800)=", Ta_fun(T_END), "Ca(10800)=", Ca_fun(T_END))
    D0 = float(moisture_diffusivity_q2(C0, T0))
    print("initial D", D0, "k", float(conductivity(C0)), "B", float(heat_capacity(C0)))

    print("manufactured-solution probe")
    mms = run_mms(4)
    print("MMS", mms)

    mesh_levels = [2, 4]
    rtol, atol_T, atol_C, max_step = 1e-9, 1e-8, 1e-10, 1.0
    solutions = {}
    for level in mesh_levels:
        r = graded_radial_nodes(level)
        fvm = CoupledRadialFVM(r)
        tag = f"G{level}"
        print(f"\n=== {tag} nodes={fvm.n_nodes} ===")
        y0 = fvm.pack(np.full(fvm.n_nodes, T0), np.full(fvm.n_nodes, C0))
        atol = np.concatenate([np.full(fvm.n_nodes, atol_T), np.full(fvm.n_nodes, atol_C)])
        times, Y = integrate_coupled(fvm, y0, T_END, Ta_fun, Ca_fun, rtol, atol, max_step)
        times_pos = times[times > 0]
        Y_pos = Y[times > 0]
        T_out, C_out = extract_output(fvm, Y_pos)
        solutions[level] = dict(
            tag=tag, fvm=fvm, times=times_pos, Y=Y_pos, Y_full=Y, times_full=times,
            T_out=T_out, C_out=C_out, n_nodes=fvm.n_nodes,
        )
        print("done", tag, "T", T_out.min(), T_out.max(), "C", C_out.min(), C_out.max(),
              "C(t=1,s)", C_out[0, -1], "C(3h,s)", C_out[-1, -1], "C(3h,c)", C_out[-1, 0])

    t = solutions[4]["times"]
    conv = {
        "T_G2_vs_G4": compare_on_output(solutions[2]["T_out"], solutions[4]["T_out"], t),
        "C_G2_vs_G4": compare_on_output(solutions[2]["C_out"], solutions[4]["C_out"], t),
    }
    fourdec = {
        "T_G2_vs_G4": four_decimal_stats(solutions[2]["T_out"], solutions[4]["T_out"], t),
        "C_G2_vs_G4": four_decimal_stats(solutions[2]["C_out"], solutions[4]["C_out"], t),
    }
    print("conv", conv)
    print("4dp", fourdec)
    paper_ok = fourdec["T_G2_vs_G4"]["paper_changed_cells"] == 0 and fourdec["C_G2_vs_G4"]["paper_changed_cells"] == 0
    abs_ok = conv["T_G2_vs_G4"]["max_abs"] <= 2e-5 and conv["C_G2_vs_G4"]["max_abs"] <= 2e-5
    need_g8 = not (paper_ok and abs_ok)
    if need_g8:
        print("G2 vs G4 not enough; running G8")
        level = 8
        r = graded_radial_nodes(level)
        fvm = CoupledRadialFVM(r)
        y0 = fvm.pack(np.full(fvm.n_nodes, T0), np.full(fvm.n_nodes, C0))
        atol = np.concatenate([np.full(fvm.n_nodes, atol_T), np.full(fvm.n_nodes, atol_C)])
        times, Y = integrate_coupled(fvm, y0, T_END, Ta_fun, Ca_fun, rtol, atol, max_step)
        times_pos = times[times > 0]
        Y_pos = Y[times > 0]
        T_out, C_out = extract_output(fvm, Y_pos)
        solutions[8] = dict(
            tag="G8", fvm=fvm, times=times_pos, Y=Y_pos, Y_full=Y, times_full=times,
            T_out=T_out, C_out=C_out, n_nodes=fvm.n_nodes,
        )
        t = solutions[8]["times"]
        conv["T_G4_vs_G8"] = compare_on_output(solutions[4]["T_out"], solutions[8]["T_out"], t)
        conv["C_G4_vs_G8"] = compare_on_output(solutions[4]["C_out"], solutions[8]["C_out"], t)
        fourdec["T_G4_vs_G8"] = four_decimal_stats(solutions[4]["T_out"], solutions[8]["T_out"], t)
        fourdec["C_G4_vs_G8"] = four_decimal_stats(solutions[4]["C_out"], solutions[8]["C_out"], t)
        print("conv G4 G8", conv["T_G4_vs_G8"], conv["C_G4_vs_G8"])
        pub_level = 8
        paper_ok = fourdec["T_G4_vs_G8"]["paper_changed_cells"] == 0 and fourdec["C_G4_vs_G8"]["paper_changed_cells"] == 0
        abs_ok = conv["T_G4_vs_G8"]["max_abs"] <= 2e-5 and conv["C_G4_vs_G8"]["max_abs"] <= 2e-5
    else:
        pub_level = 4 if fourdec["C_G2_vs_G4"]["changed_cells"] or fourdec["T_G2_vs_G4"]["changed_cells"] else 2
        if pub_level == 2:
            pub_level = 4
    pub = solutions[pub_level]
    fvm = pub["fvm"]
    times = pub["times"]
    print("publishing", pub["tag"], "paper_ok", paper_ok, "abs_ok", abs_ok)

    print("time refinement")
    y0 = fvm.pack(np.full(fvm.n_nodes, T0), np.full(fvm.n_nodes, C0))
    atol = np.concatenate([np.full(fvm.n_nodes, atol_T / 5), np.full(fvm.n_nodes, atol_C / 5)])
    t2, Y2 = integrate_coupled(fvm, y0, T_END, Ta_fun, Ca_fun, rtol / 5, atol, 0.5)
    T2, C2 = extract_output(fvm, Y2[t2 > 0])
    time_sens = {"T": compare_on_output(pub["T_out"], T2, times), "C": compare_on_output(pub["C_out"], C2, times)}
    time_4dp = {"T": four_decimal_stats(pub["T_out"], T2, times), "C": four_decimal_stats(pub["C_out"], C2, times)}
    print("time sens", time_sens)

    print("frozen-D moisture probe")
    D_freeze = D0
    Bi_m = HM * R / D_freeze
    C_ref, moisture_truncation = converged_reference(Bi_m, D_freeze, R, OUTPUT_RADII_M, times, Ca_fun, C0)
    y0f = fvm.pack(np.full(fvm.n_nodes, T0), np.full(fvm.n_nodes, C0))
    atol = np.concatenate([np.full(fvm.n_nodes, atol_T), np.full(fvm.n_nodes, atol_C)])
    tf, Yf = integrate_coupled(fvm, y0f, T_END, Ta_fun, Ca_fun, rtol, atol, max_step, D_const=D_freeze)
    _, C_frozen = extract_output(fvm, Yf[tf > 0])
    frozen = compare_on_output(C_frozen, C_ref, times)
    nl_vs_frozen = compare_on_output(pub["C_out"], C_frozen, times)
    print("frozen vs analytic", frozen)
    print("nonlinear vs frozen", nl_vs_frozen)

    n = fvm.n_nodes
    D_field = moisture_diffusivity_q2(pub["Y"][:, n:], pub["Y"][:, :n])
    D_stats = {
        "D0": D0,
        "D_min": float(np.min(D_field)),
        "D_max": float(np.max(D_field)),
        "D_min_over_D0": float(np.min(D_field) / D0),
        "D_max_over_D0": float(np.max(D_field) / D0),
        "center_C_end": float(pub["C_out"][-1, 0]),
        "surface_C_end": float(pub["C_out"][-1, -1]),
        "center_T_end": float(pub["T_out"][-1, 0]),
        "surface_T_end": float(pub["T_out"][-1, -1]),
    }
    print("D stats", D_stats)

    bounds = assert_physical(pub["Y_full"][:, :n], pub["Y_full"][:, n:], Ta_fun, Ca_fun)
    cons_C = conservation_moisture(fvm, pub["times_full"], pub["Y_full"], Ca_fun)
    cons_T = conservation_heat(fvm, pub["times_full"], pub["Y_full"], Ta_fun)
    print("bounds", bounds)
    print("conservation C", cons_C)
    print("conservation T", cons_T)

    delivery = make_gate({
        "space_abs": abs_ok,
        "space_paper": paper_ok,
        "all_finite": np.isfinite(pub["Y_full"]).all(),
        "manufactured_solution": mms["T_max_abs"] <= 2e-5 and mms["C_max_abs"] <= 2e-5,
        "time_abs": time_sens["T"]["max_abs"] <= 2e-5 and time_sens["C"]["max_abs"] <= 2e-5,
        "time_paper": time_4dp["T"]["paper_changed_cells"] == 0 and time_4dp["C"]["paper_changed_cells"] == 0,
        "temperature_bounds": bounds["T_bounds_ok"],
        "moisture_bounds": bounds["C_bounds_ok"],
        "moisture_reference_truncation": moisture_truncation["passed"],
        "frozen_reference_error": frozen["max_abs"] <= 2e-5,
        "heat_balance": cons_T["relative_residual"] <= 1e-6,
        "moisture_balance": cons_C["relative_residual"] <= 1e-6,
    })
    print("delivery", delivery)

    paper_T = {f"h{th:g}_r{rc:g}": float(pub["T_out"][int(t) - 1, int(round(rc / 0.1))])
               for th, t in zip(TABLE_HOURS, TABLE_TIMES) for rc in TABLE_RADII_CM}
    paper_C = {f"h{th:g}_r{rc:g}": float(pub["C_out"][int(t) - 1, int(round(rc / 0.1))])
               for th, t in zip(TABLE_HOURS, TABLE_TIMES) for rc in TABLE_RADII_CM}

    validation = {
        "result_version": "q2-closeout-v1",
        "interpolation": Ta_fun.kind,
        "published_mesh": pub["tag"],
        "published_n_nodes": pub["n_nodes"],
        "mesh_levels": list(solutions.keys()),
        "solver": "coupled-FVM+BDF",
        "rtol": rtol, "atol_T": atol_T, "atol_C": atol_C, "max_step": max_step,
        "mms": mms,
        "grid_convergence": conv,
        "four_decimal_grid": fourdec,
        "time_sensitivity": time_sens,
        "four_decimal_time": time_4dp,
        "delivery_gate": delivery,
        "reference_truncation_C": moisture_truncation,
        "frozen_D_moisture_vs_analytic": frozen,
        "nonlinear_C_vs_frozen_D": nl_vs_frozen,
        "D_variation": D_stats,
        "conservation_T_10800s": cons_T,
        "conservation_C_10800s": cons_C,
        "physical_bounds": bounds,
        "paper_table_T": paper_T,
        "paper_table_C": paper_C,
        "elapsed_s": time.perf_counter() - t_wall0,
        "python": sys.version,
        "end_face_2d": "skipped_by_plan",
    }
    require_delivery(validation, RESULTS_DIR / "diagnostics/q2_delivery/latest.json")
    np.savez(
        RESULTS_DIR / "q2_solution.npz",
        mesh_level=pub_level,
        n_nodes=pub["n_nodes"],
        times=times,
        r=fvm.r,
        output_radii_m=OUTPUT_RADII_M,
        checkpoint_times=np.r_[0., 1., 60., 100., TABLE_TIMES],
        T_checkpoints=pub["Y_full"][np.r_[0., 1., 60., 100., TABLE_TIMES].astype(int), :n],
        C_checkpoints=pub["Y_full"][np.r_[0., 1., 60., 100., TABLE_TIMES].astype(int), n:],
        T_out=pub["T_out"],
        C_out=pub["C_out"],
        interpolation=np.array(Ta_fun.kind),
        rtol=rtol, atol_T=atol_T, atol_C=atol_C, max_step=max_step,
    )
    write_result_workbook(RESULTS_DIR / "result2.xlsx", pub["T_out"], pub["C_out"], times)
    write_paper_csv(RESULTS_DIR / "q2_table_temperature.csv", pub["T_out"], times)
    write_paper_csv(RESULTS_DIR / "q2_table_moisture.csv", pub["C_out"], times)
    plt = setup_mpl()
    plot_figures(plt, fvm, times, pub["Y"], pub["T_out"], pub["C_out"], Ta_fun, Ca_fun)

    (RESULTS_DIR / "q2_validation.json").write_text(
        json.dumps(validation, ensure_ascii=False, indent=2, default=float),
        encoding="utf-8",
    )
    print("elapsed_s", validation["elapsed_s"])
    print("wrote results to", RESULTS_DIR)



if __name__ == "__main__":
    main()
'''

def prepare():
    code_dir = PACKAGE_ROOT / '_runtime' / f'Q{QUESTION}' / 'code'
    code_dir.mkdir(parents=True, exist_ok=True)
    for name, source in SOURCES.items():
        (code_dir / (name + '.py')).write_text(source, encoding='utf-8')
    sys.path.insert(0, str(code_dir))
    for folder in ('results', 'figures', 'reports'):
        (PACKAGE_ROOT / folder).mkdir(exist_ok=True)
    os.environ.setdefault('MPLBACKEND', 'Agg')
    for name in SOURCES:
        importlib.import_module(name)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--check', action='store_true', help='Check imports and original inputs without solving')
    parser.add_argument('--export-only', action='store_true', help='Export existing validated NPZ arrays; does not solve the model')
    args = parser.parse_args()
    prepare()
    from utils import load_oven_table
    load_oven_table()
    if QUESTION == 4:
        from problem4 import Inputs
        Inputs()
    if args.check:
        print(f'Q{QUESTION}: imports and input checks passed', flush=True)
        return
    from reproduction_export import export_saved
    if args.export_only:
        export_saved(PACKAGE_ROOT, QUESTION)
        return
    if QUESTION == 1:
        import problem1
        problem1.main()
    elif QUESTION == 2:
        import problem2
        problem2.main()
    elif QUESTION == 3:
        # The original delivery gate compares Q3 with independently computed Q2.
        if not (PACKAGE_ROOT / 'results' / 'q2_solution.npz').exists():
            print('Q3: computing Q2 first for the original overlap check', flush=True)
            import problem2
            problem2.main()
            export_saved(PACKAGE_ROOT, 2)
        import q3_closeout
        q3_closeout.main(['--no-xlsx'])
    else:
        import problem4
        import verify_q4_results
        for level, tight in ((2, False), (4, False), (4, True), (8, False), (16, False), (8, True)):
            problem4.run(level, tight)
        verify_q4_results.main()
    export_saved(PACKAGE_ROOT, QUESTION)
    print(f'Q{QUESTION}: full recomputation and Excel export completed', flush=True)


if __name__ == '__main__':
    main()
