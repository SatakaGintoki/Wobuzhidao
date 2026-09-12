"""Q2 audit. Original results and production solver are never overwritten."""
from pathlib import Path
import hashlib
import json
import re
import sys
import time
import numpy as np

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'results/verification/q2'
OUT.mkdir(parents=True, exist_ok=True)

def dump(name, data):
    (OUT/name).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    print(name, json.dumps(data, ensure_ascii=True), flush=True)

def files():
    from openpyxl import load_workbook
    import pypdfium2 as pdfium
    from PIL import Image, ImageDraw
    paths = [ROOT/'data/A.pdf', ROOT/'data/附件/附件1.xlsx', ROOT/'data/附件/附件3/result2.xlsx']
    paths += list((ROOT/'results').glob('q2*')) + [ROOT/'results/result2.xlsx']
    paths += list((ROOT/'figures').glob('q2*.pdf'))
    paths += [ROOT/'code'/n for n in ['problem2.py','radial_coupled.py','appendix3.py','utils.py','radial_fvm.py']]
    hashes = {str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in paths}
    z = np.load(ROOT/'results/q2_solution.npz')
    v = json.loads((ROOT/'results/q2_validation.json').read_text(encoding='utf-8'))
    report = dict(source_sha256=hashes, npz_keys=z.files, workbook={}, tables={}, figures=[])
    report['axes'] = dict(times_ok=bool(np.array_equal(z['times'],np.arange(1,10801))),
                         grid_increasing=bool(np.all(np.diff(z['r'])>0)), n_nodes=len(z['r']))
    template = load_workbook(ROOT/'data/附件/附件3/result2.xlsx',read_only=True)
    report['template'] = {s.title:dict(dimensions=[s.max_row,s.max_column], first_rows=list(s.values)[:4]) for s in template}
    template.close()
    wb = load_workbook(ROOT/'results/result2.xlsx',read_only=True,data_only=False)
    report['sheet_order_ok'] = wb.sheetnames == ['温度','水分浓度']
    src = (ROOT/'reports/RESULTS_REPORT.md').read_text(encoding='utf-8').split('## 问题二结果')[1]
    lines = [x for x in src.splitlines() if re.match(r'^\| [0-3]\.\d \|',x)]
    for k,(sheet,key,filename) in enumerate([('温度','T','temperature'),('水分浓度','C','moisture')]):
        ws=wb[sheet]; it=ws.iter_rows(); header=next(it)
        values=[]; times=[]; numeric=formatted=0
        for row in it:
            times.append(row[0].value)
            values.append([c.value for c in row[1:]])
            numeric += sum(c.data_type=='n' for c in row[1:])
            formatted += sum(c.number_format=='0.0000' for c in row[1:])
        a=np.array(values,dtype=float); expected=z[key+'_out']
        report['workbook'][key] = dict(dimensions=[ws.max_row,ws.max_column],
            times_ok=bool(np.array_equal(times,z['times'])),
            radii_ok=bool(np.allclose([c.value for c in header[1:]],np.arange(21)*.1,rtol=0,atol=1e-14)),
            all_finite=bool(np.isfinite(a).all()), numeric_cells=numeric, formatted_cells=formatted,
            max_difference_vs_npz=float(np.max(abs(a-expected))), minimum=float(a.min()), maximum=float(a.max()))
        expected=expected[np.arange(1800,10801,1800)-1,::5]
        csv=np.loadtxt(ROOT/f'results/q2_table_{filename}.csv',delimiter=',',skiprows=1)
        printed=np.array([[float(s.strip()) for s in line.split('|')[2:-1]] for line in lines[6*k:6*k+6]])
        j=np.array([v['paper_table_'+key][f'h{h:g}_r{r:g}'] for h in np.arange(.5,3.1,.5) for r in np.arange(5)*.5]).reshape(6,5)
        report['tables'][key]=dict(csv_max_difference=float(abs(csv[:,1:]-expected).max()),
            csv_hours_ok=bool(np.array_equal(csv[:,0],np.arange(.5,3.1,.5))),
            report_four_decimal_matches=bool(np.array_equal(printed,np.round(expected,4))),
            json_max_difference=float(abs(j-expected).max()))
    wb.close()
    previews=[]
    for p in sorted((ROOT/'figures').glob('q2*.pdf')):
        doc=pdfium.PdfDocument(str(p)); texts=[]
        for i in range(len(doc)):
            pg=doc[i]; tp=pg.get_textpage(); texts.append(tp.get_text_range()); tp.close()
            bm=pg.render(scale=2); im=bm.to_pil().convert('RGB'); im.save(OUT/f'{p.stem}_{i+1}.png')
            previews.append((p.stem,im.copy())); bm.close(); pg.close()
        report['figures'].append(dict(name=p.name,pages=len(doc),text='\n'.join(texts)))
        doc.close()
    contact=Image.new('RGB',(1600,1600),'white'); draw=ImageDraw.Draw(contact)
    for i,(name,im) in enumerate(previews):
        im.thumbnail((790,480)); x=i%2*800; y=i//2*530
        contact.paste(im,(x,y+25));draw.text((x+10,y+5),name,fill='black')
    contact.save(OUT/'figure_contact.png')
    doc=pdfium.PdfDocument(str(ROOT/'data/A.pdf'))
    for i in [1,3]:
        pg=doc[i]; bm=pg.render(scale=1.5); bm.to_pil().save(OUT/f'problem_page_{i+1}.png'); bm.close();pg.close()
    doc.close()
    report['source_files_unchanged']=all(hashlib.sha256((ROOT/p).read_bytes()).hexdigest()==h for p,h in hashes.items())
    dump('file_audit.json',report)

def numerical():
    import problem2 as p
    start=time.perf_counter()
    z=np.load(ROOT/'results/q2_solution.npz')
    Ta,Ca,knots=p.load_q2_inputs()
    report={'scope':'G4 reproduction, G4 time refinement, G8 space refinement, all 10800 seconds', 'runs':{}}
    for tag,level,rtol,at,step in [('G4_reproduction',4,1e-9,1e-8,1.),('G4_time',4,2e-10,2e-9,.5),('G8_space',8,1e-9,1e-8,1.)]:
        print('Starting',tag,flush=True)
        f=p.CoupledRadialFVM(p.graded_radial_nodes(level)); n=f.n_nodes
        y0=f.pack(np.full(n,p.T0),np.full(n,p.C0))
        ts,Y=p.integrate_coupled(f,y0,p.T_END,Ta,Ca,rtol,at,step)
        T,C=p.extract_output(f,Y[1:])
        result={key:{**p.compare_on_output(z[key+'_out'],a,ts[1:]),**p.four_decimal_stats(z[key+'_out'],a,ts[1:])} for key,a in [('T',T),('C',C)]}
        if tag=='G4_reproduction':
            result['conservation_C']=p.conservation_moisture(f,ts,Y,Ca)
            result['conservation_T']=p.conservation_heat(f,ts,Y,Ta)
            result['full_grid_bounds']={ 'T_min':float(Y[:,:n].min()),'T_max':float(Y[:,:n].max()),'C_min':float(Y[:,n:].min()),'C_max':float(Y[:,n:].max())}
            cp=np.array([0,1,60,100,1800,3600,5400,7200,9000,10800])
            np.savez_compressed(OUT/'G4_checkpoints.npz',times=ts[cp],r=f.r,Y=Y[cp])
        if tag=='G8_space':
            changed=[]
            for key,a in [('T',T),('C',C)]:
                for ti in np.arange(1800,10801,1800)-1:
                    for ri in range(0,21,5):
                        b=float(z[key+'_out'][ti,ri]); q=float(a[ti,ri])
                        if round(b,4)!=round(q,4): changed.append(dict(field=key,t_s=int(ti+1),r_cm=ri*.1,G4=b,G8=q))
            result['paper_changed_details']=changed
        np.savez_compressed(OUT/f'{tag}_outputs.npz',times=ts[1:],T_out=T,C_out=C)
        report['runs'][tag]=result
        dump('numerical_audit.json',report)
        del Y,T,C,f
    report['manufactured_solution']=p.run_mms(4)
    # A short frozen-D check separates modal truncation from early FVM error.
    f=p.CoupledRadialFVM(p.graded_radial_nodes(4)); n=f.n_nodes
    D0=float(p.moisture_diffusivity_q2(p.C0,p.T0))
    ts,Y=p.integrate_coupled(f,f.pack(np.full(n,p.T0),np.full(n,p.C0)),100.,Ta,Ca,2e-10,2e-9,.1,D_const=D0)
    _,Cf=p.extract_output(f,Y[1:]); refs={}
    for terms in [80,160,320]:
        ref=p.BesselDuhamel(p.HM*p.R/D0,D0,p.R,n_terms=terms).evaluate_many(p.OUTPUT_RADII_M,ts[1:],Ca,p.C0)
        refs[terms]=ref
        report.setdefault('frozen_D_early',{})[str(terms)]=dict(max_abs=float(abs(Cf-ref).max()),t1_surface_numeric=float(Cf[0,-1]),t1_surface_analytic=float(ref[0,-1]))
    report['frozen_D_early']['160_vs_320_max_abs']=float(abs(refs[160]-refs[320]).max())
    report['elapsed_s']=time.perf_counter()-start
    dump('numerical_audit.json',report)

def modal():
    # Closed integration of the first PCHIP derivative polynomial at t=1 s.
    # This avoids the time integrator used by the production modal reference.
    import problem2 as p
    from scipy.special import hyp1f1
    _,env,_=p.load_q2_inputs()
    coeff=env._pchip.c[:,0]
    D0=float(p.moisture_diffusivity_q2(p.C0,p.T0))
    values={}
    for n in [80,160,320,640]:
        ref=p.BesselDuhamel(p.HM*p.R/D0,D0,p.R,n_terms=n)
        beta=ref.beta
        moments=[hyp1f1(1,k+2,-beta)/(k+1) for k in range(3)]
        modes=(float(env(0))-p.C0)*np.exp(-beta)+coeff[2]*moments[0]+2*coeff[1]*moments[1]+3*coeff[0]*moments[2]
        values[str(n)]=float(env(1)-modes@ref.radial_modes(np.array([p.R]))[:,0])
    dump('modal_truncation_audit.json',dict(t_s=1,surface_C=values,
        difference_320_vs_640=abs(values['320']-values['640']),method='exact first-interval PCHIP polynomial convolution, hypergeometric moments'))

if __name__=='__main__':
    if '--numerical' in sys.argv: numerical()
    elif '--modal' in sys.argv: modal()
    else: files()
