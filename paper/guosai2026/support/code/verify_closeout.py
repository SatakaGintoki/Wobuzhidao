"""Audit the Q1/Q2 closeout against the preserved pre-closeout deliverables.

Run with numerical Python for --gate/--modal; bundled Python for --files/--render.
"""
from pathlib import Path
import ast
import hashlib
import json
import sys
import tempfile
import numpy as np

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'results/verification/closeout'
ARCHIVE = ROOT / 'results/archive/pre-closeout-20260911'


def dump(name, report):
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT/name).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=True, indent=2), flush=True)


def gate_test():
    from delivery import make_gate, require_delivery
    results = {}
    with tempfile.TemporaryDirectory(dir=ROOT/'tmp') as d:
        folder = Path(d)
        for q in (1,2):
            v = json.loads((ROOT/f'results/q{q}_validation.json').read_text(encoding='utf-8'))
            baseline = v['delivery_gate']['checks']
            assert baseline and all(baseline.values())
            for key in baseline:
                checks = dict(baseline, **{key: False})
                sentinel = folder/'result.xlsx'
                sentinel.write_bytes(b'existing approved data')
                try:
                    require_delivery({'delivery_gate': make_gate(checks)}, folder/'diagnostic.json')
                    sentinel.write_bytes(b'invalid new result')
                except RuntimeError:
                    pass
                assert sentinel.read_bytes() == b'existing approved data'
                assert key in json.loads((folder/'diagnostic.json').read_text(encoding='utf-8'))['delivery_gate']['failed_checks']
            require_delivery({'delivery_gate': make_gate(baseline)}, folder/'diagnostic.json')
            # Ensure the production main itself puts the guard ahead of all writers.
            tree = ast.parse((ROOT/f'code/problem{q}.py').read_text(encoding='utf-8'))
            main = next(n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name == 'main')
            calls = [n for n in ast.walk(main) if isinstance(n, ast.Call)]
            guard = next(n.lineno for n in calls if isinstance(n.func,ast.Name) and n.func.id=='require_delivery')
            writers = [n.lineno for n in calls if (isinstance(n.func,ast.Name) and n.func.id in ('write_result_workbook','write_paper_csv','plot_figures')) or (isinstance(n.func,ast.Attribute) and n.func.attr in ('savez','write_text'))]
            assert writers and min(writers)>guard
            results[f'q{q}'] = {'failure_cases_blocked': len(baseline), 'valid_case_passed': True, 'all_production_writers_after_guard': True}
    dump('gate_test.json', results)


def files():
    from openpyxl import load_workbook
    report = {}
    for q, end, table_t in [(1,1800,np.array([100,300,600,900,1200,1500,1800])),(2,10800,np.arange(1800,10801,1800))]:
        z = np.load(ROOT/f'results/q{q}_solution.npz')
        old = np.load(ARCHIVE/f'results/q{q}_solution.npz')
        validation = json.loads((ROOT/f'results/q{q}_validation.json').read_text(encoding='utf-8'))
        r = {'version': validation['result_version'], 'delivery_passed': validation['delivery_gate']['meets_delivery_gate'], 'fields':{}}
        assert r['delivery_passed']
        assert np.array_equal(z['times'],np.arange(1,end+1))
        assert np.all(np.diff(z['r'])>0)
        assert np.allclose(z['output_radii_m'], np.arange(21)*.001,rtol=0,atol=1e-15)
        idx = np.array([np.argmin(abs(z['r']-x)) for x in z['output_radii_m']])
        assert np.max(abs(z['r'][idx]-z['output_radii_m'])) < 1e-14
        wb = load_workbook(ROOT/f'results/result{q}.xlsx',read_only=True,data_only=False)
        before = load_workbook(ARCHIVE/f'results/result{q}.xlsx',read_only=True,data_only=False)
        assert wb.sheetnames == before.sheetnames == ['温度','水分浓度']
        for key, name, kind in [('T','温度','temperature'),('C','水分浓度','moisture')]:
            ws = wb[name]
            assert (ws.max_row,ws.max_column)==(end+1,22)
            same = True
            rows = ws.iter_rows()
            header=next(rows)
            assert np.allclose([c.value for c in header[1:]], np.arange(21)*.1,atol=1e-14,rtol=0)
            values=[]
            numeric=formatted=0
            for i,row in enumerate(rows,1):
                assert row[0].value==i
                values.append([c.value for c in row[1:]])
                numeric+=sum(c.data_type=='n' for c in row[1:])
                formatted+=sum(c.number_format=='0.0000' for c in row[1:])
            a=np.array(values)
            assert np.isfinite(a).all()
            assert numeric==formatted==end*21
            assert np.max(abs(a-z[key+'_out']))<1e-12
            for row1,row0 in zip(ws.iter_rows(),before[name].iter_rows(),strict=True):
                if any((a.value,a.number_format,a.data_type)!=(b.value,b.number_format,b.data_type) for a,b in zip(row1,row0,strict=True)):
                    same=False
            diff=float(np.max(abs(z[key+'_out']-old[key+'_out'])))
            assert diff<=2e-5
            assert np.array_equal(np.round(z[key+'_out'][table_t-1,::5],4),np.round(old[key+'_out'][table_t-1,::5],4))
            csv=np.loadtxt(ROOT/f'results/q{q}_table_{kind}.csv',delimiter=',',skiprows=1)
            assert np.max(abs(csv[:,1:]-z[key+'_out'][table_t-1,::5]))<=5.1e-9  # CSV intentionally stores 8 decimals.
            cp=z[key+'_checkpoints']
            assert cp.shape==(len(z['checkpoint_times']),len(z['r']))
            assert np.all(cp[0]==(28 if key=='T' else 2.55))
            assert np.allclose(cp[1:,idx],z[key+'_out'][z['checkpoint_times'][1:].astype(int)-1],rtol=0,atol=1e-12)
            if key+'_full' in z.files:
                assert np.max(abs(z[key+'_full'][:,idx]-z[key+'_out']))<1e-12
            r['fields'][key]={'output_cells':int(a.size),'max_change_from_previous':diff,'all_workbook_values_types_formats_unchanged':same,'paper_four_decimals_unchanged':True,'excel_npz_csv_checkpoints_consistent':True}
        wb.close();before.close()
        r['source_sha256']={str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in [ROOT/f'code/problem{q}.py',ROOT/f'results/q{q}_solution.npz',ROOT/f'results/q{q}_validation.json',ROOT/f'results/result{q}.xlsx']}
        report[f'q{q}']=r
    dump('files.json',report)


def modal():
    from scipy.special import hyp1f1
    from q1_analytic import BesselDuhamel
    report={}
    for q in (1,2):
        p=__import__(f'problem{q}')
        _,env,_=getattr(p,f'load_q{q}_inputs')()
        D=float(p.moisture_diffusivity(p.C0) if q==1 else p.moisture_diffusivity_q2(p.C0,p.T0))
        values={}
        coeff=env._pchip.c[:,0]
        for terms in (120 if q==1 else 80,320,640,1280):
            ref=BesselDuhamel(p.HM*p.R/D,D,p.R,terms)
            moments=[hyp1f1(1,k+2,-ref.beta)/(k+1) for k in range(3)]
            I=(env(0)-p.C0)*np.exp(-ref.beta)+coeff[2]*moments[0]+2*coeff[1]*moments[1]+3*coeff[0]*moments[2]
            exact=env(1)-I@ref.radial_modes(p.OUTPUT_RADII_M)
            values[terms]=exact
        ode=BesselDuhamel(p.HM*p.R/D,D,p.R,640).evaluate_many(p.OUTPUT_RADII_M,np.array([1.]),env,p.C0)[0]
        diff=float(np.max(abs(ode-values[640])))
        assert diff<1e-8
        assert np.max(abs(values[640]-values[1280]))<1e-8
        report[f'q{q}']={'surface_C_at_1s':{str(k):float(v[-1]) for k,v in values.items()},'640_vs_1280_max':float(np.max(abs(values[640]-values[1280]))),'ode_vs_exact_convolution_max':diff,'passed':True}
    dump('modal.json',report)


def render():
    import pypdfium2 as pdfium
    from PIL import Image, ImageDraw
    folder=OUT/'previews';folder.mkdir(parents=True,exist_ok=True)
    records=[]
    paths=sorted((ROOT/'figures').glob('q[12]*.pdf'))
    if (ROOT/'paper/q12_closeout.pdf').exists(): paths.append(ROOT/'paper/q12_closeout.pdf')
    for path in paths:
        doc=pdfium.PdfDocument(str(path))
        for i in range(len(doc)):
            page=doc[i];bm=page.render(scale=1.5)
            dest=folder/f'{path.stem}_{i+1}.png'
            im=bm.to_pil().convert('RGB');im.save(dest)
            records.append((dest,im.copy()))
            bm.close();page.close()
        doc.close()
    for offset in range(0,len(records),6):
        contact=Image.new('RGB',(1500,1800),'white');draw=ImageDraw.Draw(contact)
        for j,(path,im) in enumerate(records[offset:offset+6]):
            im.thumbnail((740,560));x=j%2*750;y=j//2*600
            contact.paste(im,(x,y+25));draw.text((x+5,y+5),path.stem,fill='black')
        contact.save(folder/f'contact_{offset//6+1}.png')
    dump('renders.json',{'pages':[str(p.relative_to(ROOT)) for p,_ in records]})


if __name__=='__main__':
    if '--gate' in sys.argv: gate_test()
    elif '--modal' in sys.argv: modal()
    elif '--render' in sys.argv: render()
    else: files()
