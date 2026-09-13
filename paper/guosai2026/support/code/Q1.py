# -*- coding: utf-8 -*-
"""问题 1：从原始附件复算答案数据。

运行 python Q1.py；检查环境 python Q1.py --check。
本文件内嵌本问所需的全部项目模块，模块源码在下方按文件名分段列出。
运行时自动展开到相对目录 _runtime/Q1/code，以支持源码哈希核验。
输入、输出均相对于本文件的位置，与当前工作目录无关。
方程、物性、网格、容差和验证判据沿用原正式求解程序。
"""
import argparse
import importlib
import os
from pathlib import Path
import sys

PACKAGE_ROOT = Path(__file__).resolve().parent
QUESTION = 1
SOURCES = {}

# ========================================================================
# Module: utils.py
# ========================================================================
SOURCES['utils'] = r'''"""Shared paths and attachment readers for Problem A."""
from __future__ import annotations

from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.interpolate import PchipInterpolator

ROOT = Path(__file__).resolve().parents[3]
DATA_DIR = ROOT / "data"
ATTACH_DIR = DATA_DIR / "附件"
RESULTS_DIR = ROOT / "results"
FIGURES_DIR = ROOT / "figures"
REPORTS_DIR = ROOT / "reports"

# Appendix 2 (Problem 1)
R = 0.02
L = 0.25
RHO = 820.0
CP = 2600.0
K = 0.36
H = 25.0
HM = 8.0e-7
T0 = 28.0
C0 = 2.55
D_PREFACTOR = 7.0e-9
D_EXP_COEF = 0.89
T_END = 1800.0
ALPHA = K / (RHO * CP)


def moisture_diffusivity(C: np.ndarray | float, *, protect: bool = True) -> np.ndarray | float:
    """D(C) = 7e-9 exp(-0.89/C). Nonpositive C uses D=0 only as a solver-domain guard."""
    C_arr = np.asarray(C, dtype=float)
    out = np.zeros_like(C_arr, dtype=float)
    positive = C_arr > 0.0
    if np.any(positive):
        out[positive] = D_PREFACTOR * np.exp(-D_EXP_COEF / C_arr[positive])
    if not protect and np.any(~positive):
        raise ValueError("moisture diffusivity evaluated at nonpositive C without protection")
    if np.isscalar(C):
        return float(out)
    return out


def load_oven_table() -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    path = ATTACH_DIR / "附件1.xlsx"
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = list(wb.active.values)
    wb.close()
    data = np.asarray(rows[1:], dtype=float)
    t, Ta, Ca = data[:, 0], data[:, 1], data[:, 2]
    if t[0] != 0.0:
        raise ValueError(f"oven table must start at t=0, got {t[0]}")
    return t, Ta, Ca


class PiecewiseLinear:
    """Piecewise linear interpolant; kept for comparison, not used in Q1 production."""

    def __init__(self, t: np.ndarray, y: np.ndarray):
        self.t = np.asarray(t, dtype=float)
        self.y = np.asarray(y, dtype=float)
        self.kind = "linear"
        if self.t.ndim != 1 or self.y.shape != self.t.shape:
            raise ValueError("interpolant requires matching 1-D arrays")
        if np.any(np.diff(self.t) <= 0):
            raise ValueError("breakpoints must be strictly increasing")

    def __call__(self, t: float | np.ndarray) -> float | np.ndarray:
        return np.interp(t, self.t, self.y)

    def slope(self, t: float) -> float:
        t = float(t)
        if t <= self.t[0] or t >= self.t[-1]:
            if t >= self.t[-1]:
                dt = self.t[-1] - self.t[-2]
                return (self.y[-1] - self.y[-2]) / dt
            return (self.y[1] - self.y[0]) / (self.t[1] - self.t[0])
        i = int(np.searchsorted(self.t, t, side="right") - 1)
        i = min(i, len(self.t) - 2)
        return (self.y[i + 1] - self.y[i]) / (self.t[i + 1] - self.t[i])


class PchipOven:
    """Monotone piecewise cubic Hermite interpolant (PCHIP) through attachment knots."""

    def __init__(self, t: np.ndarray, y: np.ndarray):
        self.t = np.asarray(t, dtype=float)
        self.y = np.asarray(y, dtype=float)
        self.kind = "pchip"
        if self.t.ndim != 1 or self.y.shape != self.t.shape:
            raise ValueError("interpolant requires matching 1-D arrays")
        if np.any(np.diff(self.t) <= 0):
            raise ValueError("breakpoints must be strictly increasing")
        self._pchip = PchipInterpolator(self.t, self.y, extrapolate=False)
        self._deriv = self._pchip.derivative()

    def __call__(self, t: float | np.ndarray) -> float | np.ndarray:
        tt = np.clip(np.asarray(t, dtype=float), self.t[0], self.t[-1])
        val = self._pchip(tt)
        if np.ndim(t) == 0:
            return float(val)
        return np.asarray(val, dtype=float)

    def slope(self, t: float) -> float:
        tt = float(np.clip(t, self.t[0], self.t[-1]))
        return float(self._deriv(tt))
'''

# ========================================================================
# Module: delivery.py
# ========================================================================
SOURCES['delivery'] = r'''"""Fail-closed numerical delivery checks shared by Q1 and Q2."""
import json
from pathlib import Path


def make_gate(checks):
    checks = {name: bool(value) for name, value in checks.items()}
    return {"checks": checks, "failed_checks": [k for k, v in checks.items() if not v],
            "meets_delivery_gate": bool(checks) and all(checks.values())}


def require_delivery(validation, diagnostic_path):
    """Persist diagnostics first; never reach any production writer on failure."""
    path = Path(diagnostic_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(validation, ensure_ascii=False, indent=2, default=float), encoding="utf-8")
    gate = validation["delivery_gate"]
    checks = gate.get("checks", {})
    if not checks or not all(value is True for value in checks.values()) or gate.get("meets_delivery_gate") is not True:
        raise RuntimeError("Delivery blocked; see " + str(path) + ": " + str(gate.get("failed_checks", [])))
'''

# ========================================================================
# Module: q1_analytic.py
# ========================================================================
SOURCES['q1_analytic'] = r'''"""Independent Bessel–Duhamel reference for the linear radial heat (or frozen-D moisture) problem."""
from __future__ import annotations

from typing import Any

import numpy as np
from scipy.integrate import solve_ivp
from scipy.optimize import brentq
from scipy.special import j0, j1
from scipy.sparse import diags


def robin_eigenvalues(Bi: float, n_roots: int) -> np.ndarray:
    """Positive roots of lambda * J1(lambda) = Bi * J0(lambda)."""
    if Bi <= 0:
        raise ValueError("Biot number must be positive")
    n_roots = int(n_roots)

    def f(lam: float) -> float:
        return lam * j1(lam) - Bi * j0(lam)

    lam_max = (n_roots + 4) * np.pi
    grid = np.linspace(1e-10, lam_max, max(8000, 200 * n_roots))
    vals = grid * j1(grid) - Bi * j0(grid)
    roots: list[float] = []
    for i in range(len(grid) - 1):
        v0, v1 = vals[i], vals[i + 1]
        if not np.isfinite(v0) or not np.isfinite(v1):
            continue
        a, b = float(grid[i]), float(grid[i + 1])
        if v0 == 0.0:
            cand = a
        elif v0 * v1 < 0.0:
            cand = float(brentq(f, a, b, xtol=1e-14, maxiter=200))
        else:
            continue
        if not roots or abs(cand - roots[-1]) > 1e-10:
            roots.append(cand)
        if len(roots) >= n_roots:
            break
    if len(roots) < n_roots:
        raise RuntimeError(f"found only {len(roots)} Robin Bessel roots, need {n_roots}")
    return np.asarray(roots, dtype=float)


def expansion_coefficients(lam: np.ndarray) -> np.ndarray:
    j0l = j0(lam)
    j1l = j1(lam)
    return 2.0 * j1l / (lam * (j0l**2 + j1l**2))


class BesselDuhamel:
    def __init__(self, Bi: float, diffusivity: float, R: float, n_terms: int = 120):
        self.Bi = float(Bi)
        self.alpha = float(diffusivity)
        self.R = float(R)
        self.lam = robin_eigenvalues(self.Bi, n_terms)
        self.A = expansion_coefficients(self.lam)
        self.beta = self.alpha * self.lam**2 / self.R**2

    def radial_modes(self, r: np.ndarray) -> np.ndarray:
        """Return array shape (n_terms, n_r) of A_n J0(lambda_n r/R)."""
        r = np.asarray(r, dtype=float)
        return self.A[:, None] * j0(self.lam[:, None] * (r[None, :] / self.R))

    def evaluate(self, r: np.ndarray, t: float, env: Any, u0: float) -> np.ndarray:
        return self.evaluate_many(r, np.array([float(t)]), env, u0)[0]

    def evaluate_many(self, r: np.ndarray, times: np.ndarray, env: Any, u0: float) -> np.ndarray:
        """T = env(t) - sum A_n J0 I_n, with I' = -beta I + env'(t), I(0)=env(0)-u0.

        Valid for any C1 interpolant that exposes env.y, env.slope, and env(t).
        """
        r = np.asarray(r, dtype=float)
        times = np.asarray(times, dtype=float)
        if times.size == 0:
            return np.empty((0, r.size), dtype=float)
        modes = self.radial_modes(r)
        beta = self.beta
        I0 = np.full(beta.size, float(env.y[0] - u0))
        t_end = float(np.max(times))
        out = np.empty((times.size, r.size), dtype=float)
        zero = times == 0.0
        if np.all(zero):
            out[:] = u0
            return out
        t_eval = times[~zero]
        sol = solve_ivp(
            lambda t, I: -beta * I + env.slope(t),
            (0.0, max(t_end, 1e-30)),
            I0,
            method="BDF",
            jac=diags(-beta, format="csc"),
            rtol=1e-11,
            atol=1e-13,
            t_eval=t_eval,
            max_step=1.0,
            vectorized=False,
        )
        if not sol.success:
            raise RuntimeError(f"Bessel Duhamel ODE failed: {sol.message}")
        Ihist = sol.y.T
        env_t = np.array([float(env(t)) for t in t_eval], dtype=float)
        nonzero_vals = env_t[:, None] - Ihist @ modes
        out[zero] = u0
        out[~zero] = nonzero_vals
        return out


def converged_reference(Bi, diffusivity, R, r, times, env, u0, start_terms=320, tolerance=1e-8):
    """Double the modal truncation and explicitly check every requested point."""
    previous = BesselDuhamel(Bi, diffusivity, R, start_terms).evaluate_many(r, times, env, u0)
    history = []
    for terms in (start_terms * 2, start_terms * 4, start_terms * 8):
        current = BesselDuhamel(Bi, diffusivity, R, terms).evaluate_many(r, times, env, u0)
        error = float(np.max(np.abs(current - previous)))
        history.append({"terms": terms, "previous_terms": terms // 2, "max_abs_difference": error})
        if np.isfinite(error) and error <= tolerance:
            return current, {"passed": True, "tolerance": tolerance, "history": history, "terms": terms}
        previous = current
    return current, {"passed": False, "tolerance": tolerance, "history": history, "terms": terms}
'''

# ========================================================================
# Module: radial_fvm.py
# ========================================================================
SOURCES['radial_fvm'] = r'''"""Node-centered radial finite-volume operators for Problem 1."""
from __future__ import annotations

import numpy as np
from scipy.sparse import diags

from utils import CP, H, HM, K, R, RHO, moisture_diffusivity

OUTPUT_RADII_M = 0.001 * np.arange(21)


def graded_radial_nodes(level: int) -> np.ndarray:
    """Nodes denser near the surface; every 0.1 cm output radius remains a node.

    Interval subdivisions (before multiplying by ``level``):
    0–1.5 cm: 8; 1.5–1.8 cm: 16; 1.8–1.9 cm: 32; 1.9–2.0 cm: 64.
    Surface spacing at level L is 1e-3 / (64 L) metres.
    """
    level = int(level)
    if level < 1:
        raise ValueError("graded mesh level must be a positive integer")
    n_sub = np.full(20, 8 * level, dtype=int)
    n_sub[15:] = 16 * level
    n_sub[18] = 32 * level
    n_sub[19] = 64 * level
    pts = [0.0]
    for j in range(20):
        r0 = 0.001 * j
        r1 = 0.001 * (j + 1)
        m = int(n_sub[j])
        for k in range(1, m + 1):
            pts.append(r0 + (r1 - r0) * k / m)
    r = np.asarray(pts, dtype=float)
    r[0] = 0.0
    r[-1] = R
    return r


class RadialFVM:
    def __init__(self, N: int | None = None, r: np.ndarray | None = None):
        if r is not None:
            self.r = np.asarray(r, dtype=float)
            if self.r.ndim != 1 or self.r.size < 3:
                raise ValueError("r must be a 1-D array with at least 3 nodes")
            if abs(self.r[0]) > 1e-15 or abs(self.r[-1] - R) > 1e-14:
                raise ValueError("radial nodes must start at 0 and end at R")
            if np.any(np.diff(self.r) <= 0.0):
                raise ValueError("radial nodes must be strictly increasing")
        else:
            if N is None or int(N) < 2:
                raise ValueError("N must be at least 2")
            self.r = np.linspace(0.0, R, int(N) + 1)
        self.n_nodes = int(self.r.size)
        self.N = self.n_nodes - 1
        self.dr_face = np.diff(self.r)
        self.r_half = 0.5 * (self.r[:-1] + self.r[1:])
        a = np.empty(self.n_nodes)
        b = np.empty(self.n_nodes)
        a[0] = 0.0
        b[0] = self.r_half[0]
        a[-1] = self.r_half[-1]
        b[-1] = R
        if self.n_nodes > 2:
            a[1:-1] = self.r_half[:-1]
            b[1:-1] = self.r_half[1:]
        self.W = 0.5 * (b**2 - a**2)
        self.dr = float(self.dr_face.min())
        if abs(self.W.sum() - 0.5 * R**2) > 1e-14:
            raise RuntimeError("radial weights do not sum to R^2/2")

        self.M_heat = self._heat_matrix()
        self.M_heat_sparse = diags(
            [np.diag(self.M_heat, -1), np.diag(self.M_heat, 0), np.diag(self.M_heat, 1)],
            offsets=[-1, 0, 1],
            shape=self.M_heat.shape,
            format="csc",
        )
        self._heat_force_coef = (R * H) / (RHO * CP * self.W[-1])
        self._heat_force_index = self.n_nodes - 1

    def output_indices(self) -> np.ndarray:
        """Indices of 0, 0.1, ..., 2.0 cm nodes."""
        idx = np.empty(OUTPUT_RADII_M.size, dtype=int)
        for j, rc in enumerate(OUTPUT_RADII_M):
            k = int(np.argmin(np.abs(self.r - rc)))
            if abs(self.r[k] - rc) > 1e-12:
                raise ValueError(f"grid misses output radius {rc}, nearest {self.r[k]}")
            idx[j] = k
        return idx

    def _interior_face_coeff(self) -> np.ndarray:
        return self.r_half / self.dr_face

    def _heat_matrix(self) -> np.ndarray:
        n = self.n_nodes
        M = np.zeros((n, n))
        face = K * self._interior_face_coeff()
        den = RHO * CP * self.W
        M[0, 0] = -face[0] / den[0]
        M[0, 1] = face[0] / den[0]
        for i in range(1, n - 1):
            M[i, i - 1] = face[i - 1] / den[i]
            M[i, i + 1] = face[i] / den[i]
            M[i, i] = -(face[i - 1] + face[i]) / den[i]
        i = n - 1
        M[i, i - 1] = face[i - 1] / den[i]
        M[i, i] = (-face[i - 1] - R * H) / den[i]
        return M

    def heat_rhs(self, t: float, T: np.ndarray, Ta: float) -> np.ndarray:
        out = self.M_heat @ T
        out[self._heat_force_index] += self._heat_force_coef * Ta
        return out

    def heat_jac(self, t: float, T: np.ndarray):
        return self.M_heat_sparse

    def heat_jac_sparsity(self):
        n = self.n_nodes
        ones_m = np.ones(n - 1, dtype=float)
        ones_d = np.ones(n, dtype=float)
        return diags([ones_m, ones_d, ones_m], offsets=[-1, 0, 1], shape=(n, n), format="csc")

    def moisture_rhs(self, t: float, C: np.ndarray, Ca: float, D_const: float | None = None) -> np.ndarray:
        if D_const is None:
            D_node = moisture_diffusivity(C, protect=True)
        else:
            D_node = np.full_like(C, float(D_const))
        D_face = harmonic_mean(D_node[:-1], D_node[1:])
        F = -self.r_half * D_face * np.diff(C) / self.dr_face
        Cdot = np.empty_like(C)
        Cdot[0] = -F[0] / self.W[0]
        Cdot[1:-1] = (F[:-1] - F[1:]) / self.W[1:-1]
        F_surf = R * HM * (C[-1] - Ca)
        Cdot[-1] = (F[-1] - F_surf) / self.W[-1]
        return Cdot

    def moisture_jac_sparsity(self):
        return self.heat_jac_sparsity()

    def surface_heat_influx_density(self, T_s: float, Ta: float) -> float:
        """Heat into the solid per unit surface area, W/m^2."""
        return H * (Ta - T_s)

    def surface_moisture_outflux(self, C_s: float, Ca: float) -> float:
        """Outward moisture flux in the normalized (no rho_d) units, m/s * kg/kg."""
        return HM * (C_s - Ca)


def harmonic_mean(a: np.ndarray, b: np.ndarray) -> np.ndarray:
    out = np.zeros_like(a, dtype=float)
    both = (a > 0.0) & (b > 0.0)
    out[both] = 2.0 * a[both] * b[both] / (a[both] + b[both])
    return out
'''

# ========================================================================
# Module: reproduction_export.py
# ========================================================================
SOURCES['reproduction_export'] = r'''"""Portable Excel export from validated solver arrays; no external runtime."""
from copy import copy
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


def write_template(root, question, times, fields):
    root = Path(root)
    template = root / 'data' / '附件' / '附件3' / f'result{question}.xlsx'
    wb = load_workbook(template)
    if len(wb.worksheets) != len(fields):
        raise ValueError('Template sheet count does not match the output fields')
    times = np.asarray(times)
    if times.ndim != 1 or not np.isfinite(times).all() or not (np.diff(times) > 0).all():
        raise ValueError('Invalid output time sequence')
    for ws, values in zip(wb.worksheets, fields):
        values = np.asarray(values)
        expected = 22 if question == 4 else 21
        if values.shape != (len(times), expected):
            raise ValueError('Output shape does not match the requested radial grid')
        if np.isinf(values).any() or (question != 4 and np.isnan(values).any()):
            raise ValueError('Unexpected nonfinite output')
        header = ws.cell(1, 1).value
        surface = ws.cell(1, ws.max_column).value
        number_style = copy(ws.cell(2, 2)._style)
        header_style = copy(ws.cell(1, 2)._style)
        ws.delete_rows(1, ws.max_row)
        ws.append([header] + [j / 10 for j in range(21)] + ([surface] if question == 4 else []))
        for cell in ws[1]:
            cell._style = copy(header_style)
        for t, row in zip(times, values):
            ws.append([float(t)] + [None if np.isnan(v) else float(v) for v in row])
        for row in ws.iter_rows(min_row=2, min_col=2):
            for cell in row:
                cell._style = copy(number_style)
                cell.number_format = '0.0000'
        ws.freeze_panes = 'B2'
    wb.properties.creator = 'Team'
    wb.properties.lastModifiedBy = 'Team'
    dest = root / 'results' / f'result{question}.xlsx'
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    wb.close()
    print(f'Excel saved: results/result{question}.xlsx', flush=True)


def export_saved(root, question):
    root = Path(root)
    validation = json.loads((root / 'results' / f'q{question}_validation.json').read_text(encoding='utf-8'))
    if question == 4:
        if validation.get('pass') is not True or not validation.get('checks') or not all(v is True for v in validation['checks'].values()):
            raise RuntimeError('Q4 numerical validation did not pass')
    else:
        from delivery import require_delivery
        require_delivery(validation, root / 'results' / f'q{question}_export_validation.json')
        if question == 3:
            from q3_closeout import require_q3_delivery
            require_q3_delivery(validation, root / 'results' / 'q3_export_validation.json')
    with np.load(root / 'results' / f'q{question}_solution.npz') as z:
        if question in (1, 2):
            write_template(root, question, z['times'], [z['T_out'], z['C_out']])
        elif question == 3:
            write_template(root, question, z['excel_times'], [z['C_out']])
        else:
            write_template(root, question, z['times_s'][1:], [z['C_physical'][1:]])
'''

# ========================================================================
# Module: problem1.py
# ========================================================================
SOURCES['problem1'] = r'''"""Problem 1: radial FVM + BDF for preheating temperature and moisture."""
from __future__ import annotations

import json
import sys
import time
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from scipy.integrate import solve_ivp, quad

CODE_DIR = Path(__file__).resolve().parent
if str(CODE_DIR) not in sys.path:
    sys.path.insert(0, str(CODE_DIR))

from q1_analytic import BesselDuhamel, converged_reference
from delivery import make_gate, require_delivery
from radial_fvm import RadialFVM
from radial_fvm import graded_radial_nodes
from utils import (
    ALPHA,
    C0,
    CP,
    FIGURES_DIR,
    H,
    HM,
    K,
    L,
    R,
    RESULTS_DIR,
    RHO,
    T0,
    T_END,
    PchipOven,
    load_oven_table,
    moisture_diffusivity,
)

OUTPUT_TIMES = np.arange(1.0, T_END + 0.5, 1.0)
TABLE_TIMES = np.array([100, 300, 600, 900, 1200, 1500, 1800], dtype=float)
TABLE_RADII_CM = np.array([0.0, 0.5, 1.0, 1.5, 2.0])
OUTPUT_RADII_M = 0.001 * np.arange(21)  # 0, 0.1, ..., 2.0 cm
PAPER_TIMES = TABLE_TIMES.copy()


def load_q1_inputs() -> tuple[PchipOven, PchipOven, np.ndarray]:
    t, Ta, Ca = load_oven_table()
    mask = t <= T_END + 1e-12
    t, Ta, Ca = t[mask], Ta[mask], Ca[mask]
    if t.size != 31 or abs(t[-1] - T_END) > 1e-12:
        raise ValueError(f"expected 31 oven records on [0, 1800], got {t.size}, t_end={t[-1]}")
    if abs(Ta[0] - T0) > 1e-12:
        raise ValueError(f"Ta(0)={Ta[0]} != T0={T0}")
    Ta_fun = PchipOven(t, Ta)
    Ca_fun = PchipOven(t, Ca)
    for tj, yj in zip(t, Ta):
        if abs(float(Ta_fun(tj)) - float(yj)) > 1e-12:
            raise RuntimeError(f"PCHIP misses Ta knot t={tj}")
    for tj, yj in zip(t, Ca):
        if abs(float(Ca_fun(tj)) - float(yj)) > 1e-12:
            raise RuntimeError(f"PCHIP misses Ca knot t={tj}")
    return Ta_fun, Ca_fun, t


def integrate_mol(
    rhs,
    y0: np.ndarray,
    t_breaks: np.ndarray,
    *,
    rtol: float,
    atol: float,
    max_step: float,
    jac=None,
    jac_sparsity=None,
    field_name: str = "field",
):
    y = np.asarray(y0, dtype=float).copy()
    times = [0.0]
    states = [y.copy()]
    pieces: list[tuple[float, float, object]] = []
    t_end = float(t_breaks[-1])
    for i in range(len(t_breaks) - 1):
        t0 = float(t_breaks[i])
        t1 = float(min(t_breaks[i + 1], t_end))
        if t1 <= t0:
            continue
        t_eval = np.arange(np.floor(t0) + 1.0, np.floor(t1) + 1.0 + 1e-12, 1.0)
        t_eval = t_eval[(t_eval > t0 + 1e-14) & (t_eval <= t1 + 1e-14)]
        kwargs = dict(
            fun=rhs,
            t_span=(t0, t1),
            y0=y,
            method="BDF",
            rtol=rtol,
            atol=atol,
            max_step=max_step,
            dense_output=True,
            vectorized=False,
        )
        if t_eval.size:
            kwargs["t_eval"] = t_eval
        if jac is not None:
            kwargs["jac"] = jac
        if jac_sparsity is not None:
            kwargs["jac_sparsity"] = jac_sparsity
        sol = solve_ivp(**kwargs)
        if not sol.success:
            raise RuntimeError(f"{field_name} BDF failed on [{t0}, {t1}]: {sol.message}")
        pieces.append((t0, t1, sol))
        if sol.t.size:
            for k, tk in enumerate(sol.t):
                if times and abs(tk - times[-1]) < 1e-12:
                    states[-1] = sol.y[:, k].copy()
                    continue
                times.append(float(tk))
                states.append(sol.y[:, k].copy())
        y = sol.y[:, -1].copy() if sol.y.size else y
    t_arr = np.asarray(times, dtype=float)
    Y = np.stack(states, axis=0)
    return t_arr, Y, pieces


def dense_state(pieces, t: float) -> np.ndarray:
    t = float(t)
    for t0, t1, sol in pieces:
        if t0 - 1e-12 <= t <= t1 + 1e-12:
            return np.asarray(sol.sol(np.clip(t, t0, t1)), dtype=float)
    raise ValueError(f"time {t} not covered by dense output")


def extract_output_grid(fvm: RadialFVM, Y: np.ndarray) -> np.ndarray:
    idx = fvm.output_indices()
    return Y[:, idx]


def conservation_heat(fvm: RadialFVM, pieces, Ta_fun: PchipOven, T_init: np.ndarray, t: float) -> dict:
    T_t = dense_state(pieces, t)
    E = 2.0 * np.pi * L * RHO * CP * np.dot(fvm.W, T_t - T_init)
    flux = lambda x: float(Ta_fun(x)) - float(dense_state(pieces, x)[-1])
    Qraw, qerr = quad(flux, 0, t, points=Ta_fun.t[(Ta_fun.t > 0) & (Ta_fun.t < t)], epsabs=1e-9, epsrel=1e-10, limit=500)
    Q = 2.0 * np.pi * R * L * H * Qraw
    denom = max(abs(Q), abs(E), 1e-12)
    return {
        "E_relative_J": E,
        "Q_surface_J": Q,
        "relative_residual": abs(E - Q) / denom,
        "quadrature_abs_error": qerr, "quadrature": "adaptive QUADPACK on PCHIP intervals",
    }


def conservation_moisture(fvm: RadialFVM, pieces, Ca_fun: PchipOven, C_init: np.ndarray, t: float) -> dict:
    C_t = dense_state(pieces, t)
    MC = float(np.dot(fvm.W, C_t))
    MC0 = float(np.dot(fvm.W, C_init))
    flux = lambda x: HM * (float(dense_state(pieces, x)[-1]) - float(Ca_fun(x)))
    I, qerr = quad(flux, 0, t, points=Ca_fun.t[(Ca_fun.t > 0) & (Ca_fun.t < t)], epsabs=1e-14, epsrel=1e-10, limit=500)
    rhs = MC0 - R * I
    denom = max(abs(MC0 - MC), abs(R * I), 1e-18)
    return {
        "MC": MC,
        "MC0": MC0,
        "minus_R_int_hm": rhs,
        "relative_residual": abs(MC - rhs) / denom,
        "quadrature_abs_error": qerr, "quadrature": "adaptive QUADPACK on PCHIP intervals",
    }


def compare_on_output(A: np.ndarray, B: np.ndarray, times: np.ndarray) -> dict:
    diff = np.abs(A - B)
    k = int(np.argmax(diff))
    n_t, n_r = diff.shape
    i, j = divmod(k, n_r)
    return {
        "max_abs": float(diff.max()),
        "max_time_s": float(times[i]),
        "max_radius_cm": 0.1 * j,
        "mean_abs": float(diff.mean()),
        "paper_table_max_abs": float(np.abs(A[np.isin(times, TABLE_TIMES)][:, ::5] - B[np.isin(times, TABLE_TIMES)][:, ::5]).max())
        if A.shape[1] == 21
        else None,
    }


def four_decimal_stats(A: np.ndarray, B: np.ndarray, times: np.ndarray) -> dict:
    Ar = np.round(A, 4)
    Br = np.round(B, 4)
    changed = Ar != Br
    paper_t = np.isin(times, TABLE_TIMES)
    paper_r = slice(None, None, 5) if A.shape[1] == 21 else slice(None)
    paper_changed = Ar[paper_t][:, paper_r] != Br[paper_t][:, paper_r]
    early = times <= 100.0
    return {
        "changed_cells": int(np.sum(changed)),
        "total_cells": int(A.size),
        "paper_changed_cells": int(np.sum(paper_changed)),
        "paper_total_cells": int(paper_changed.size),
        "early100_changed_cells": int(np.sum(changed[early])),
        "max_abs": float(np.max(np.abs(A - B))),
        "paper_table_max_abs": float(np.max(np.abs(A[paper_t][:, paper_r] - B[paper_t][:, paper_r]))),
    }


def solve_temperature(fvm: RadialFVM, Ta_fun: PchipOven, t_breaks, rtol, atol, max_step):
    T_init = np.full(fvm.n_nodes, T0)
    def rhs(t, T):
        return fvm.heat_rhs(t, T, float(Ta_fun(t)))
    def jac(t, T):
        return fvm.heat_jac(t, T)
    return integrate_mol(
        rhs, T_init, t_breaks,
        rtol=rtol, atol=atol, max_step=max_step,
        jac=jac, jac_sparsity=fvm.heat_jac_sparsity(),
        field_name="temperature",
    )


def solve_moisture(fvm: RadialFVM, Ca_fun: PchipOven, t_breaks, rtol, atol, max_step, D_const=None):
    C_init = np.full(fvm.n_nodes, C0)
    def rhs(t, C):
        return fvm.moisture_rhs(t, C, float(Ca_fun(t)), D_const=D_const)
    return integrate_mol(
        rhs, C_init, t_breaks,
        rtol=rtol, atol=atol, max_step=max_step,
        jac_sparsity=fvm.moisture_jac_sparsity(),
        field_name="moisture",
    )


def assert_physical(T_out, C_out, Ta_fun, Ca_fun) -> dict:
    Ta_max = float(np.max(Ta_fun.y))
    Ta_min = float(np.min(Ta_fun.y))
    Ca_min = float(np.min(Ca_fun.y))
    info = {
        "T_min": float(T_out.min()),
        "T_max": float(T_out.max()),
        "C_min": float(C_out.min()),
        "C_max": float(C_out.max()),
        "T_below_28": float(max(0.0, 28.0 - T_out.min())),
        "T_above_Ta_max": float(max(0.0, T_out.max() - Ta_max)),
        "C_above_C0": float(max(0.0, C_out.max() - C0)),
        "C_below_Ca_min": float(max(0.0, Ca_min - C_out.min())),
        "Ta_range": [Ta_min, Ta_max],
        "Ca_range": [Ca_min, float(np.max(Ca_fun.y))],
    }
    info["T_bounds_ok"] = info["T_below_28"] <= 1e-6 and info["T_above_Ta_max"] <= 1e-6
    info["C_bounds_ok"] = info["C_above_C0"] <= 1e-8 and info["C_below_Ca_min"] <= 1e-8
    info["C_nonnegative"] = bool(C_out.min() >= -1e-8)
    return info


def write_result_workbook(path: Path, T_out: np.ndarray, C_out: np.ndarray, times: np.ndarray):
    radii_cm = 0.1 * np.arange(21)
    wb = Workbook()
    for title, arr in (("温度", T_out), ("水分浓度", C_out)):
        ws = wb.create_sheet(title)
        ws["A1"] = "时间\\到药材中心的距离"
        for j, rc in enumerate(radii_cm, start=2):
            ws.cell(1, j, float(rc))
        for i, t in enumerate(times):
            ws.cell(i + 2, 1, int(round(float(t))))
            for j in range(21):
                cell = ws.cell(i + 2, j + 2, float(arr[i, j]))
                cell.number_format = "0.0000"
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_paper_csv(path: Path, arr_out: np.ndarray, times: np.ndarray, radii_cm=TABLE_RADII_CM):
    t_idx = {float(t): i for i, t in enumerate(times)}
    r_idx = {round(0.1 * j, 10): j for j in range(21)}
    lines = ["t_s," + ",".join(f"r_{rc:g}cm" for rc in radii_cm)]
    for t in TABLE_TIMES:
        row = [f"{t:.0f}"]
        i = t_idx[float(t)]
        for rc in radii_cm:
            j = r_idx[float(rc)]
            row.append(f"{arr_out[i, j]:.8f}")
        lines.append(",".join(row))
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def setup_mpl():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({
        "font.sans-serif": ["Microsoft YaHei", "SimHei", "Noto Sans CJK SC", "DejaVu Sans"],
        "axes.unicode_minus": False,
        "font.size": 9,
        "axes.labelsize": 10,
        "legend.fontsize": 8,
        "xtick.labelsize": 8,
        "ytick.labelsize": 8,
        "axes.linewidth": 0.8,
        "pdf.fonttype": 42,
        "ps.fonttype": 42,
    })
    return plt


def plot_figures(plt, fvm, times, T_full, C_full, T_out, C_out, Ta_fun, Ca_fun):
    FIGURES_DIR.mkdir(parents=True, exist_ok=True)
    r_cm = fvm.r * 100.0
    r_out_cm = OUTPUT_RADII_M * 100.0
    t_marks = PAPER_TIMES
    idx = {float(t): i for i, t in enumerate(times)}

    fig, ax = plt.subplots(figsize=(6.4, 3.8))
    for t in t_marks:
        ax.plot(r_cm, T_full[idx[float(t)]], lw=1.2, label=fr"$t={int(t)}\,\mathrm{{s}}$")
    ax.set_xlabel("半径 $r$ (cm)")
    ax.set_ylabel("温度 $T$ (°C)")
    ax.set_xlim(0, 2)
    ax.legend(frameon=False, ncol=2)
    fig.tight_layout()
    fig.savefig(FIGURES_DIR / "q1_T_profiles.pdf")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(6.4, 3.8))
    for t in t_marks:
        ax.plot(r_cm, C_full[idx[float(t)]], lw=1.2, label=fr"$t={int(t)}\,\mathrm{{s}}$")
    ax.set_xlabel("半径 $r$ (cm)")
    ax.set_ylabel(r"干基含水率 $C$ (kg/kg)")
    ax.set_xlim(0, 2)
    ax.legend(frameon=False, ncol=2)
    fig.tight_layout()
    fig.savefig(FIGURES_DIR / "q1_C_profiles.pdf")
    plt.close(fig)

    t_plot = np.concatenate([[0.0], times])
    T_c = np.concatenate([[T0], T_out[:, 0]])
    T_s = np.concatenate([[T0], T_out[:, -1]])
    Ta = np.array([Ta_fun(t) for t in t_plot])
    fig, ax = plt.subplots(figsize=(6.4, 3.8))
    ax.plot(t_plot, Ta, lw=1.0, ls="--", label="烘房 $T_a$")
    ax.plot(t_plot, T_s, lw=1.2, label="表面")
    ax.plot(t_plot, T_c, lw=1.2, label="中心")
    ax.set_xlabel("时间 $t$ (s)")
    ax.set_ylabel("温度 (°C)")
    ax.set_xlim(0, 1800)
    ax.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(FIGURES_DIR / "q1_T_center_surface.pdf")
    plt.close(fig)

    C_c = np.concatenate([[C0], C_out[:, 0]])
    C_s = np.concatenate([[C0], C_out[:, -1]])
    fig, ax = plt.subplots(figsize=(6.4, 3.8))
    ax.plot(t_plot, C_s, lw=1.2, label="表面")
    ax.plot(t_plot, C_c, lw=1.2, label="中心")
    ax.set_xlabel("时间 $t$ (s)")
    ax.set_ylabel(r"干基含水率 $C$ (kg/kg)")
    ax.set_xlim(0, 1800)
    ax.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(FIGURES_DIR / "q1_C_center_surface.pdf")
    plt.close(fig)

    fig, ax1 = plt.subplots(figsize=(6.4, 3.6))
    ax2 = ax1.twinx()
    t_dense = np.linspace(0.0, T_END, 1801)
    ax1.plot(t_dense, np.array([Ta_fun(t) for t in t_dense]), color="C0", lw=1.2, label="烘房温度")
    ax2.plot(t_dense, np.array([Ca_fun(t) for t in t_dense]), color="C1", lw=1.2, ls="--", label="烘房水分")
    ax1.scatter(Ta_fun.t, Ta_fun.y, s=14, color="C0", zorder=3)
    ax2.scatter(Ca_fun.t, Ca_fun.y, s=14, color="C1", zorder=3)
    ax1.set_xlabel("时间 $t$ (s)")
    ax1.set_ylabel("烘房温度 (°C)", color="C0")
    ax2.set_ylabel("烘房水分 (kg/kg)", color="C1")
    ax1.set_xlim(0, 1800)
    h1, l1 = ax1.get_legend_handles_labels()
    h2, l2 = ax2.get_legend_handles_labels()
    ax1.legend(h1 + h2, l1 + l2, frameon=False)
    fig.tight_layout()
    fig.savefig(FIGURES_DIR / "q1_oven_input.pdf")
    plt.close(fig)
    _ = r_out_cm


def main():
    t_wall0 = time.perf_counter()
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    Ta_fun, Ca_fun, knots = load_q1_inputs()
    # PCHIP is C1; integrate as one span rather than 60 s linear kinks.
    t_breaks = np.array([0.0, T_END], dtype=float)
    print(
        "oven points", knots.size,
        "kind", Ta_fun.kind,
        "Ta(1800)=", Ta_fun(T_END),
        "Ca(1800)=", Ca_fun(T_END),
    )
    print("alpha", ALPHA, "D(C0)", moisture_diffusivity(C0), "Bi_T", H * R / K)

    mesh_levels = [2, 4, 8]
    rtol, atol_T, atol_C, max_step = 1e-10, 1e-10, 1e-12, 0.5
    solutions = {}
    for level in mesh_levels:
        r = graded_radial_nodes(level)
        fvm = RadialFVM(r=r)
        tag = f"G{level}"
        print(
            f"\n=== {tag} nodes={fvm.n_nodes} surface_dr={fvm.dr_face[-1]:.4e} m ==="
        )
        print(f"=== {tag} temperature ===")
        tT, YT, piecesT = solve_temperature(fvm, Ta_fun, t_breaks, rtol, atol_T, max_step)
        print(f"=== {tag} moisture ===")
        tC, YC, piecesC = solve_moisture(fvm, Ca_fun, t_breaks, rtol, atol_C, max_step)
        if not np.allclose(tT, tC):
            raise RuntimeError("temperature and moisture output times differ")
        times = tT[tT > 0]
        YT_pos = YT[tT > 0]
        YC_pos = YC[tC > 0]
        T_out = extract_output_grid(fvm, YT_pos)
        C_out = extract_output_grid(fvm, YC_pos)
        solutions[level] = dict(
            tag=tag,
            level=level,
            fvm=fvm,
            times=times,
            YT=YT_pos,
            YC=YC_pos,
            T_out=T_out,
            C_out=C_out,
            piecesT=piecesT,
            piecesC=piecesC,
            T_init=np.full(fvm.n_nodes, T0),
            C_init=np.full(fvm.n_nodes, C0),
            n_nodes=fvm.n_nodes,
            surface_dr=float(fvm.dr_face[-1]),
        )
        print(
            "done", tag,
            "T range", T_out.min(), T_out.max(),
            "C range", C_out.min(), C_out.max(),
            "C(t=1,s)", C_out[0, -1],
            "C(t=100,s)", C_out[99, -1],
        )

    conv = {}
    fourdec = {}
    for coarse, fine in ((2, 4), (4, 8)):
        t = solutions[fine]["times"]
        conv[f"T_G{coarse}_vs_G{fine}"] = compare_on_output(
            solutions[coarse]["T_out"], solutions[fine]["T_out"], t
        )
        conv[f"C_G{coarse}_vs_G{fine}"] = compare_on_output(
            solutions[coarse]["C_out"], solutions[fine]["C_out"], t
        )
        fourdec[f"T_G{coarse}_vs_G{fine}"] = four_decimal_stats(
            solutions[coarse]["T_out"], solutions[fine]["T_out"], t
        )
        fourdec[f"C_G{coarse}_vs_G{fine}"] = four_decimal_stats(
            solutions[coarse]["C_out"], solutions[fine]["C_out"], t
        )
        print("conv T", coarse, fine, conv[f"T_G{coarse}_vs_G{fine}"])
        print("conv C", coarse, fine, conv[f"C_G{coarse}_vs_G{fine}"])
        print("4dp T", coarse, fine, fourdec[f"T_G{coarse}_vs_G{fine}"])
        print("4dp C", coarse, fine, fourdec[f"C_G{coarse}_vs_G{fine}"])

    fd_T = fourdec["T_G4_vs_G8"]
    fd_C = fourdec["C_G4_vs_G8"]
    paper_4dp_stable = (
        fd_T["paper_changed_cells"] == 0 and fd_C["paper_changed_cells"] == 0
    )
    full_4dp_stable = fd_T["changed_cells"] == 0 and fd_C["changed_cells"] == 0
    abs_2e5 = (
        conv["T_G4_vs_G8"]["max_abs"] <= 2e-5
        and conv["C_G4_vs_G8"]["max_abs"] <= 2e-5
    )
    pub_level = 4 if paper_4dp_stable and full_4dp_stable else 8
    pub = solutions[pub_level]
    fvm = pub["fvm"]
    times = pub["times"]
    print(
        "publishing", pub["tag"],
        "n_nodes", pub["n_nodes"],
        "paper_4dp_stable", paper_4dp_stable,
        "full_4dp_stable", full_4dp_stable,
        "abs_2e5", abs_2e5,
    )

    print("time refinement on published mesh")
    tT2, YT2, _ = solve_temperature(fvm, Ta_fun, t_breaks, rtol / 5, atol_T / 5, 0.25)
    tC2, YC2, _ = solve_moisture(fvm, Ca_fun, t_breaks, rtol / 5, atol_C / 5, 0.25)
    T2 = extract_output_grid(fvm, YT2[tT2 > 0])
    C2 = extract_output_grid(fvm, YC2[tC2 > 0])
    time_sens = {
        "T": compare_on_output(pub["T_out"], T2, times),
        "C": compare_on_output(pub["C_out"], C2, times),
    }
    time_4dp = {
        "T": four_decimal_stats(pub["T_out"], T2, times),
        "C": four_decimal_stats(pub["C_out"], C2, times),
    }
    print("time sens", time_sens)
    print("time 4dp", time_4dp)

    print("Bessel temperature reference")
    Bi = H * R / K
    T_ref, heat_truncation = converged_reference(Bi, ALPHA, R, OUTPUT_RADII_M, times, Ta_fun, T0, start_terms=120)
    bessel_T = compare_on_output(pub["T_out"], T_ref, times)
    print("Bessel T", bessel_T)

    print("frozen-D moisture probe")
    D0 = float(moisture_diffusivity(C0))
    Bi_m = HM * R / D0
    C_ref, moisture_truncation = converged_reference(Bi_m, D0, R, OUTPUT_RADII_M, times, Ca_fun, C0)
    tCf, YCf, _ = solve_moisture(fvm, Ca_fun, t_breaks, rtol, atol_C, max_step, D_const=D0)
    C_frozen = extract_output_grid(fvm, YCf[tCf > 0])
    frozen = compare_on_output(C_frozen, C_ref, times)
    print("frozen D vs analytic", frozen)
    nl_vs_frozen = compare_on_output(pub["C_out"], C_frozen, times)
    print("nonlinear vs frozen D", nl_vs_frozen)
    D_field = moisture_diffusivity(pub["YC"])
    D_stats = {
        "D0": D0,
        "D_min": float(np.min(D_field)),
        "D_max": float(np.max(D_field)),
        "D_min_over_D0": float(np.min(D_field) / D0),
        "D_max_over_D0": float(np.max(D_field) / D0),
        "C_min_field": float(np.min(pub["YC"])),
        "C_max_field": float(np.max(pub["YC"])),
        "surface_C_min": float(np.min(pub["C_out"][:, -1])),
        "center_C_end": float(pub["C_out"][-1, 0]),
        "surface_C_end": float(pub["C_out"][-1, -1]),
        "surface_C_t1": float(pub["C_out"][0, -1]),
        "surface_C_t100": float(pub["C_out"][99, -1]),
    }
    print("D stats", D_stats)

    bounds = assert_physical(pub["YT"], pub["YC"], Ta_fun, Ca_fun)
    print("bounds", bounds)

    cons_T = conservation_heat(fvm, pub["piecesT"], Ta_fun, pub["T_init"], T_END)
    cons_C = conservation_moisture(fvm, pub["piecesC"], Ca_fun, pub["C_init"], T_END)
    print("conservation T", cons_T)
    print("conservation C", cons_C)

    delivery = make_gate({
        "space_abs": abs_2e5,
        "space_paper": paper_4dp_stable,
        "all_finite": np.isfinite(pub["YT"]).all() and np.isfinite(pub["YC"]).all(),
        "heat_reference_truncation": heat_truncation["passed"],
        "heat_reference_error": bessel_T["max_abs"] <= 2e-5,
        "time_abs": time_sens["T"]["max_abs"] <= 2e-5 and time_sens["C"]["max_abs"] <= 2e-5,
        "time_paper": time_4dp["T"]["paper_changed_cells"] == 0 and time_4dp["C"]["paper_changed_cells"] == 0,
        "temperature_bounds": bounds["T_bounds_ok"],
        "moisture_bounds": bounds["C_bounds_ok"],
        "moisture_reference_truncation": moisture_truncation["passed"],
        "frozen_reference_error": frozen["max_abs"] <= 2e-5,
        "heat_balance": cons_T["relative_residual"] <= 1e-6,
        "moisture_balance": cons_C["relative_residual"] <= 1e-6,
    })
    print("delivery", delivery)

    paper_T = {
        f"t{int(t)}_r{rc:g}": float(pub["T_out"][int(t) - 1, int(round(rc / 0.1))])
        for t in TABLE_TIMES for rc in TABLE_RADII_CM
    }
    paper_C = {
        f"t{int(t)}_r{rc:g}": float(pub["C_out"][int(t) - 1, int(round(rc / 0.1))])
        for t in TABLE_TIMES for rc in TABLE_RADII_CM
    }

    validation = {
        "result_version": "q1-closeout-v1",
        "interpolation": Ta_fun.kind,
        "published_mesh": pub["tag"],
        "published_n_nodes": pub["n_nodes"],
        "published_surface_dr_m": pub["surface_dr"],
        "mesh_levels": mesh_levels,
        "mesh_n_nodes": {f"G{lv}": solutions[lv]["n_nodes"] for lv in mesh_levels},
        "solver": "FVM+BDF",
        "rtol": rtol,
        "atol_T": atol_T,
        "atol_C": atol_C,
        "max_step": max_step,
        "t_breaks": t_breaks.tolist(),
        "grid_convergence": conv,
        "four_decimal_grid": fourdec,
        "time_sensitivity": time_sens,
        "four_decimal_time": time_4dp,
        "delivery_gate": delivery,
        "reference_truncation_T": heat_truncation,
        "bessel_temperature": bessel_T,
        "reference_truncation_C": moisture_truncation,
        "frozen_D_moisture_vs_analytic": frozen,
        "nonlinear_C_vs_frozen_D": nl_vs_frozen,
        "D_variation": D_stats,
        "conservation_T_1800s": cons_T,
        "conservation_C_1800s": cons_C,
        "physical_bounds": bounds,
        "paper_table_T": paper_T,
        "paper_table_C": paper_C,
        "elapsed_s": time.perf_counter() - t_wall0,
        "python": sys.version,
    }
    require_delivery(validation, RESULTS_DIR / "diagnostics/q1_delivery/latest.json")
    np.savez(
        RESULTS_DIR / "q1_solution.npz",
        N=pub["n_nodes"] - 1,
        mesh_level=pub_level,
        n_nodes=pub["n_nodes"],
        times=times,
        r=fvm.r,
        checkpoint_times=np.r_[0., TABLE_TIMES],
        T_checkpoints=np.vstack([pub["T_init"], pub["YT"][TABLE_TIMES.astype(int)-1]]),
        C_checkpoints=np.vstack([pub["C_init"], pub["YC"][TABLE_TIMES.astype(int)-1]]),
        T_full=pub["YT"],
        C_full=pub["YC"],
        T_out=pub["T_out"],
        C_out=pub["C_out"],
        output_radii_m=OUTPUT_RADII_M,
        interpolation=np.array(Ta_fun.kind),
        rtol=rtol,
        atol_T=atol_T,
        atol_C=atol_C,
        max_step=max_step,
        surface_dr=pub["surface_dr"],
    )
    write_result_workbook(RESULTS_DIR / "result1.xlsx", pub["T_out"], pub["C_out"], times)
    write_paper_csv(RESULTS_DIR / "q1_table_temperature.csv", pub["T_out"], times)
    write_paper_csv(RESULTS_DIR / "q1_table_moisture.csv", pub["C_out"], times)

    plt = setup_mpl()
    plot_figures(plt, fvm, times, pub["YT"], pub["YC"], pub["T_out"], pub["C_out"], Ta_fun, Ca_fun)

    (RESULTS_DIR / "q1_validation.json").write_text(
        json.dumps(validation, ensure_ascii=False, indent=2, default=float),
        encoding="utf-8",
    )
    print("elapsed_s", validation["elapsed_s"])
    print("wrote results to", RESULTS_DIR)



if __name__ == "__main__":
    main()
'''

def prepare():
    code_dir = PACKAGE_ROOT / '_runtime' / f'Q{QUESTION}' / 'code'
    code_dir.mkdir(parents=True, exist_ok=True)
    for name, source in SOURCES.items():
        (code_dir / (name + '.py')).write_text(source, encoding='utf-8')
    sys.path.insert(0, str(code_dir))
    for folder in ('results', 'figures', 'reports'):
        (PACKAGE_ROOT / folder).mkdir(exist_ok=True)
    os.environ.setdefault('MPLBACKEND', 'Agg')
    for name in SOURCES:
        importlib.import_module(name)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--check', action='store_true', help='Check imports and original inputs without solving')
    parser.add_argument('--export-only', action='store_true', help='Export existing validated NPZ arrays; does not solve the model')
    args = parser.parse_args()
    prepare()
    from utils import load_oven_table
    load_oven_table()
    if QUESTION == 4:
        from problem4 import Inputs
        Inputs()
    if args.check:
        print(f'Q{QUESTION}: imports and input checks passed', flush=True)
        return
    from reproduction_export import export_saved
    if args.export_only:
        export_saved(PACKAGE_ROOT, QUESTION)
        return
    if QUESTION == 1:
        import problem1
        problem1.main()
    elif QUESTION == 2:
        import problem2
        problem2.main()
    elif QUESTION == 3:
        # The original delivery gate compares Q3 with independently computed Q2.
        if not (PACKAGE_ROOT / 'results' / 'q2_solution.npz').exists():
            print('Q3: computing Q2 first for the original overlap check', flush=True)
            import problem2
            problem2.main()
            export_saved(PACKAGE_ROOT, 2)
        import q3_closeout
        q3_closeout.main(['--no-xlsx'])
    else:
        import problem4
        import verify_q4_results
        for level, tight in ((2, False), (4, False), (4, True), (8, False), (16, False), (8, True)):
            problem4.run(level, tight)
        verify_q4_results.main()
    export_saved(PACKAGE_ROOT, QUESTION)
    print(f'Q{QUESTION}: full recomputation and Excel export completed', flush=True)


if __name__ == '__main__':
    main()
